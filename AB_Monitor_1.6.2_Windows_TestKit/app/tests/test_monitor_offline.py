#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-License-Identifier: GPL-3.0-or-later
# SPDX-FileCopyrightText: 2026 William W. Wallace <wwallace@nrao.edu>
# %% ===========================================================================
#    MODULE HEADER
# =============================================================================
"""Offline tests of ab_power_meter_monitor 1.6.1 against the meter emulator.

===============================================================================
 Title          : AB Power Meter Monitor - offline integration tests
 Filename       : tests/test_monitor_offline.py
 Version        : 0.0.2 (Semantic Versioning 2.0.0, MAJOR.MINOR.PATCH)
 Author         : William W. Wallace
 Company        : NRAO/GBO
 Email          : wwallace@nrao.edu
 Created        : 2026-10-26
 Modified       : 2026-10-26
 Language       : Python 3.8 through 3.12
 Requires       : the monitor's dependencies (numpy, h5py, astropy, openpyxl,
                  psutil); pytest optional (also runs as a plain script)
 Licence        : GPL-3.0-or-later (GNU General Public License v3.0 or later)
 Status         : Internal draft
===============================================================================

What is checked (each test resets the module state first)
--------------------------------------------------------
1  Every enabled output holds exactly one row per successful poll, across a
   RAM flush in the middle of the run (no loss, no duplicates); the .vszh5
   columns equal the CSV columns value for value.
2  A second run appending to the same files (store/cursors reset = new
   process) continues the files - the 1.5.x "stops after the first poll"
   bug is gone - and append_files = 0 starts files afresh.
3  FITS: one HDU per table with its full name and only its own rows (the
   1.5.x EXTNAME collision multiplied rows).
4  Duplicate parameter names in one table are all kept ("Reserved [#3]").
5  RAM-flush triggers: store limit always clears; system-RAM % only when the
   store is big enough; minimum interval respected; keep-tail honoured.
6  .vszh5 locked by a reader: a normal poll keeps the rows for later; a
   flush writes a rescue document; afterwards the main document catches up.
7  A forced persist (flush / Save Now / stop) that waits for a busy writer
   never hands it a row twice (1.6.1 race fix), checked by sequence number.
8  Optional (VEUSZ_EXE set): the .vszh5 loads and renders in real Veusz.

Runs on Linux and on Windows builds of Python (Windows Python 3.12 tested
under Wine).

Run:  python tests/test_monitor_offline.py     or     pytest tests/

REVISION HISTORY
----------------
0.0.2  2026-10-26  Busy-writer regression test (monitor 1.6.1).
0.0.1  2026-10-26  First version (monitor 1.6.0).
===============================================================================
"""

# %% IMPORTS ===================================================================
from __future__ import annotations

import csv
import glob
import logging
import os
import socket
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

import h5py
import numpy as np

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
import ab_power_meter_monitor as abm  # noqa: E402
from ab_meter_emulator import start_emulators  # noqa: E402

HOSTS = ["127.0.0.2", "127.0.0.3"]


def _free_port() -> int:
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        return probe.getsockname()[1]


PORT = _free_port()
start_emulators(HOSTS, PORT)


# %% HELPERS ===================================================================
def _reset(outputs=("csv", "md", "vszh5", "fits", "xlsx"), append: int = 1) -> dict:
    """Fresh module state, as if a new process started; returns a run cfg."""
    with abm._TS_LOCK:
        abm.TIME_SERIES_STORE.clear()
        abm._TS_META.clear()
    with abm._CURSOR_LOCK:
        abm._CURSORS.clear()
    abm._INFLIGHT.clear()
    abm._MD_HEADER_CACHE.clear()
    abm._RUN_STATE["active"] = False
    abm._FLUSH_STATE.update(count=0, last_mono=None, last_reason="")
    abm.METER_URL_TEMPLATE = "http://{ip}:%d/{page}" % PORT
    abm.IP_LIST = list(HOSTS)
    abm.HTTP_TIMEOUT_SEC = 2
    abm.HTTP_RETRY_COUNT = 1
    abm.ENABLE_RAM_FLUSH = 0
    abm.APPEND_OUTPUT_FILES = append
    abm.FULL_WRITE_EVERY_N_CYCLES = 3
    flags = {"csv": "ENABLE_CSV", "md": "ENABLE_LOG_APPEND", "vszh5": "ENABLE_VEUSZ",
             "fits": "ENABLE_FITS", "xlsx": "ENABLE_XLSX"}
    for name, flag in flags.items():
        setattr(abm, flag, 1 if name in outputs else 0)
    logging.getLogger("ABMonitor").setLevel(logging.WARNING)
    return abm.default_config()


def _use_dir(folder: str) -> None:
    """Point every output at ``folder`` (sub-folders follow OUTPUT_BASE_DIR)."""
    abm.OUTPUT_BASE_DIR = folder


def _csv_rows(path: str) -> list:
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.reader(handle))[2:]


def _md_rows(path: str) -> int:
    with open(path, encoding="utf-8") as handle:
        return sum(1 for line in handle if line.startswith("| ") and not line.startswith("| Timestamp_Local"))


def _vsz_lengths(path: str) -> dict:
    with h5py.File(path, "r") as hdf:
        return {name: hdf[f"Veusz/Data/{name}/data"].shape[0] for name in hdf["Veusz/Data"]}


# %% TESTS =====================================================================
def test_no_loss_no_duplicates_across_flush() -> str:
    folder = tempfile.mkdtemp(prefix="abm_t1_")
    _use_dir(folder)
    cfg = _reset()
    abm.begin_run(cfg)
    polls = 0
    for _ in range(4):
        assert abm.poll_cycle(cfg)["ok_ips"] == HOSTS
        polls += 1
    flush = abm.perform_ram_flush(cfg, "test mid-run flush")
    assert flush["rows_kept_unpersisted"] == 0 and abm.TIME_SERIES_STORE == {}, flush
    for _ in range(3):
        abm.poll_cycle(cfg)
        polls += 1
    abm.end_run(cfg)
    ip = "127_0_0_2"
    for table in abm.TABLE_NAMES.values():
        rows = _csv_rows(f"{folder}/csv/ABMeter_{ip}_{table}.csv")
        assert len(rows) == polls, (table, len(rows))
        assert _md_rows(f"{folder}/logs/ABMeter_{ip}_{table}.md") == polls, table
    lengths = _vsz_lengths(f"{folder}/veusz/ABMeter_{ip}.vszh5")
    assert lengths and set(lengths.values()) == {polls}, lengths
    # The .vszh5 column equals the CSV column value for value.
    rows = _csv_rows(f"{folder}/csv/ABMeter_{ip}_Real_Time_Power_Table.csv")
    with open(f"{folder}/csv/ABMeter_{ip}_Real_Time_Power_Table.csv", newline="") as handle:
        header = next(csv.reader(handle))
    column = header.index("Total Real Power")
    csv_values = np.array([float(r[column]) for r in rows])
    with h5py.File(f"{folder}/veusz/ABMeter_{ip}.vszh5", "r") as hdf:
        vsz_values = hdf["Veusz/Data/2_Real_Time_Power_Table_Total_Real_Power/data"][()]
    assert np.allclose(csv_values, vsz_values), (csv_values, vsz_values)
    from astropy.io import fits
    with fits.open(f"{folder}/fits/ABMeter_{ip}.fits") as hdul:
        assert all(len(h.data) == polls for h in hdul[1:]), [len(h.data) for h in hdul[1:]]
    import openpyxl
    book = openpyxl.load_workbook(f"{folder}/xlsx/ABMeter_{ip}.xlsx", read_only=True)
    assert all(sum(1 for _ in ws.iter_rows()) - 2 == polls for ws in book.worksheets)
    book.close()
    return folder


def test_second_run_appends_and_fresh_mode() -> None:
    folder = tempfile.mkdtemp(prefix="abm_t2_")
    _use_dir(folder)
    cfg = _reset()
    abm.begin_run(cfg)
    for _ in range(3):
        abm.poll_cycle(cfg)
    abm.end_run(cfg)
    cfg = _reset()                       # "new process", same files, append mode
    abm.begin_run(cfg)
    for _ in range(2):
        abm.poll_cycle(cfg)
    abm.end_run(cfg)
    path = f"{folder}/csv/ABMeter_127_0_0_3_Voltage_Current_Table.csv"
    assert len(_csv_rows(path)) == 5, len(_csv_rows(path))
    assert set(_vsz_lengths(f"{folder}/veusz/ABMeter_127_0_0_3.vszh5").values()) == {5}
    cfg = _reset(append=0)               # append_files = 0: start afresh once per run
    abm.begin_run(cfg)
    for _ in range(2):
        abm.poll_cycle(cfg)
    abm.end_run(cfg)
    assert len(_csv_rows(path)) == 2, len(_csv_rows(path))
    assert set(_vsz_lengths(f"{folder}/veusz/ABMeter_127_0_0_3.vszh5").values()) == {2}


def test_fits_tables_stay_separate() -> None:
    folder = tempfile.mkdtemp(prefix="abm_t3_")
    _use_dir(folder)
    cfg = _reset(outputs=("fits",))
    abm.begin_run(cfg)
    for _ in range(6):
        abm.poll_cycle(cfg)
        abm.write_fits(None, cfg["fits_dir"])       # rewrite every poll, as 1.5.x GUI did
    abm.end_run(cfg)
    from astropy.io import fits
    with fits.open(f"{folder}/fits/ABMeter_127_0_0_2.fits") as hdul:
        names = [h.header["EXTNAME"] for h in hdul[1:]]
        assert len(names) == len(set(names)) == 11, names
        assert all(len(h.data) == 6 for h in hdul[1:]), [len(h.data) for h in hdul[1:]]
        assert "EPOCH_UTC" in hdul["Real_Time_Power_Table"].data.names


def test_duplicate_parameter_names_kept() -> None:
    folder = tempfile.mkdtemp(prefix="abm_t4_")
    _use_dir(folder)
    cfg = _reset(outputs=("csv",))
    abm.begin_run(cfg)
    abm.poll_cycle(cfg)
    abm.end_run(cfg)
    columns = abm.TIME_SERIES_STORE["127.0.0.2"]["Diagnostic_Table"]["columns"]
    assert "Reserved" in columns and "Reserved [#3]" in columns, list(columns)
    assert columns["Reserved"] == [0.0] and columns["Reserved [#3]"] == [1.0]


def test_flush_triggers() -> None:
    folder = tempfile.mkdtemp(prefix="abm_t5_")
    _use_dir(folder)
    cfg = _reset(outputs=("csv",))
    abm.ENABLE_RAM_FLUSH = 1
    cfg = abm.default_config()
    cfg.update(mem_flush_threshold_mb=1e9, mem_free_min_mb=0, mem_proc_rss_limit_mb=0,
               mem_ram_pct_limit=1, mem_min_store_mb_for_clear=1.0, mem_flush_min_interval_sec=0)
    abm.begin_run(cfg)
    status = abm.poll_cycle(cfg)
    # RAM % is "exceeded" (limit 1 %) but the store is tiny: no clearing.
    assert status["flush"] is None and abm.TIME_SERIES_STORE, status
    cfg["mem_min_store_mb_for_clear"] = 0.0
    status = abm.poll_cycle(cfg)
    assert status["flush"] is not None and abm.TIME_SERIES_STORE == {}, status["flush"]
    # Minimum interval: no second automatic flush right away.
    cfg["mem_flush_min_interval_sec"] = 3600
    status = abm.poll_cycle(cfg)
    assert status["flush"] is None and abm.TIME_SERIES_STORE
    # Store-size trigger with keep-tail.
    cfg.update(mem_ram_pct_limit=0, mem_flush_threshold_mb=1e-6, mem_flush_min_interval_sec=0,
               ram_flush_keep_samples=2)
    abm.poll_cycle(cfg)
    abm.poll_cycle(cfg)
    rows = len(abm.TIME_SERIES_STORE["127.0.0.2"]["Real_Time_Power_Table"]["timestamps_local"])
    assert rows == 2, rows
    abm.end_run(cfg)
    # Nothing was written twice although rows were kept in memory.
    assert len(_csv_rows(f"{folder}/csv/ABMeter_127_0_0_2_Real_Time_Power_Table.csv")) == 5


def test_vszh5_locked_by_reader() -> None:
    folder = tempfile.mkdtemp(prefix="abm_t6_")
    _use_dir(folder)
    cfg = _reset(outputs=("vszh5",))
    abm.begin_run(cfg)
    abm.poll_cycle(cfg)
    path = f"{folder}/veusz/ABMeter_127_0_0_2.vszh5"
    # Another process holds the file open read-only (as Veusz does while loading).
    holder = subprocess.Popen([sys.executable, "-c",
                               "import h5py,sys,time; f=h5py.File(sys.argv[1],'r'); "
                               "print('open', flush=True); time.sleep(30)", path],
                              stdout=subprocess.PIPE, text=True)
    try:
        assert holder.stdout.readline().strip() == "open"
        status = abm.poll_cycle(cfg)                      # rows must be kept, not lost
        assert set(_vsz_lengths(path).values()) == {1}, status
        flush = abm.perform_ram_flush(cfg, "test while locked")   # must persist -> rescue
        rescue = glob.glob(f"{folder}/veusz/ABMeter_127_0_0_2_*_rescue.vszh5")
        assert rescue and flush["rows_kept_unpersisted"] == 0, (rescue, flush)
        assert set(_vsz_lengths(rescue[0]).values()) == {1}
    finally:
        holder.kill()
        holder.wait()
    abm.poll_cycle(cfg)
    abm.end_run(cfg)
    assert set(_vsz_lengths(path).values()) == {2}       # 1 before + 1 after the lock


def test_busy_writer_never_gets_rows_twice() -> None:
    """1.6.1 race fix: a forced persist that waits for a busy writer resends nothing.

    The XLSX writer is made slower than a poll, so every poll finds it busy and
    every forced persist (as at flush / Save Now / stop) has to wait for it.
    Each sequence number must reach the writer exactly once.
    """
    folder = tempfile.mkdtemp(prefix="abm_t7_")
    _use_dir(folder)
    cfg = _reset(outputs=("csv", "xlsx"))
    cfg["full_write_every_n_cycles"] = 1            # XLSX dispatched in the background every poll
    real_writer = abm._write_xlsx_batches
    seen: dict = {}
    repeated: list = []
    guard = threading.Lock()

    def slow_writer(batches, out_dir, run_cfg, must_persist):
        with guard:                                  # record every seq handed to the writer
            for key, batch in batches.items():
                known = seen.setdefault(key, set())
                repeated.extend((key, s) for s in batch["seq"] if s in known)
                known.update(batch["seq"])
        time.sleep(0.8)                              # slower than one emulator poll
        return real_writer(batches, out_dir, run_cfg, must_persist)

    abm._write_xlsx_batches = slow_writer
    try:
        abm.begin_run(cfg)
        polls = 4
        for _ in range(polls):
            abm.poll_cycle(cfg)                      # finds the XLSX writer busy
            abm.persist_outputs(cfg, full=True, must_persist=True, reason="test forced")
        abm.end_run(cfg)
    finally:
        abm._write_xlsx_batches = real_writer
    assert not repeated, repeated[:10]
    import openpyxl
    book = openpyxl.load_workbook(f"{folder}/xlsx/ABMeter_127_0_0_2.xlsx", read_only=True)
    counts = [sum(1 for _ in ws.iter_rows()) - 2 for ws in book.worksheets]
    book.close()
    assert counts and set(counts) == {polls}, counts
    assert len(_csv_rows(f"{folder}/csv/ABMeter_127_0_0_2_Real_Time_Power_Table.csv")) == polls


def test_render_in_real_veusz(folder: str = "") -> None:
    exe = os.environ.get("VEUSZ_EXE", "")
    if not exe:
        print("  (VEUSZ_EXE not set - real-Veusz render skipped)")
        return
    import vszh5_writer as vw
    folder = folder or test_no_loss_no_duplicates_across_flush()
    document = f"{folder}/veusz/ABMeter_127_0_0_2.vszh5"
    out = vw.render_pages(document, outdir=f"{folder}/render", png_pages="all",
                          veusz_exe=exe, method="program", dpi=60, timeout=300)
    pages = vw.page_names(document)
    pngs = [p for p in out if p.suffix == ".png"]
    assert len(pngs) == len(pages) and any(p.suffix == ".pdf" for p in out), (len(pngs), len(pages))
    print(f"  real Veusz rendered {len(pngs)} page(s) + PDF in {folder}/render")


# %% RUN AS A SCRIPT ===========================================================
if __name__ == "__main__":
    started = time.monotonic()
    first_folder = test_no_loss_no_duplicates_across_flush()
    print("PASS test_no_loss_no_duplicates_across_flush")
    for test in (test_second_run_appends_and_fresh_mode, test_fits_tables_stay_separate,
                 test_duplicate_parameter_names_kept, test_flush_triggers,
                 test_vszh5_locked_by_reader, test_busy_writer_never_gets_rows_twice):
        test()
        print("PASS", test.__name__)
    test_render_in_real_veusz(first_folder)
    print("PASS test_render_in_real_veusz")
    print(f"all tests passed in {time.monotonic() - started:.1f} s")
