#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-License-Identifier: GPL-3.0-or-later
# SPDX-FileCopyrightText: 2026 William W. Wallace <wwallace@nrao.edu>
# %% ===========================================================================
#    MODULE HEADER
# =============================================================================
"""Allen-Bradley site power-meter web-table poller, data logger and GUI.

===============================================================================
 Title          : AB Power Meter Monitor (NRAO / GBO site infrastructure)
 Filename       : ab_power_meter_monitor.py
 Version        : 1.6.2 (Semantic Versioning 2.0.0, MAJOR.MINOR.PATCH)
 Author         : William W. Wallace
 Company        : NRAO/GBO
 Email          : wwallace@nrao.edu   (alt: naval.antennas@gmail.com)
 Phone          : +1 (304) 456-2216
 Created        : 2026-05 (first internal release; see git history)
 Modified       : 2026-10-26
 Language       : Python 3.12 (primary, tested); 3.8 - 3.11 supported.
                  The source stays at the Python 3.7 language level, so a
                  3.7.11 host still runs it with requirements-py37.txt pins.
 Requires       : requests, beautifulsoup4, lxml                  (always)
                  psutil                     (memory protection; recommended)
                  numpy + h5py + vszh5_writer.py   (Veusz .vszh5, ENABLE_VEUSZ)
                  numpy + astropy            (FITS, ENABLE_FITS)
                  openpyxl                   (XLSX, ENABLE_XLSX)
                  PySide6 + qtpy + matplotlib     (GUI, ENABLE_GUI)
                  Veusz itself is NOT needed to write files - only to view.
 Companion      : vszh5_writer.py >= 0.0.18 (same folder), ab_meter_caller.py,
                  ab_prometheus_exporter.py, ab_stop.py
 Licence        : GPL-3.0-or-later (GNU General Public License v3.0 or later)
 Copyright      : Copyright (C) 2026 William W. Wallace
 Status         : Release candidate - verified offline against the AB meter
                  emulator in tests/, against real Veusz 4.2.1, on Python 3.12
===============================================================================

DESCRIPTION
-----------
Polls Allen-Bradley (Rockwell 1403) power meters over HTTP, parses the 11 HTML
data tables (pages 0-10) of every meter, keeps the numeric values as a
columnar time series in memory (TIME_SERIES_STORE) and persists them to any
combination of

  * CSV            one file per device per table, appended every poll
  * Markdown       one GFM pipe-table log per device per table, appended
  * Veusz .vszh5   one document per device, APPENDED IN PLACE every poll
                   (vszh5_writer.VeuszAppender - no Veusz install needed)
  * FITS           one NRAO-style file per device (rewritten on a cadence)
  * XLSX           one workbook per device (rewritten on a cadence)

and shows live previews in a PySide6 GUI (qtpy), or runs headless as a
service. A Prometheus exporter (ab_prometheus_exporter.py) and a Python
caller (ab_meter_caller.py) use this module as a library.

Device URL: METER_URL_TEMPLATE, default "http://{ip}/{page}" (page 0 ... 10).
Each meter is accessed SERIALLY (its web server is single threaded); different
meters are polled in parallel, one thread per IP address.

MEMORY AND PERSISTENCE MODEL (1.6.0)
------------------------------------
Every stored row gets a process-wide sequence number. Each output writer
remembers the last sequence number it persisted, per device and table, so a
poll persists only the rows that are new to that writer - O(new rows), never
O(file size) - and nothing is written twice or skipped, whether the store is
cleared, trimmed or the program restarts.

A RAM flush (automatic or "Flush RAM Now") first persists every enabled
output, waits for the writes to finish, and only then drops the rows that
EVERY enabled output has persisted (optionally keeping the newest N rows for
the GUI preview). The freed memory is then handed back to the operating
system (MEM_RETURN_TO_OS, 1.6.1): Linux glibc malloc_trim(0) [R4]; Windows
HeapSetInformation(HeapOptimizeResources) [R8] decommits free heap memory and
EmptyWorkingSet() [R10] trims the working set (MEM_TRIM_WORKING_SET), so Task
Manager's "Memory" column drops after a flush.
If no file output is enabled, the rows are dropped, or spilled to a pickle
first when RAM_FLUSH_SPILL_PICKLE = 1.

Flush triggers (any one, when ENABLE_RAM_FLUSH = 1):
  * store estimate     > MEM_FLUSH_THRESHOLD_MB
  * process RSS        > MEM_PROC_RSS_LIMIT_MB      (0 = off)
  * system RAM used   >= MEM_RAM_PCT_LIMIT %        } only when the store is
  * system free RAM    < MEM_FREE_MIN_MB            } >= MEM_MIN_STORE_MB_FOR_CLEAR
with at least MEM_FLUSH_MIN_INTERVAL_SEC between automatic flushes, so a
machine whose RAM is used by OTHER processes no longer makes the monitor
flush and clear on every single poll.

Store size for a table with n rows and k parameters (estimate, CPython 64-bit):

    bytes ~= n * (32 * k + 144)

    (32 B per value: 24 B float object + 8 B list slot; 144 B per row for
     the timestamp string, sequence number and epoch.)

USAGE
-----
    python ab_power_meter_monitor.py                  # GUI or headless per ENABLE_GUI
    python ab_power_meter_monitor.py --headless --count 0 --interval 30
    python ab_power_meter_monitor.py --gui

    import ab_power_meter_monitor as abm              # as a library
    abm.begin_run(); status = abm.poll_cycle(); abm.end_run()

Stop a headless run: Ctrl-C, SIGTERM, or `python ab_stop.py`
(writes the STOP_COLLECTION file the loop checks every 0.5 s).

TABLE MAP (page -> name)
------------------------
     0 Device_Configuration_Table          6 Diagnostic_Table
     1 Communications_Configuration_Table 7 Voltage_Current_Snapshot_Log_Table
     2 Voltage_Current_Table              8 Power_Snapshot_Log_Table
     3 Real_Time_Power_Table              9 MinMax_Log_Table
     4 Cumulative_Power_Table            10 Diagnostic_Table_Extended
     5 Demand_Data_Table

REFERENCES
----------
[R1] Qt Group, "QTabWidget Class," Qt 6 documentation (clear()/removeTab()
     do not delete the page widgets). https://doc.qt.io/qt-6/qtabwidget.html
[R2] Qt Group, "QPlainTextEdit Class," (maximumBlockCount for log viewers).
     https://doc.qt.io/qt-6/qplaintextedit.html
[R3] Python Software Foundation, "datetime," Python 3 docs (utcnow() is
     deprecated since 3.12). https://docs.python.org/3/library/datetime.html
[R4] Linux man-pages, "malloc_trim(3)."
     https://man7.org/linux/man-pages/man3/malloc_trim.3.html
[R5] h5py developers, "Datasets" (resizable, chunked datasets).
     https://docs.h5py.org/en/stable/high/dataset.html
[R6] Microsoft, "Excel specifications and limits" (1,048,576 rows).
     https://support.microsoft.com/en-us/excel/excel-specifications-and-limits
[R7] openpyxl, "Optimised Modes" (read-only / write-only workbooks).
     https://openpyxl.readthedocs.io/en/stable/optimized.html
[R8] Microsoft, "HeapSetInformation function (heapapi.h)" (HeapOptimizeResources:
     LFH caches optimised, memory decommitted if possible; Windows 8.1+).
     https://learn.microsoft.com/en-us/windows/win32/api/heapapi/nf-heapapi-heapsetinformation
[R9] Microsoft, "HeapCompact function (heapapi.h)" (returns the largest free
     block; "does not compact the heap any further"). The Universal CRT's
     _heapmin() is only HeapCompact(__acrt_heap, 0), so it is NOT used.
     https://learn.microsoft.com/en-us/windows/win32/api/heapapi/nf-heapapi-heapcompact
[R10] Microsoft, "EmptyWorkingSet function (psapi.h)" (removes as many pages as
     possible from the working set; K32EmptyWorkingSet in kernel32).
     https://learn.microsoft.com/en-us/windows/win32/api/psapi/nf-psapi-emptyworkingset

REVISION HISTORY
----------------
1.6.2  2026-10-26  Windows robustness.
                   * CSV / Markdown file open in another program (Excel opens a
                     CSV with "deny write"): 1.6.x logged an ERROR for every
                     table on every poll, and a RAM flush could not drop those
                     rows while the file stayed open. Now: one WARNING per file
                     per 10 min, rows kept in memory and written when the file
                     is free; a flush / Save Now / stop writes them to
                     "<file>_<stamp>_rescue.csv|.md", merged back (in order,
                     never twice) on the next successful write - as FITS, XLSX
                     and .vszh5 already did.
                   * SIGBREAK (Ctrl+Break, CTRL_BREAK_EVENT) now stops the
                     headless run gracefully with the final writes (Windows
                     used to kill the process).
1.6.1  2026-10-26  RAM flush now returns the freed memory to WINDOWS too. 1.6.0
                   only called glibc malloc_trim, so on Windows the rows were
                   dropped but the process kept the memory (Task Manager showed
                   no drop). After a flush, return_memory_to_os() now calls
                   HeapSetInformation(NULL, HeapOptimizeResources) [R8]
                   (decommit free heap memory) and EmptyWorkingSet() [R10]
                   (working-set trim). _heapmin() is deliberately not used:
                   in the Universal CRT it is HeapCompact(), which does not
                   release anything on a normal system [R9]. Every call is
                   best effort, bound once through private ctypes instances
                   with explicit 64-bit argument types, serialised by a lock,
                   and can never raise into the poll loop.
                   * New switches MEM_RETURN_TO_OS, MEM_TRIM_WORKING_SET (both
                     on; GUI check boxes in "Memory protection").
                   * New API: return_memory_to_os(), describe_memory_release(),
                     process_commit_mb(); release_memory(cfg=None) honours
                     the switches. Flush results carry "os_release" and
                     "os_release_text"; run_status() adds "commit_mb".
                   * Flush log / GUI line show the release steps and the RSS
                     (working set) and commit-charge change.
                   FIX (found running the test-suite on Windows Python 3.12)
                   * Duplicate-row race in persist_outputs(): writer cursors
                     were advanced in a future done-callback, which
                     concurrent.futures runs AFTER waking the threads blocked
                     in result(). A flush / Save Now / stop that waited for a
                     busy writer could read the old cursor and send the same
                     rows again (seen as extra XLSX rows). Cursors now advance
                     inside the writer task (_run_writer), each writer's cursors
                     are re-read after its previous write is confirmed done, and
                     a writer still running after WRITER_TIMEOUT_SEC is never
                     started a second time on the same file.
1.6.0  2026-10-26  Memory protection, persistence and Veusz output reworked.
                   Public functions, switches and the TIME_SERIES_STORE layout
                   are unchanged (additions only), so callers keep working.
                   FIXES
                   * GUI memory leak: QTabWidget.clear() never deleted the
                     preview tabs [R1]; every poll leaked ~70 widgets and
                     their pixmaps (measured: +1,440 widgets and +255 MB in
                     ~45 s at a 2 s period). Tabs are now created once per
                     plot and updated in place; stale tabs are deleteLater()'d.
                   * GUI log console grew without bound (QTextEdit + DEBUG
                     records). Now a QPlainTextEdit with maximumBlockCount
                     [R2] fed through a thread-safe signal bridge, INFO level.
                   * GUI thread was saturated rendering ~40 preview PNGs per
                     poll, and Poll Now blocked the GUI for the whole poll.
                     Only the visible tab is rendered (min/max decimated);
                     polling, persistence and flushing run in a worker thread.
                   * FITS: EXTNAME was tname[:8], so Voltage_Current_Table /
                     Voltage_Current_Snapshot_Log_Table and Diagnostic_Table /
                     Diagnostic_Table_Extended collided and the reader merged
                     them - rows DOUBLED on every rewrite (measured: 54-59 rows
                     after 6 polls). EXTNAME is now the full table name; old
                     files are read by their TBLNAME keyword. Same-second
                     timestamp de-duplication (which dropped real rows) is gone.
                   * CSV / Markdown: in append mode across runs, rows stopped
                     being written after the first poll (row counting mixed
                     old-run and current-run rows). Writers now use per-row
                     sequence cursors; files are never re-read to count rows.
                   * Store-size flush wrote files but never cleared the store;
                     the GUI never checked it at all. One flush path now
                     serves GUI, headless and caller, and always clears.
                   * RAM-% flush fired on EVERY poll once other processes
                     held the RAM (store cleared each poll). Added minimum
                     interval and a minimum-store-size condition.
                   * gc.collect() does not return freed heap to the OS;
                     glibc malloc_trim(0) is now called after a flush [R4].
                   * ALL_DEVICE_DATA was cleared/rebuilt by the poll thread
                     while the GUI thread iterated it (dictionary changed size
                     during iteration). Now swapped under a lock.
                   * Duplicate IPs in IP_LIST polled one meter from two threads
                     at once and duplicated its rows; the list is de-duplicated.
                   * Duplicate parameter names inside one table overwrote each
                     other; later ones are stored as "Name [#N]".
                   * HTTP_TIMEOUT_SEC / retry defaults were bound at import, so
                     caller overrides were ignored. Read at call time now.
                   * "inf"/"nan" cell text raised OverflowError and lost the
                     whole table; non-finite values are kept as text.
                   * The named first-device dicts flipped between devices
                     (thread completion order); now follow IP_LIST order.
                   * Silent headless mode added a NullHandler to the root
                     logger on every run (unbounded list in caller mode).
                   * datetime.utcnow() (deprecated in Python 3.12 [R3])
                     replaced by timezone-aware UTC.
                   * FITS/XLSX were rewritten in place; a crash mid-write
                     destroyed the file. Now written to a temp file and
                     moved into place atomically.
                   CHANGES
                   * Veusz output written natively with vszh5_writer.py
                     (h5py) - the veusz package is no longer required. One
                     document per device is appended in place every poll
                     (VEUSZ_APPEND_EVERY_CYCLE); rollover at VSZH5_MAX_FILE_MB.
                     "Open in Veusz" launches the Veusz program on the file.
                   * FITS/XLSX rewritten every FULL_WRITE_EVERY_N_CYCLES polls
                     and at flush/stop; XLSX streamed with openpyxl read-only
                     / write-only workbooks [R7]; rollover before the Excel
                     row limit [R6] (XLSX_MAX_ROWS_PER_FILE) and for FITS
                     (FITS_MAX_ROWS_PER_FILE). FITS gains an EPOCH_UTC column.
                   * Polls can be cancelled between pages (stop is prompt);
                     each meter is still read strictly one page at a time.
                   * New public API: begin_run(), poll_cycle(), end_run(),
                     persist_outputs(), finalize_outputs(), perform_ram_flush(),
                     ram_flush_reason(), trim_store(), store_snapshot(),
                     ts_store_size_estimate_mb(), process_rss_mb(),
                     set_console_logging(), dedupe_ips(), open_in_veusz(),
                     load_spill().
                   * New switches: METER_URL_TEMPLATE, VEUSZ_APPEND_EVERY_CYCLE,
                     VSZH5_MAX_FILE_MB, VSZH5_TIME_AXIS, VSZH5_PANELS_PER_PAGE,
                     FULL_WRITE_EVERY_N_CYCLES, FITS_MAX_ROWS_PER_FILE,
                     XLSX_MAX_ROWS_PER_FILE, FSYNC_OUTPUTS, MEM_PROC_RSS_LIMIT_MB,
                     MEM_FLUSH_MIN_INTERVAL_SEC, MEM_MIN_STORE_MB_FOR_CLEAR,
                     RAM_FLUSH_KEEP_SAMPLES, RAM_FLUSH_SPILL_PICKLE,
                     PREVIEW_MAX_POINTS, GUI_LOG_MAX_LINES,
                     GUI_KEEP_SAMPLES_AFTER_FLUSH; env ABMETER_OUTPUT_DIR.
                   * VEUSZ_WRITE_ON_FLUSH is kept for compatibility but has no
                     effect (the document is appended continuously).
                   * GUI: Save Now, Flush RAM Now, memory status line, theme
                     follows into the preview plots, final writes finish
                     before the window closes (no QThread.terminate()).
1.5.5  2026-10-11  Restore sys.stdout/stderr even if the headless loop raises a
                   non-KeyboardInterrupt exception in silent mode.
1.5.4  2026-10-11  Import-safe CLI front-end (--headless, --count, --interval,
                   --ips, --ram-pct, --no-ram-flush).
1.5.3  2026-08-13  Docs: requirements files, Python 3.7.11 note, advisory check.
1.5.2  2026-08-13  ENABLE_RAM_FLUSH master switch (headless + GUI).
1.5.1  2026-07-22  ENABLE_SIGNAL_HANDLERS; embedding hosts are not torn down.
1.5.0  2026-06-09  Production release (see git history for 1.0 - 1.4).
===============================================================================
"""

# %% ===========================================================================
#    IMPORTS
# =============================================================================
from __future__ import annotations   # PEP 563: annotations are never evaluated

# --- standard library ---------------------------------------------------------
import atexit
import bisect
import concurrent.futures
import csv
import datetime
import gc
import glob
import io
import itertools
import json
import logging
import math
import os
import pickle
import queue
import re
import signal
import subprocess
import sys
import tempfile
import threading
import time
import traceback
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

# --- third party, always required ----------------------------------------------
import requests                         # HTTP client (one keep-alive Session per meter)
from bs4 import BeautifulSoup           # HTML parsing; libxml2 underneath via "lxml"

# --- third party, optional (each feature degrades gracefully without it) ---------
try:
    import psutil                       # RAM / RSS figures for the flush triggers
except ImportError:                     # pragma: no cover - reported in the GUI/log
    psutil = None

try:
    import numpy as np                  # FITS, .vszh5 and preview decimation
except ImportError:                     # pragma: no cover
    np = None

# matplotlib renders the GUI previews off-screen with the Agg backend. It must
# be selected once, before any Qt or pyplot code runs; the interactive dialog
# embeds its own FigureCanvasQTAgg, which does not depend on this choice.
try:
    import matplotlib
    matplotlib.use("Agg")
except ImportError:                     # pragma: no cover - GUI previews skipped
    matplotlib = None

__version__ = "1.6.2"                   # Semantic Versioning 2.0.0
__author__ = "William W. Wallace"
__email__ = "wwallace@nrao.edu"
__license__ = "GPL-3.0-or-later"

# %% ===========================================================================
#    PYTHON VERSION CHECK (advisory - never aborts)
# =============================================================================
# The source stays at the 3.7 language level (no walrus, no match, no builtin
# generic annotations evaluated at run time). 3.12 is the tested primary target.
# On 3.7.x the dependencies must be pinned (requirements-py37.txt).
MIN_PYTHON = (3, 7, 11)
RECOMMENDED_PYTHON = (3, 12)
if sys.version_info < MIN_PYTHON:
    sys.stderr.write(
        "WARNING: ab_power_meter_monitor requires Python >= {0}; running {1}. "
        "Behaviour below the minimum is unsupported.\n".format(
            ".".join(map(str, MIN_PYTHON)), ".".join(map(str, sys.version_info[:3]))))

# %% ===========================================================================
#    SWITCHES
#  ██████╗  ██████╗ ██╗    ██╗███████╗██████╗     ███████╗██╗    ██╗██╗████████╗ ██████╗██╗  ██╗███████╗███████╗
#  ██╔══██╗██╔═══██╗██║    ██║██╔════╝██╔══██╗    ██╔════╝██║    ██║██║╚══██╔══╝██╔════╝██║  ██║██╔════╝██╔════╝
#  ██████╔╝██║   ██║██║ █╗ ██║█████╗  ██████╔╝    ███████╗██║ █╗ ██║██║   ██║   ██║     ███████║█████╗  ███████╗
#  ██╔═══╝ ██║   ██║██║███╗██║██╔══╝  ██╔══██╗    ╚════██║██║███╗██║██║   ██║   ██║     ██╔══██║██╔══╝  ╚════██║
#  ██║     ╚██████╔╝╚███╔███╔╝███████╗██║  ██║    ███████║╚███╔███╔╝██║   ██║   ╚██████╗██║  ██║███████╗███████║
#  ╚═╝      ╚═════╝  ╚══╝╚══╝ ╚══════╝╚═╝  ╚═╝    ╚══════╝ ╚══╝╚══╝ ╚═╝   ╚═╝    ╚═════╝╚═╝  ╚═╝╚══════╝╚══════╝
#
#  ALL RUNTIME BEHAVIOUR IS CONTROLLED BY THE VARIABLES IN THIS SECTION.
#  0 = off / False, 1 = on / True. A caller may override any of them by
#  assignment after import (abm.ENABLE_CSV = 1); they are read at call time.
# =============================================================================

# %% Feature / output enable switches  (0 = off, 1 = on) -----------------------
ENABLE_GUI = 1          # Show the PySide6 main window (0 = headless)
ENABLE_FITS = 0         # NRAO-style FITS file per device
ENABLE_CSV = 0          # One CSV per device per table (appended every poll)
ENABLE_XLSX = 0         # One Excel workbook per device
ENABLE_LOG_APPEND = 0   # One Markdown (.md) table log per device per table
ENABLE_LOG_FILE = 0     # ab_monitor.log (Python logging file handler)
ENABLE_VEUSZ = 0        # One Veusz .vszh5 document per device (native, h5py)

# Veusz document update cadence:
#   1 = append the new rows to the .vszh5 on EVERY poll (live document; default)
#   0 = append only at RAM flush, Save Now and stop (fewer writes)
VEUSZ_APPEND_EVERY_CYCLE = 1
# DEPRECATED (1.6.0): kept so existing callers that assign it do not fail. It
# has no effect - the document is now appended continuously (no snapshots).
VEUSZ_WRITE_ON_FLUSH = 0

# Install SIGINT/SIGTERM handlers that turn Ctrl-C / kill into a graceful stop.
#   1 = standalone / CLI use (default).
#   0 = REQUIRED when embedded in a host that owns its own signal handling
#       (exporter, web server, supervised service) - otherwise the host's
#       SIGTERM is swallowed into _HEADLESS_STOP. Stop with ab_stop.py then.
ENABLE_SIGNAL_HANDLERS = 1

# %% Headless loop control -------------------------------------------------------
HEADLESS_LOOP_COUNT = 1          # 0 = run until stopped; N = run N polls then exit
HEADLESS_MAX_CONSEC_FAILS = 3    # consecutive all-device failures before exit (0 = never)

# %% Headless console output (any combination is valid) ---------------------------
HEADLESS_CONSOLE_DICTS_ONLY = 0  # 1 = print only the dict JSON blocks to stdout
HEADLESS_PRINT_EACH_SAMPLE = 0   # 1 = print the 11 named dicts after every poll
HEADLESS_PRINT_CUMULATIVE = 0    # 1 = print the whole store at stop / flush
HEADLESS_SILENT = 1              # 1 = no console output at all (files unaffected)

# %% Memory protection (RAM flush) -------------------------------------------------
# Master enable for AUTOMATIC flushing. 0 = never flush/clear mid-run (the store
# then grows for the whole run - watch RAM). Manual "Flush RAM Now", Save Now
# and the end-of-run writes are not affected.
ENABLE_RAM_FLUSH = 1
MEM_RAM_PCT_LIMIT = 60           # flush when SYSTEM RAM use >= N % (see below)
MEM_FLUSH_THRESHOLD_MB = 256     # flush when the store estimate > N MB
MEM_FREE_MIN_MB = 512            # flush when system free RAM < N MB
MEM_PROC_RSS_LIMIT_MB = 0        # flush when THIS process's RSS > N MB (0 = off)
MEM_FLUSH_MIN_INTERVAL_SEC = 120  # at least this long between automatic flushes
# The two SYSTEM-wide triggers (RAM %, free MB) only clear the store when it
# holds at least this much - clearing a tiny store cannot help when OTHER
# processes use the RAM, and used to clear the store on every poll.
MEM_MIN_STORE_MB_FOR_CLEAR = 1.0
RAM_FLUSH_KEEP_SAMPLES = 0       # newest rows per table kept after a flush (0 = clear)
RAM_FLUSH_SPILL_PICKLE = 0       # 1 = pickle the store before clearing when NO
#                                  file output is enabled (otherwise it is dropped)
# Hand the freed memory back to the OPERATING SYSTEM after every RAM flush
# (1.6.1). Without this the rows are gone from the store, but the process keeps
# the freed heap, so Task Manager / top show no drop.
#   Linux   : glibc malloc_trim(0) [R4]
#   Windows : decommit free low-fragmentation-heap memory of every heap
#             (HeapSetInformation(NULL, HeapOptimizeResources), Windows 8.1+)
#             [R8] and, with MEM_TRIM_WORKING_SET = 1, trim the working set
#             (EmptyWorkingSet) [R10].
#   macOS / other: gc.collect() only.
MEM_RETURN_TO_OS = 1
# Windows only. 1 = also remove the process's pages from its working set after
# a flush: the physical RAM goes back to Windows at once and Task Manager's
# "Memory" column drops. Pages still in use fault back in on their next access
# (a few ms, once). 0 = leave the working set to the Windows memory manager.
MEM_TRIM_WORKING_SET = 1

# %% Persistence cadence and file sizes ---------------------------------------------
FULL_WRITE_EVERY_N_CYCLES = 10   # FITS / XLSX rewrite every N polls (0 = flush/stop only)
FSYNC_OUTPUTS = 1                # 1 = os.fsync() appended CSV/MD files (crash-safe)
FITS_MAX_ROWS_PER_FILE = 100000  # start a new FITS file beyond N rows (0 = never)
XLSX_MAX_ROWS_PER_FILE = 50000   # start a new workbook beyond N rows (Excel max 1,048,574 data rows)
VSZH5_MAX_FILE_MB = 256          # start a new .vszh5 beyond N MB (0 = never)
VSZH5_TIME_AXIS = "local"        # "local" (DST-aware wall clock) or "utc"
VSZH5_PANELS_PER_PAGE = 6        # graphs per per-table page in the .vszh5 (2 columns)
VSZH5_FILE_LOCKING = None        # h5py locking=: None (HDF5 default), True, False, "best-effort"

# %% GUI ----------------------------------------------------------------------------
PREVIEW_MAX_POINTS = 2000        # preview points per series (min/max decimation)
GUI_LOG_MAX_LINES = 2000         # status console keeps only the newest N lines
GUI_KEEP_SAMPLES_AFTER_FLUSH = 120  # GUI default for RAM_FLUSH_KEEP_SAMPLES (preview continuity)

# %% Devices --------------------------------------------------------------------------
# Explicit list of meters. Duplicates are removed automatically (a duplicate
# used to poll one meter from two threads at once and store its rows twice).
IP_LIST: List[str] = [
    "10.16.130.50",
    "10.16.130.54",
]
# URL of one table page. {ip} and {page} are substituted. Change only for a
# proxy / port forward / test emulator, e.g. "http://{ip}:8080/{page}".
METER_URL_TEMPLATE = "http://{ip}/{page}"

# %% Polling / timing --------------------------------------------------------------------
SAMPLE_PERIOD_SEC = 30           # seconds between the starts of successive polls
HTTP_TIMEOUT_SEC = 8             # per-request timeout (s) - not a port number
HTTP_RETRY_COUNT = 2             # attempts per page (1 = no retry)
HTTP_RETRY_DELAY_SEC = 2.0       # seconds between attempts
APPEND_OUTPUT_FILES = 1          # 1 = append to existing files; 0 = start fresh each run

# %% Output paths --------------------------------------------------------------------------
# All output goes under OUTPUT_BASE_DIR. The environment variable
# ABMETER_OUTPUT_DIR overrides the default (handy for systemd units), e.g.
#   Environment=ABMETER_OUTPUT_DIR=/srv/abmeter
_DEFAULT_BASE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ab_meter_output")
OUTPUT_BASE_DIR = os.environ.get("ABMETER_OUTPUT_DIR", "").strip() or _DEFAULT_BASE_DIR
OUTPUT_DIR = OUTPUT_BASE_DIR                          # back-compat alias
LOG_DIR = os.path.join(OUTPUT_BASE_DIR, "logs")       # log file and Markdown tables
FITS_DIR = os.path.join(OUTPUT_BASE_DIR, "fits")
CSV_DIR = os.path.join(OUTPUT_BASE_DIR, "csv")
XLSX_DIR = os.path.join(OUTPUT_BASE_DIR, "xlsx")
VEUSZ_DIR = os.path.join(OUTPUT_BASE_DIR, "veusz")
SPILL_DIR = os.path.join(OUTPUT_BASE_DIR, "spill")    # RAM_FLUSH_SPILL_PICKLE files
# Touch this file (or run ab_stop.py) to stop a headless run cleanly.
STOP_SIGNAL_FILE = os.path.join(OUTPUT_BASE_DIR, "STOP_COLLECTION")
# Remember the import-time base so default_config() can tell whether a caller
# changed OUTPUT_BASE_DIR without changing the sub-directories.
_IMPORT_BASE_DIR = OUTPUT_BASE_DIR

# %% Table page index -> canonical name --------------------------------------------------------
TABLE_NAMES: Dict[int, str] = {
    0: "Device_Configuration_Table",
    1: "Communications_Configuration_Table",
    2: "Voltage_Current_Table",
    3: "Real_Time_Power_Table",
    4: "Cumulative_Power_Table",
    5: "Demand_Data_Table",
    6: "Diagnostic_Table",
    7: "Voltage_Current_Snapshot_Log_Table",
    8: "Power_Snapshot_Log_Table",
    9: "MinMax_Log_Table",
    10: "Diagnostic_Table_Extended",
}

# %% Unit inference: lower-case substring of a parameter name -> unit label ----------------------
UNIT_MAP: List[Tuple[str, str]] = [
    ("current", "A"),
    ("voltage", "V"),
    ("frequency", "Hz"),
    ("kw hour", "kWh"),
    ("kvar hour", "kVARh"),
    ("real power", "W"),
    ("reactive power", "VAR"),
    ("apparent power", "VA"),
    ("true pf", "%"),
    ("displacement pf", "%"),
    ("distortion pf", "%"),
    ("demand current", "A"),
    ("demand power", "W"),
    ("demand apparent", "VA"),
    ("demand reactive", "VAR"),
    ("elapsed time", "s"),
    ("period", "min"),
    ("interval", "s"),
    ("pulse width", "ms"),
]

# %% Overlay groups: parameters sharing one unit, plotted together --------------------------------
# Keys are human-readable titles (GUI tabs, Veusz overlay pages and legends).
VEUSZ_OVERLAY_GROUPS: Dict[str, List[str]] = {
    "Current A": ["current"],
    "Voltage L-L (V)": ["l1-l2 voltage", "l2-l3 voltage", "l3-l1 voltage",
                        "3 phase average voltage l-l", "pos. seq. voltage",
                        "neg. seq. voltage", "aux voltage"],
    "Voltage L-N (V)": ["l1-n voltage", "l2-n voltage", "l3-n voltage",
                        "3 phase average voltage l-n"],
    "Real Power (W)": ["l1 real power", "l2 real power", "l3 real power", "total real power"],
    "Reactive Power (VAR)": ["l1 reactive power", "l2 reactive power",
                             "l3 reactive power", "total reactive power"],
    "Apparent Power (VA)": ["l1 apparent power", "l2 apparent power",
                            "l3 apparent power", "total apparent power"],
    "True PF (%)": ["l1 true pf", "l2 true pf", "l3 true pf", "total true pf"],
    "Displacement PF (%)": ["l1 displacement pf", "l2 displacement pf",
                            "l3 displacement pf", "total displacement pf"],
    "Distortion PF (%)": ["l1 distortion pf", "l2 distortion pf",
                          "l3 distortion pf", "total distortion pf"],
}


# %% ===========================================================================
#    LOGGING
# =============================================================================
# Handlers are identified by name, so re-running the setup (main() is called
# once per poll by older callers) never stacks duplicate handlers.
_CONSOLE_HANDLER_NAME = "abmonitor_console"
_FILE_HANDLER_NAME = "abmonitor_file"
_LOG_FORMAT = logging.Formatter(fmt="%(asctime)s | %(levelname)-8s | %(name)s | %(message)s",
                                datefmt="%Y-%m-%dT%H:%M:%S")


def _named_handler(log: logging.Logger, name: str) -> Optional[logging.Handler]:
    """The handler called ``name`` attached to ``log``, or None."""
    return next((h for h in log.handlers if getattr(h, "name", None) == name), None)


def _setup_logging(log_dir: str, append: bool = True,
                   enable_log_file: bool = True) -> logging.Logger:
    """Initialise (or re-initialise) the "ABMonitor" logger. Idempotent.

    Parameters
    ----------
    log_dir : str
        Folder of ab_monitor.log; not created when ``enable_log_file`` is False.
    append : bool
        Append to an existing log file (True) or truncate it (False).
    enable_log_file : bool
        Attach the file handler (DEBUG and above). The console handler (INFO
        and above, stdout) is always ensured unless set_console_logging(False)
        removed it on purpose.
    """
    log = logging.getLogger("ABMonitor")
    log.setLevel(logging.DEBUG)

    # Console handler - added once, unless a host disabled console logging.
    if _named_handler(log, _CONSOLE_HANDLER_NAME) is None and _CONSOLE_STATE["enabled"]:
        console = logging.StreamHandler(sys.stdout)
        console.set_name(_CONSOLE_HANDLER_NAME)
        console.setLevel(_CONSOLE_STATE["level"])
        console.setFormatter(_LOG_FORMAT)
        log.addHandler(console)

    # File handler - replaced when the folder changes, removed when disabled.
    existing = _named_handler(log, _FILE_HANDLER_NAME)
    if enable_log_file:
        os.makedirs(log_dir, exist_ok=True)
        log_path = os.path.abspath(os.path.join(log_dir, "ab_monitor.log"))
        if existing is not None and getattr(existing, "baseFilename", "") != log_path:
            log.removeHandler(existing)
            existing.close()
            existing = None
        if existing is None:
            file_handler = logging.FileHandler(log_path, mode="a" if append else "w",
                                               encoding="utf-8")
            file_handler.set_name(_FILE_HANDLER_NAME)
            file_handler.setLevel(logging.DEBUG)
            file_handler.setFormatter(_LOG_FORMAT)
            log.addHandler(file_handler)
            log.info("Log file: %s", log_path)
    elif existing is not None:
        log.removeHandler(existing)
        existing.close()
    return log


# Console logging state, so set_console_logging() survives a later
# _setup_logging() call (older callers re-run main() on every poll).
_CONSOLE_STATE: Dict[str, Any] = {"enabled": True, "level": logging.INFO}


def set_console_logging(enabled: bool = True, level: int = logging.INFO) -> None:
    """Enable/disable the monitor's own stdout log handler (persistent setting).

    A host with its own logging (the Prometheus exporter, a service) calls
    set_console_logging(False) so ABMonitor records reach the console once,
    through the host's root handler, instead of twice.
    """
    _CONSOLE_STATE["enabled"] = bool(enabled)
    _CONSOLE_STATE["level"] = int(level)
    log = logging.getLogger("ABMonitor")
    handler = _named_handler(log, _CONSOLE_HANDLER_NAME)
    if handler is not None and not enabled:
        log.removeHandler(handler)
    elif enabled:
        if handler is None:
            _setup_logging(LOG_DIR, enable_log_file=_named_handler(log, _FILE_HANDLER_NAME) is not None)
        else:
            handler.setLevel(int(level))


# Module logger, console-only at import time: importing the module never
# creates folders or files. main()/begin_run() attach the file handler.
logger: logging.Logger = _setup_logging(LOG_DIR, append=True, enable_log_file=False)


# %% ===========================================================================
#    SMALL HELPERS: TIME, TEXT, NAMES
# =============================================================================
def _utc_now() -> datetime.datetime:
    """Timezone-aware current UTC time (replaces deprecated utcnow() [R3])."""
    return datetime.datetime.now(datetime.timezone.utc)


def _utc_iso_z(moment: Optional[datetime.datetime] = None) -> str:
    """ISO-8601 UTC string with a Z suffix, seconds resolution."""
    return (moment or _utc_now()).strftime("%Y-%m-%dT%H:%M:%SZ")


def _local_text(epoch: float) -> str:
    """Local wall-clock 'YYYY-MM-DD HH:MM:SS' of a Unix time (the store's format)."""
    return datetime.datetime.fromtimestamp(epoch).strftime("%Y-%m-%d %H:%M:%S")


def _epoch_from_local_text(text: str) -> float:
    """Inverse of _local_text (ambiguous in the DST fall-back hour); NaN on error."""
    try:
        return datetime.datetime.strptime(text, "%Y-%m-%d %H:%M:%S").timestamp()
    except (TypeError, ValueError, OverflowError, OSError):
        return float("nan")


def _stamp() -> str:
    """Compact local time stamp for file names: YYYYmmdd_HHMMSS."""
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def _fits_ascii(value: str, *, comment: bool = False) -> str:
    """Printable 7-bit ASCII for a FITS header value / COMMENT / HISTORY card.

    FITS restricts header characters to 0x20-0x7E. Common typographic
    characters are transliterated first; anything else becomes '-'. String
    values are cut to 68 characters, COMMENT/HISTORY text to 72.
    """
    value = str(value)
    value = value.replace("—", " - ").replace("–", " - ")   # em/en dash
    value = value.replace("°", "deg").replace("μ", "u").replace("Ω", "Ohm")
    sanitised = "".join(c if 0x20 <= ord(c) <= 0x7E else "-" for c in value)
    return sanitised[:72] if comment else sanitised[:68]


def infer_unit(param_name: str) -> str:
    """Physical unit guessed from a parameter name via UNIT_MAP ('' if unknown)."""
    lower = param_name.lower()
    for key, unit in UNIT_MAP:
        if key in lower:
            return unit
    return ""


def dedupe_ips(ip_list: Iterable[str]) -> List[str]:
    """Stripped, non-empty IPs in first-seen order, without duplicates."""
    seen: Set[str] = set()
    out: List[str] = []
    for raw in ip_list or []:
        ip = str(raw).strip()
        if ip and ip not in seen:
            seen.add(ip)
            out.append(ip)
    return out


def _ident(text: str) -> str:
    """ASCII identifier for Veusz dataset/page names: [A-Za-z0-9_], never empty."""
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", str(text)).strip("_")
    return cleaned or "x"


def _veusz_safe(name: str) -> str:
    """Kept for compatibility (1.5.x): underscores, word characters, <= 64 chars.

    New code uses _ident() plus collision-free naming instead, because the
    64-character cut made long parameter names collide.
    """
    s = name.replace(" ", "_").replace(".", "_").replace("/", "_")
    return re.sub(r"[^\w]", "", s)[:64]


def _safe_ip(ip: str) -> str:
    """IP address in a file name: dots and colons replaced by underscores."""
    return ip.replace(".", "_").replace(":", "_")


def _octet(ip: str) -> str:
    """Last octet (or last group) of an address, used to prefix dataset names."""
    return re.split(r"[.:]", ip)[-1] or _safe_ip(ip)


def _table_order(tname: str) -> Tuple[int, str]:
    """Sort key placing tables in meter page order, unknown names last."""
    for page, name in TABLE_NAMES.items():
        if name == tname:
            return (page, tname)
    return (len(TABLE_NAMES), tname)


# %% ===========================================================================
#    HTML FETCH & PARSE
# =============================================================================
def fetch_table_html(
    ip: str,
    page: int,
    timeout: Optional[float] = None,
    retries: Optional[int] = None,
    retry_delay: Optional[float] = None,
    session: Optional["requests.Session"] = None,
    stop_event: Optional[threading.Event] = None,
) -> Optional[str]:
    """Fetch the raw HTML of one meter page (METER_URL_TEMPLATE), with retries.

    Parameters
    ----------
    ip, page : str, int
        Meter address and page index (0 ... 10, the TABLE_NAMES keys).
    timeout, retries, retry_delay : optional
        Per-request timeout (s), total attempts (1 = no retry) and the pause
        between attempts. None = the module switch at CALL time (1.5.x bound
        them at import, so caller overrides were silently ignored).
    session : requests.Session, optional
        Keep-alive session owned by the per-meter thread; pages of one meter
        reuse its TCP connection.
    stop_event : threading.Event, optional
        When set, no further attempt is started and the retry pause ends at
        once, so a stop request is honoured within one HTTP timeout.

    Returns
    -------
    str or None
        Page HTML, or None when every attempt failed or a stop was requested.
    """
    timeout = HTTP_TIMEOUT_SEC if timeout is None else timeout
    retries = HTTP_RETRY_COUNT if retries is None else retries
    retry_delay = HTTP_RETRY_DELAY_SEC if retry_delay is None else retry_delay
    url = METER_URL_TEMPLATE.format(ip=ip, page=page)
    getter = session.get if session is not None else requests.get
    attempts = max(1, int(retries))
    for attempt in range(1, attempts + 1):
        if stop_event is not None and stop_event.is_set():
            return None
        try:
            response = getter(url, timeout=timeout)
            try:
                response.raise_for_status()
                text = response.text
            finally:
                response.close()            # release the connection to the pool at once
            if attempt > 1:
                logger.info("Fetched %s on attempt %d/%d.", url, attempt, attempts)
            else:
                logger.debug("Fetched %s - %d characters", url, len(text))
            return text
        except Exception as exc:             # requests.RequestException, OSError, ...
            if attempt < attempts:
                logger.warning("Fetch failed %s (%d/%d): %s - retry in %.1f s",
                               url, attempt, attempts, exc, retry_delay)
                if stop_event is not None:
                    if stop_event.wait(retry_delay):   # wakes immediately on stop
                        return None
                else:
                    time.sleep(retry_delay)
            else:
                logger.warning("Fetch failed %s - all %d attempt(s) exhausted: %s",
                               url, attempts, exc)
    return None


def _parse_cell_value(raw: str) -> Any:
    """Numeric value of a table cell when it is a finite number, else the text.

    Integers written without a decimal point stay int (as in 1.5.x). "inf",
    "nan" and anything float() rejects are kept as text, so one odd cell can
    no longer raise OverflowError and lose the whole table.
    """
    try:
        value = float(raw)
    except (TypeError, ValueError, OverflowError):
        return raw
    if not math.isfinite(value):
        return raw
    if "." not in raw and value == int(value):      # 1.5.x rule, kept for parity
        return int(value)
    return value


def parse_html_table(html: Optional[str], table_name: str, ip: str, page: int) -> Dict[str, Any]:
    """Parse one meter page into the table dict used everywhere in the program.

    Rows have the form ``<tr><td>#</td><td>Parameter Name</td><td>Value</td></tr>``.

    Returns
    -------
    dict
        ``_meta``     : table_name, source_ip, page_index, fetch_utc (ISO Z),
                        fetch_epoch (Unix seconds, 1.6.0), error (if any)
        ``#N``        : parameter name of row N
        ``#N_value``  : numeric value, or the cell text
        ``#N_unit``   : unit inferred from the name ('' if unknown)
    """
    now = _utc_now()
    result: Dict[str, Any] = {
        "_meta": {
            "table_name": table_name,
            "source_ip": ip,
            "page_index": page,
            "fetch_utc": _utc_iso_z(now),
            "fetch_epoch": now.timestamp(),
        }
    }
    if html is None:
        result["_meta"]["error"] = "No HTML received"
        return result

    soup = None
    try:
        soup = BeautifulSoup(html, "lxml")
        for row in soup.find_all("tr"):
            cells = row.find_all("td")
            if len(cells) < 3:
                continue                        # header row (<th>/<font>) or decoration
            idx_text = cells[0].get_text(strip=True)
            param_name = cells[1].get_text(strip=True)
            raw_value = cells[2].get_text(strip=True)
            if param_name in ("Parameter Name", "#") or idx_text == "#":
                continue
            key = f"#{idx_text}"
            result[key] = param_name
            result[f"{key}_value"] = _parse_cell_value(raw_value)
            result[f"{key}_unit"] = infer_unit(param_name)
        logger.debug("Parsed table '%s' - %d rows", table_name, (len(result) - 1) // 3)
    except Exception as exc:
        logger.error("Parse error for table '%s': %s\n%s", table_name, exc, traceback.format_exc())
        result["_meta"]["error"] = str(exc)
    finally:
        if soup is not None:
            # BeautifulSoup trees are parent/child reference cycles; decompose()
            # breaks them now instead of waiting for a GC cycle.
            soup.decompose()
    return result


# %% ===========================================================================
#    POLLING: one thread per meter, pages strictly serial per meter
# =============================================================================
def poll_all_devices(
    ip_list: Sequence[str],
    table_names: Dict[int, str],
    retries: Optional[int] = None,
    retry_delay: Optional[float] = None,
    stop_event: Optional[threading.Event] = None,
    timeout: Optional[float] = None,
) -> Dict[str, Dict[str, Any]]:
    """Poll all meters in parallel, reading each meter's pages one at a time.

    Parallelism
    -----------
    One worker thread per meter. Inside a thread the pages are fetched
    SERIALLY on one keep-alive session: the meter firmware serves a single
    request at a time and drops or resets requests when several arrive at
    once. Meters themselves are independent, so they are read concurrently.
    The page HTML is parsed in the same thread right after it arrives, which
    overlaps parsing of one meter with network waits of the others.

    Duplicate IPs are removed first (a duplicate used to read one meter from
    two threads at once, breaking the serial rule and storing rows twice).

    Parameters
    ----------
    ip_list : sequence of str
        Meter addresses.
    table_names : dict
        Page index -> table name; pages are read in ascending page order.
    retries, retry_delay, timeout : optional
        Per-page fetch settings (None = module switches at call time).
    stop_event : threading.Event, optional
        When set, remaining pages are skipped and the poll returns promptly.

    Returns
    -------
    dict
        ``"<ip>_<table>"`` -> table dict (see parse_html_table), plus
        ``"__poll_meta__"`` (callers pop it before update_named_dicts()):
        total_ips, ok_ips, failed_ips, all_failed, cancelled, pages_ok.
    """
    clean_ips = dedupe_ips(ip_list)
    results: Dict[str, Dict[str, Any]] = {}
    if len(clean_ips) != len([ip for ip in ip_list if str(ip).strip()]):
        logger.warning("IP list contained duplicates; polling each meter once: %s", clean_ips)
    if not clean_ips:
        results["__poll_meta__"] = {"total_ips": 0, "ok_ips": [], "failed_ips": [],
                                    "all_failed": True, "cancelled": False, "pages_ok": {}}
        return results

    ordered_pages: List[Tuple[int, str]] = sorted(table_names.items())
    lock = threading.Lock()
    pages_ok: Dict[str, int] = {ip: 0 for ip in clean_ips}

    def _poll_ip(ip: str) -> None:
        """Thread target: every page of one meter, in order, one at a time."""
        with requests.Session() as session:
            for page, tname in ordered_pages:
                if stop_event is not None and stop_event.is_set():
                    break                      # cancelled: leave the rest unread
                html: Optional[str] = None
                try:
                    html = fetch_table_html(ip, page, timeout=timeout, retries=retries,
                                            retry_delay=retry_delay, session=session,
                                            stop_event=stop_event)
                    parsed = parse_html_table(html, tname, ip, page)
                except Exception as exc:       # never let one page kill the thread
                    logger.error("poll %s page %d (%s): %s", ip, page, tname, exc)
                    now = _utc_now()
                    # Keep a full _meta so the entry is filed under the right
                    # device/table (1.5.x filed it under "unknown").
                    parsed = {"_meta": {"table_name": tname, "source_ip": ip, "page_index": page,
                                        "fetch_utc": _utc_iso_z(now),
                                        "fetch_epoch": now.timestamp(), "error": str(exc)}}
                ok = html is not None and "error" not in parsed.get("_meta", {})
                html = None                    # drop the page text before the next fetch
                with lock:
                    results[f"{ip}_{tname}"] = parsed
                    if ok:
                        pages_ok[ip] += 1

    try:
        with concurrent.futures.ThreadPoolExecutor(
                max_workers=len(clean_ips), thread_name_prefix="ab_poll") as pool:
            futures = {pool.submit(_poll_ip, ip): ip for ip in clean_ips}
            for future in concurrent.futures.as_completed(futures):
                try:
                    future.result()
                except Exception as exc:
                    logger.error("poll thread for %s raised: %s", futures[future], exc)
    except Exception as exc:                   # e.g. RuntimeError at interpreter exit
        logger.error("poll_all_devices executor error: %s", exc)

    ok_ips = [ip for ip in clean_ips if pages_ok[ip] > 0]
    failed = [ip for ip in clean_ips if pages_ok[ip] == 0]
    results["__poll_meta__"] = {
        "total_ips": len(clean_ips),
        "ok_ips": ok_ips,
        "failed_ips": failed,
        "all_failed": not ok_ips,
        "cancelled": bool(stop_event is not None and stop_event.is_set()),
        "pages_ok": dict(pages_ok),
    }
    if failed:
        logger.warning("poll_all_devices: %d meter(s) returned no data: %s", len(failed), failed)
    if not ok_ips:
        logger.error("poll_all_devices: ALL %d meter(s) failed.", len(clean_ips))
    return results


# %% ===========================================================================
#    NAMED TABLE DICTS  (latest snapshot; first device in IP_LIST order)
# =============================================================================
Device_Configuration_Table: Dict[str, Any] = {}
Communications_Configuration_Table: Dict[str, Any] = {}
Voltage_Current_Table: Dict[str, Any] = {}
Real_Time_Power_Table: Dict[str, Any] = {}
Cumulative_Power_Table: Dict[str, Any] = {}
Demand_Data_Table: Dict[str, Any] = {}
Diagnostic_Table: Dict[str, Any] = {}
Voltage_Current_Snapshot_Log_Table: Dict[str, Any] = {}
Power_Snapshot_Log_Table: Dict[str, Any] = {}
MinMax_Log_Table: Dict[str, Any] = {}
Diagnostic_Table_Extended: Dict[str, Any] = {}

# Latest snapshot of every device: {ip: {table_name: table_dict}}. Replaced as a
# whole under _ADD_LOCK each poll; readers take a snapshot under the same lock.
ALL_DEVICE_DATA: Dict[str, Dict[str, Dict[str, Any]]] = {}
_ADD_LOCK = threading.RLock()


def all_device_data_snapshot() -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Shallow copy of ALL_DEVICE_DATA ({ip: {table: dict}}) taken under its lock."""
    with _ADD_LOCK:
        return {ip: dict(tables) for ip, tables in ALL_DEVICE_DATA.items()}


def update_named_dicts(all_data: Dict[str, Dict[str, Any]],
                       ip_list: Optional[Sequence[str]] = None) -> None:
    """Install a poll result as the latest snapshot and append it to the store.

    Parameters
    ----------
    all_data : dict
        poll_all_devices() output (``__poll_meta__`` may be present; ignored).
    ip_list : sequence of str, optional
        The run's IP list (default IP_LIST). Every listed meter gets an entry
        in ALL_DEVICE_DATA, empty when it returned nothing this poll, and the
        named first-device dicts follow THIS order (1.5.x used the order in
        which threads finished, so they could flip between meters).
    """
    global Device_Configuration_Table, Communications_Configuration_Table
    global Voltage_Current_Table, Real_Time_Power_Table, Cumulative_Power_Table
    global Demand_Data_Table, Diagnostic_Table, Voltage_Current_Snapshot_Log_Table
    global Power_Snapshot_Log_Table, MinMax_Log_Table, Diagnostic_Table_Extended

    # Build the new snapshot privately, then swap it in under the lock, so a
    # reader never sees a half-built dict.
    fresh: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for key, data in all_data.items():
        if key == "__poll_meta__" or not isinstance(data, dict):
            continue
        meta = data.get("_meta", {})
        ip = meta.get("source_ip") or data.get("__ip__") or "unknown"
        tname = meta.get("table_name") or data.get("__table__") or "unknown"
        fresh.setdefault(ip, {})[tname] = data
    order = dedupe_ips(ip_list if ip_list is not None else IP_LIST)
    for ip in order:
        fresh.setdefault(ip, {})

    first_ip = next((ip for ip in order if fresh.get(ip)), next(iter(fresh), None))
    with _ADD_LOCK:
        ALL_DEVICE_DATA.clear()
        ALL_DEVICE_DATA.update(fresh)
        dev = fresh.get(first_ip, {}) if first_ip is not None else {}
        Device_Configuration_Table = dev.get("Device_Configuration_Table", {})
        Communications_Configuration_Table = dev.get("Communications_Configuration_Table", {})
        Voltage_Current_Table = dev.get("Voltage_Current_Table", {})
        Real_Time_Power_Table = dev.get("Real_Time_Power_Table", {})
        Cumulative_Power_Table = dev.get("Cumulative_Power_Table", {})
        Demand_Data_Table = dev.get("Demand_Data_Table", {})
        Diagnostic_Table = dev.get("Diagnostic_Table", {})
        Voltage_Current_Snapshot_Log_Table = dev.get("Voltage_Current_Snapshot_Log_Table", {})
        Power_Snapshot_Log_Table = dev.get("Power_Snapshot_Log_Table", {})
        MinMax_Log_Table = dev.get("MinMax_Log_Table", {})
        Diagnostic_Table_Extended = dev.get("Diagnostic_Table_Extended", {})

    # The store is fed from the private dict, never from the shared one.
    accumulate_poll(fresh)
    logger.info("Named dicts updated from %d device(s).", len(fresh))


# %% ===========================================================================
#    NUMERIC SERIES OF ONE TABLE
# =============================================================================
def extract_numeric_series(table_dict: Dict[str, Any]) -> Dict[str, Tuple[float, str]]:
    """``{parameter: (value, unit)}`` for every finite numeric row of a table dict.

    A parameter name that occurs twice in one table (e.g. "Reserved") is kept
    for every row: the second and later occurrences are named "Name [#N]"
    with N the meter's row number (1.5.x silently kept only the last one).
    """
    series: Dict[str, Tuple[float, str]] = {}
    for key, name in table_dict.items():
        if (not isinstance(key, str) or key.startswith("_")
                or key.endswith("_value") or key.endswith("_unit")):
            continue
        value = table_dict.get(f"{key}_value")
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        if isinstance(value, float) and not math.isfinite(value):
            continue
        label = str(name)
        if label in series:                        # duplicate name inside this table
            label = f"{label} [{key}]"
            counter = 2
            while label in series:
                label = f"{name} [{key}.{counter}]"
                counter += 1
        series[label] = (float(value), table_dict.get(f"{key}_unit", ""))
    return series


# %% ===========================================================================
#    TIME-SERIES STORE
# =============================================================================
# Public layout, unchanged since 1.x (callers and the exporter rely on it):
#   TIME_SERIES_STORE[ip][table] = {
#       "timestamps_local": [str, ...],          # local wall clock, one per row
#       "columns":          {param: [float|None, ...]},
#       "units":            {param: str},
#   }
# Private row metadata kept beside it, aligned with timestamps_local:
#   _TS_META[ip][table] = {"seq": [int, ...], "epoch": [float, ...]}
# "seq" is a process-wide increasing row number; writers remember the last
# seq they persisted. If a caller trims or clears the store directly, the
# metadata is re-aligned to the tail automatically (_aligned_meta).
TIME_SERIES_STORE: Dict[str, Dict[str, Dict[str, Any]]] = {}
_TS_META: Dict[str, Dict[str, Dict[str, List[Any]]]] = {}
# Re-entrant so helpers that lock can be called while the lock is held.
# Callers using "with abm._TS_LOCK:" keep working (RLock is a context manager).
_TS_LOCK = threading.RLock()
_SEQ_COUNTER = itertools.count(1)        # next() is atomic under the GIL and the lock

# Number of RAM flushes in the current run. Kept under its 1.5.x name.
_VEUSZ_FLUSH_COUNT: int = 0


def _aligned_meta(ip: str, tname: str, tstore: Dict[str, Any]) -> Dict[str, List[Any]]:
    """Row metadata of one table, repaired to match timestamps_local (lock held).

    Rows removed from the head of the store by outside code are dropped from
    the metadata too. Rows without metadata (added by outside code) get seq 0,
    i.e. they count as already persisted and are never written twice.
    """
    n_rows = len(tstore.get("timestamps_local", []))
    meta = _TS_META.setdefault(ip, {}).get(tname)
    if meta is None:
        timestamps = tstore.get("timestamps_local", [])
        meta = {"seq": [0] * n_rows, "epoch": [_epoch_from_local_text(t) for t in timestamps]}
        _TS_META[ip][tname] = meta
        return meta
    for key in ("seq", "epoch"):
        values = meta[key]
        if len(values) > n_rows:
            del values[:len(values) - n_rows]
        elif len(values) < n_rows:
            values[:0] = [0 if key == "seq" else float("nan")] * (n_rows - len(values))
    return meta


def _fetch_epoch(table_dict: Dict[str, Any], fallback: float) -> float:
    """Unix time a table was fetched: _meta fetch_epoch, else fetch_utc, else fallback."""
    meta = table_dict.get("_meta", {})
    epoch = meta.get("fetch_epoch")
    if isinstance(epoch, (int, float)) and math.isfinite(epoch):
        return float(epoch)
    text = meta.get("fetch_utc", "")
    try:
        return datetime.datetime.strptime(text, "%Y-%m-%dT%H:%M:%SZ").replace(
            tzinfo=datetime.timezone.utc).timestamp()
    except (TypeError, ValueError):
        return fallback


def accumulate_poll(all_device_data: Dict[str, Dict[str, Dict[str, Any]]]) -> None:
    """Append one poll ({ip: {table: table_dict}}) to TIME_SERIES_STORE.

    Each table row is stamped with that page's own fetch time (local wall
    clock text, as before) and gets a sequence number and Unix time in the
    private metadata. A parameter seen for the first time is back-filled with
    None for earlier rows; a parameter missing from a poll gets None.
    """
    fallback = time.time()
    with _TS_LOCK:
        for ip, tables in all_device_data.items():
            for tname, table_dict in tables.items():
                series = extract_numeric_series(table_dict)
                if not series:
                    continue                    # failed page or no numeric rows
                epoch = _fetch_epoch(table_dict, fallback)
                ip_store = TIME_SERIES_STORE.setdefault(ip, {})
                tstore = ip_store.get(tname)
                if tstore is None:
                    tstore = {"timestamps_local": [], "columns": {}, "units": {}}
                    ip_store[tname] = tstore
                    _TS_META.setdefault(ip, {})[tname] = {"seq": [], "epoch": []}
                meta = _aligned_meta(ip, tname, tstore)
                n_rows = len(tstore["timestamps_local"])
                columns = tstore["columns"]
                for param, (_value, unit) in series.items():
                    if param not in columns:
                        columns[param] = [None] * n_rows      # back-fill earlier rows
                        tstore["units"][param] = unit
                tstore["timestamps_local"].append(_local_text(epoch))
                for param, column in columns.items():
                    if len(column) != n_rows:                  # repair outside edits
                        if len(column) < n_rows:
                            column.extend([None] * (n_rows - len(column)))
                        else:
                            del column[:len(column) - n_rows]
                    entry = series.get(param)
                    column.append(entry[0] if entry is not None else None)
                meta["seq"].append(next(_SEQ_COUNTER))
                meta["epoch"].append(epoch)
    logger.debug("Accumulator updated - %d device(s)", len(all_device_data))


def _drop_head_rows(ip: str, tname: str, count: int) -> int:
    """Remove the oldest ``count`` rows of one table (lock held). Returns rows removed."""
    tstore = TIME_SERIES_STORE.get(ip, {}).get(tname)
    if tstore is None or count <= 0:
        return 0
    meta = _aligned_meta(ip, tname, tstore)
    count = min(count, len(tstore["timestamps_local"]))
    del tstore["timestamps_local"][:count]
    for column in tstore["columns"].values():
        del column[:count]
    del meta["seq"][:count]
    del meta["epoch"][:count]
    return count


def trim_store(max_samples: int) -> int:
    """Keep only the newest ``max_samples`` rows of every table (<= 0: no trim).

    Trims the public store and its private metadata together - use this
    instead of slicing TIME_SERIES_STORE lists by hand. Returns rows removed.
    """
    if max_samples <= 0:
        return 0
    removed = 0
    with _TS_LOCK:
        for ip, tables in TIME_SERIES_STORE.items():
            for tname, tstore in tables.items():
                excess = len(tstore["timestamps_local"]) - int(max_samples)
                if excess > 0:
                    removed += _drop_head_rows(ip, tname, excess)
    return removed


def _clear_time_series_store() -> None:
    """Clear the whole store and its metadata, then release memory (1.5.x API)."""
    with _TS_LOCK:
        TIME_SERIES_STORE.clear()
        _TS_META.clear()
    release_memory()
    logger.info("[RAM Flush] TIME_SERIES_STORE cleared.")


def store_snapshot(ip: Optional[str] = None, tname: Optional[str] = None,
                   max_points: Optional[int] = None) -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Independent copy of (part of) the store, same layout plus "epoch_utc".

    Parameters
    ----------
    ip, tname : optional
        Restrict to one device and/or one table.
    max_points : int, optional
        Keep at most about this many rows per table: an evenly spaced subset
        that always contains the newest row (for previews; files never use it).
    """
    out: Dict[str, Dict[str, Dict[str, Any]]] = {}
    with _TS_LOCK:
        for dev, tables in TIME_SERIES_STORE.items():
            if ip is not None and dev != ip:
                continue
            for name, tstore in tables.items():
                if tname is not None and name != tname:
                    continue
                meta = _aligned_meta(dev, name, tstore)
                n_rows = len(tstore["timestamps_local"])
                if max_points and n_rows > max_points:
                    step = int(math.ceil(n_rows / float(max_points)))
                    index = list(range(n_rows - 1, -1, -step))[::-1]
                    pick = lambda seq, _i=index: [seq[i] for i in _i]   # noqa: E731
                else:
                    pick = list
                out.setdefault(dev, {})[name] = {
                    "timestamps_local": pick(tstore["timestamps_local"]),
                    "columns": {p: pick(v) for p, v in tstore["columns"].items()},
                    "units": dict(tstore["units"]),
                    "epoch_utc": pick(meta["epoch"]),
                }
    return out


# %% ===========================================================================
#    MEMORY FIGURES
# =============================================================================
#: Bytes per stored value and per row for the store estimate (CPython 64-bit):
#: float object 24 B + list slot 8 B; per row a 19-character str (~68 B), an
#: int (~28 B), a float (24 B) and three list slots.
_BYTES_PER_VALUE = 32
_BYTES_PER_ROW = 144


def ts_store_size_estimate_mb() -> float:
    """Fast estimate of the store's memory (MB): O(tables), not O(values).

    bytes ~= sum over tables of n_rows * (32 * n_params + 144). Used every poll;
    ts_store_size_mb() walks every object and is kept for diagnostics only.
    """
    total = 0
    with _TS_LOCK:
        for tables in TIME_SERIES_STORE.values():
            for tstore in tables.values():
                n_rows = len(tstore["timestamps_local"])
                total += n_rows * (_BYTES_PER_VALUE * len(tstore["columns"]) + _BYTES_PER_ROW)
    return total / (1024.0 * 1024.0)


def ts_store_size_mb() -> float:
    """Exact-ish store size (MB) by walking every object with sys.getsizeof.

    Slow (visits every value) - diagnostics only; the flush logic uses
    ts_store_size_estimate_mb().
    """
    seen: Set[int] = set()

    def _size(obj: Any) -> int:
        if id(obj) in seen:
            return 0
        seen.add(id(obj))
        size = sys.getsizeof(obj)
        if isinstance(obj, dict):
            size += sum(_size(k) + _size(v) for k, v in obj.items())
        elif isinstance(obj, (list, tuple)):
            size += sum(_size(item) for item in obj)
        return size

    with _TS_LOCK:
        return (_size(TIME_SERIES_STORE) + _size(_TS_META)) / (1024.0 * 1024.0)


def system_free_ram_mb() -> float:
    """Available system RAM in MB (9999 when psutil is missing: never triggers)."""
    try:
        if psutil is not None:
            return psutil.virtual_memory().available / (1024.0 * 1024.0)
    except Exception:
        pass
    return 9999.0


def system_ram_used_pct() -> float:
    """System RAM use in percent (0.0 when psutil is missing: never triggers)."""
    try:
        if psutil is not None:
            return float(psutil.virtual_memory().percent)
    except Exception:
        pass
    return 0.0


def process_rss_mb() -> float:
    """Resident set size of THIS process in MB (0.0 when psutil is missing).

    On Windows psutil reports the working set here (Task Manager "Memory").
    """
    try:
        if psutil is not None:
            return psutil.Process(os.getpid()).memory_info().rss / (1024.0 * 1024.0)
    except Exception:
        pass
    return 0.0


def process_commit_mb() -> float:
    """Windows commit charge (private bytes) of THIS process in MB; 0.0 elsewhere.

    psutil maps memory_info().vms to the Windows "pagefile" counter, the
    private committed memory (Task Manager "Commit size"). It only falls
    when heap memory is decommitted or released, not when the working set is
    trimmed, so it shows what the heap release actually achieved.
    """
    if sys.platform != "win32":
        return 0.0
    try:
        if psutil is not None:
            return psutil.Process(os.getpid()).memory_info().vms / (1024.0 * 1024.0)
    except Exception:
        pass
    return 0.0


# %% Returning freed memory to the operating system ------------------------------
# Why this is needed: dropping rows frees Python objects, but every allocator
# underneath keeps (some of) the freed memory for re-use:
#   * pymalloc (objects <= 512 B, e.g. the float values) returns a 1 MiB arena
#     with VirtualFree / munmap only when EVERY object in it is free;
#   * the C runtime heap (larger blocks, e.g. list buffers) keeps freed blocks
#     committed - glibc in its arenas, Windows in the NT heap and its
#     low-fragmentation-heap (LFH) caches;
#   * on Windows, even released pages stay in the process's WORKING SET (what
#     Task Manager shows as "Memory") until the memory manager trims it.
# release_memory() therefore runs gc.collect() and then the OS-specific calls
# below. All calls are best effort: a failure is logged at DEBUG level and
# never raises. They are serialised by _RELEASE_LOCK (a flush from the GUI
# button and an automatic one cannot overlap).
_RELEASE_LOCK = threading.Lock()
_OS_MEM_API: Optional[Dict[str, Any]] = None      # ctypes bindings, built once


def _os_memory_api() -> Dict[str, Any]:
    """Bind the OS memory-release functions once (cached in _OS_MEM_API).

    Private ctypes.WinDLL / CDLL instances are used, so the argtypes and
    restype set here never leak into ctypes.windll / ctypes.cdll objects that
    other libraries share. 64-bit HANDLE and SIZE_T arguments are declared
    explicitly (ctypes would otherwise pass a C int and truncate them).
    """
    global _OS_MEM_API
    if _OS_MEM_API is not None:
        return _OS_MEM_API
    import ctypes
    api: Dict[str, Any] = {"errors": []}
    if sys.platform == "win32":
        from ctypes import wintypes
        try:
            kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
            get_current_process = kernel32.GetCurrentProcess
            get_current_process.argtypes = []
            get_current_process.restype = wintypes.HANDLE
            api["process"] = get_current_process()      # pseudo handle, all access rights
            # BOOL HeapSetInformation(HANDLE, HEAP_INFORMATION_CLASS, PVOID, SIZE_T) [R8]
            heap_set = kernel32.HeapSetInformation
            heap_set.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, ctypes.c_size_t]
            heap_set.restype = wintypes.BOOL
            api["heap_set_information"] = heap_set

            class _HeapOptimizeResourcesInformation(ctypes.Structure):
                """HEAP_OPTIMIZE_RESOURCES_INFORMATION {DWORD Version; DWORD Flags;} (winnt.h)."""
                _fields_ = [("Version", wintypes.DWORD), ("Flags", wintypes.DWORD)]
            api["heap_optimize_info_type"] = _HeapOptimizeResourcesInformation
            # BOOL K32EmptyWorkingSet(HANDLE) - kernel32 on Windows 7+ [R10]
            empty_ws = getattr(kernel32, "K32EmptyWorkingSet", None)
            if empty_ws is None:                             # pre-Windows-7: psapi.dll
                empty_ws = ctypes.WinDLL("psapi", use_last_error=True).EmptyWorkingSet
            empty_ws.argtypes = [wintypes.HANDLE]
            empty_ws.restype = wintypes.BOOL
            api["empty_working_set"] = empty_ws
        except (OSError, AttributeError) as exc:
            api["errors"].append(f"kernel32: {exc}")
        # Not bound on purpose: the CRT's _heapmin(). In the Universal CRT
        # (python.exe since 3.5) it is HeapCompact(__acrt_heap, 0), which only
        # reports the largest free block and "does not compact the heap any
        # further" [R9] - it would log "ok" while releasing nothing.
    elif sys.platform.startswith("linux"):
        try:
            malloc_trim = ctypes.CDLL("libc.so.6").malloc_trim   # glibc only [R4]
            malloc_trim.argtypes = [ctypes.c_size_t]
            malloc_trim.restype = ctypes.c_int
            api["malloc_trim"] = malloc_trim
        except (OSError, AttributeError) as exc:         # musl, sandboxing ...
            api["errors"].append(f"libc: {exc}")
    _OS_MEM_API = api
    return api


def return_memory_to_os(trim_working_set: bool = True) -> Dict[str, Any]:
    """Hand freed heap memory back to the operating system (best effort).

    Call after the large objects are gone (perform_ram_flush() does). Runs
    gc.collect() first. Never raises.

    Windows (in this order)
      1. HeapSetInformation(NULL, HeapOptimizeResources = 3,
         {Version = 1, Flags = 0}) - every heap with an LFH has its caches
         optimised and free memory decommitted "if possible" [R8]
         (Windows 8.1+; the same call Firefox makes under memory pressure).
      2. EmptyWorkingSet(GetCurrentProcess()) when trim_working_set -
         "removes as many pages as possible from the working set" [R10]:
         the physical RAM goes back to Windows at once.
      Arenas of CPython's own allocator (pymalloc) are already returned with
      VirtualFree as soon as they are completely empty.
    Linux
      malloc_trim(0) - returns free memory at the top of the heap and free
      whole pages inside the arenas to the kernel [R4].

    Returns
    -------
    dict
        platform, released (True when any call succeeded), steps
        {name: True/False}, errors [text], rss_mb_before/after,
        commit_mb_before/after (Windows), skipped (another release running).
    """
    report: Dict[str, Any] = {"platform": sys.platform, "released": False, "steps": {},
                              "errors": [], "skipped": False}
    if not _RELEASE_LOCK.acquire(blocking=False):
        report["skipped"] = True                     # a concurrent release does the job
        return report
    try:
        import ctypes
        report["rss_mb_before"], report["commit_mb_before"] = process_rss_mb(), process_commit_mb()
        gc.collect()
        api = _os_memory_api()
        report["errors"].extend(api["errors"])
        steps = report["steps"]
        if sys.platform == "win32":
            heap_set = api.get("heap_set_information")
            if heap_set is not None:
                # HeapOptimizeResources = 3; HEAP_OPTIMIZE_RESOURCES_CURRENT_VERSION = 1;
                # HeapHandle NULL = every heap of the process that has an LFH [R8].
                info = api["heap_optimize_info_type"](1, 0)
                steps["heap_optimize"] = bool(heap_set(None, 3, ctypes.byref(info), ctypes.sizeof(info)))
                if not steps["heap_optimize"]:          # e.g. Windows 7: class 3 unknown
                    report["errors"].append(f"HeapSetInformation: Win32 error {ctypes.get_last_error()}")
            empty_ws = api.get("empty_working_set")
            if trim_working_set and empty_ws is not None and api.get("process") is not None:
                steps["working_set_trim"] = bool(empty_ws(api["process"]))
                if not steps["working_set_trim"]:
                    report["errors"].append(f"EmptyWorkingSet: Win32 error {ctypes.get_last_error()}")
        elif api.get("malloc_trim") is not None:
            # malloc_trim returns 1 if memory was released, 0 if there was none to release.
            steps["malloc_trim"] = bool(api["malloc_trim"](0))
        report["released"] = any(steps.values())
        report["rss_mb_after"], report["commit_mb_after"] = process_rss_mb(), process_commit_mb()
    except Exception as exc:                          # never let memory hygiene stop a run
        report["errors"].append(f"{type(exc).__name__}: {exc}")
    finally:
        _RELEASE_LOCK.release()
    if report["errors"]:
        logger.debug("return_memory_to_os: %s", "; ".join(report["errors"]))
    return report


def describe_memory_release(report: Dict[str, Any]) -> str:
    """One-line human summary of a return_memory_to_os() report (logs, GUI)."""
    if not report:
        return "OS release off"
    if report.get("skipped"):
        return "OS release already running"
    steps = ", ".join(f"{name} {'ok' if ok else 'no'}" for name, ok in report.get("steps", {}).items())
    text = f"OS release [{steps or 'gc only'}]"
    if "rss_mb_after" in report:
        text += f"; RSS {report['rss_mb_before']:.0f} -> {report['rss_mb_after']:.0f} MB"
    if report.get("commit_mb_before"):
        text += f", commit {report['commit_mb_before']:.0f} -> {report['commit_mb_after']:.0f} MB"
    return text


def release_memory(cfg: Optional[Dict[str, Any]] = None) -> bool:
    """gc.collect() and, with MEM_RETURN_TO_OS, return_memory_to_os() (1.5.x API).

    Returns True when the operating system took memory back (Linux: malloc_trim
    released some; Windows: any release call succeeded). Use
    return_memory_to_os() directly for the detailed report.
    """
    cfg = default_config() if cfg is None else cfg
    if not bool(cfg.get("mem_return_to_os", MEM_RETURN_TO_OS)):
        gc.collect()
        return False
    return bool(return_memory_to_os(bool(cfg.get("mem_trim_working_set", MEM_TRIM_WORKING_SET)))
                .get("released"))


# %% ===========================================================================
#    PERSISTENCE ENGINE (sequence cursors, one in-flight write per output)
# =============================================================================
# Each output ("writer": csv, md, vszh5, fits, xlsx) keeps, per output folder
# and per (ip, table), the highest row sequence number it has persisted. A
# persist call copies only rows newer than the LOWEST cursor of the writers it
# runs (one small snapshot, taken once under the store lock), then every writer
# gets the rows newer than its own cursor. Writers run in parallel on a
# persistent thread pool (they write different files); a writer whose previous
# write is still running is skipped for that poll (never two writes to one
# file at the same time), and its cursor only moves when a write succeeded.
_PERSIST_LOCK = threading.RLock()     # one persist / flush at a time (re-entrant)
_CURSOR_LOCK = threading.Lock()
_CURSORS: Dict[Tuple[str, str], Dict[Tuple[str, str], int]] = {}
_INFLIGHT: Dict[Tuple[str, str], "concurrent.futures.Future"] = {}
_FRESH_LOCK = threading.Lock()
_RUN_STATE: Dict[str, Any] = {"active": False, "run_id": 0, "cycle": 0,
                              "fresh": set(), "started_utc": ""}
_FLUSH_STATE: Dict[str, Any] = {"count": 0, "last_mono": None, "last_reason": "",
                                "last_warn_mono": None}
_POOL_LOCK = threading.Lock()
_WRITER_POOL: Optional[concurrent.futures.ThreadPoolExecutor] = None
_FLUSH_MGR: Optional[concurrent.futures.ThreadPoolExecutor] = None
#: Longest wait (s) for one writer when the caller needs the data on disk.
WRITER_TIMEOUT_SEC = 600.0
#: Writers that append (cheap): their result is always awaited.
_INCREMENTAL_WRITERS = ("csv", "md", "vszh5")


def _writer_pool() -> concurrent.futures.ThreadPoolExecutor:
    """Persistent pool for the file writers (created on first use)."""
    global _WRITER_POOL
    with _POOL_LOCK:
        if _WRITER_POOL is None:
            _WRITER_POOL = concurrent.futures.ThreadPoolExecutor(
                max_workers=5, thread_name_prefix="ab_writer")
        return _WRITER_POOL


def _flush_manager() -> concurrent.futures.ThreadPoolExecutor:
    """Single-thread executor behind the 1.5.x flush_outputs_parallel() API."""
    global _FLUSH_MGR
    with _POOL_LOCK:
        if _FLUSH_MGR is None:
            _FLUSH_MGR = concurrent.futures.ThreadPoolExecutor(
                max_workers=1, thread_name_prefix="ab_flush_mgr")
        return _FLUSH_MGR


def _shutdown_pools() -> None:
    """At interpreter exit: let running writes finish, accept no new ones."""
    for pool in (_FLUSH_MGR, _WRITER_POOL):
        if pool is not None:
            try:
                pool.shutdown(wait=True)
            except Exception:
                pass


atexit.register(_shutdown_pools)


def _cursor_key(writer: str, out_dir: str) -> Tuple[str, str]:
    """Cursor identity: the writer and the absolute, case-normalised folder."""
    return (writer, os.path.normcase(os.path.abspath(out_dir)))


def _is_fresh_pending(path: str, cfg: Dict[str, Any]) -> bool:
    """True when this run must start ``path`` afresh (append_files = 0, not yet done)."""
    if bool(cfg.get("append_files", APPEND_OUTPUT_FILES)):
        return False
    with _FRESH_LOCK:
        return os.path.abspath(path) not in _RUN_STATE["fresh"]


def _mark_fresh_done(path: str) -> None:
    """Record that ``path`` was started afresh in this run."""
    with _FRESH_LOCK:
        _RUN_STATE["fresh"].add(os.path.abspath(path))


def _fsync(handle: Any, cfg: Dict[str, Any]) -> None:
    """Flush a text file to disk when FSYNC_OUTPUTS is on (survives power loss)."""
    if bool(cfg.get("fsync_outputs", FSYNC_OUTPUTS)):
        handle.flush()
        os.fsync(handle.fileno())


def _temp_sibling(path: str) -> str:
    """Private temporary file next to ``path`` (same file system for os.replace)."""
    folder = os.path.dirname(os.path.abspath(path))
    os.makedirs(folder, exist_ok=True)
    handle, name = tempfile.mkstemp(prefix=os.path.basename(path) + ".", suffix=".part", dir=folder)
    os.close(handle)
    return name


def _remove_quietly(path: str) -> None:
    """os.remove() that ignores a missing file or a sharing violation."""
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except OSError:
        pass


# %% Files held open by another program (Excel on Windows) -------------------------
# Excel opens a CSV / XLSX with "deny write": other programs may read it, but an
# append, a rename or a replace fails (PermissionError; Win32 error 5 access
# denied, 32 sharing violation, 33 lock violation). Such a file is treated as
# "in use": its rows stay in memory (cursor not advanced) and are written when
# the file is free again. A forced persist (RAM flush, Save Now, stop) writes
# them to a "<file>_<stamp>_rescue<ext>" file instead, which is merged back
# into the main file on the next successful write (1.6.2).
_IN_USE_WINERRORS = {5, 32, 33}
_IN_USE_WARNED: Dict[str, float] = {}       # path -> monotonic time of the last warning


def _file_in_use(exc: BaseException) -> bool:
    """True when ``exc`` means another program holds the file open (or it is read-only)."""
    return isinstance(exc, PermissionError) or getattr(exc, "winerror", None) in _IN_USE_WINERRORS


def _warn_in_use(path: str, exc: BaseException, every_sec: float = 600.0) -> None:
    """Warn that ``path`` is in use - at most once every ``every_sec`` per file."""
    now = time.monotonic()
    last = _IN_USE_WARNED.get(path)
    if last is None or now - last >= every_sec:
        _IN_USE_WARNED[path] = now
        logger.warning("%s is open in another program (%s) - its new rows are kept in memory and "
                       "written when it is closed (a flush / Save Now / stop writes a rescue file).",
                       path, exc)


def _rescue_name(path: str) -> str:
    """A new, unused "<stem>_<YYYYmmdd_HHMMSS>[_N]_rescue<ext>" next to ``path``."""
    root, ext = os.path.splitext(path)
    candidate, counter = f"{root}_{_stamp()}_rescue{ext}", 2
    while os.path.exists(candidate):
        candidate, counter = f"{root}_{_stamp()}_{counter}_rescue{ext}", counter + 1
    return candidate


def _rescue_files(path: str) -> List[str]:
    """Rescue files of exactly ``path`` (oldest first).

    A regular expression, not a glob: "Diagnostic_Table_*" would also match the
    rescue files of "Diagnostic_Table_Extended".
    """
    folder = os.path.dirname(path) or "."
    root, ext = os.path.splitext(os.path.basename(path))
    pattern = re.compile(re.escape(root) + r"_(\d{8}_\d{6})(?:_(\d+))?_rescue" + re.escape(ext) + r"$")
    try:
        found = [(m.group(1), int(m.group(2) or 1), n)
                 for n in os.listdir(folder) for m in (pattern.match(n),) if m]
    except OSError:
        return []
    # Sort by (time stamp, counter): a plain name sort would put "..._2_rescue"
    # BEFORE the first rescue file of the same second ("_2" < "_r").
    return [os.path.join(folder, n) for _stamp_text, _count, n in sorted(found)]


def _claim_rescue(rescue_path: str) -> Optional[str]:
    """Rename a rescue file to "<name>.merging" before its rows are merged.

    The renamed file no longer matches _rescue_files(), so a merge that
    succeeded is never repeated even if deleting the file afterwards fails (the
    file is then simply left on disk). None if the file cannot be renamed now
    (e.g. it is open in Excel itself) - it is then merged on a later write.
    """
    claimed = rescue_path + ".merging"
    try:
        os.replace(rescue_path, claimed)
    except OSError:
        return None
    return claimed


def _rollover(path: str, reason: str) -> Optional[str]:
    """Rename ``path`` to <stem>_<stamp>_closed<ext>; None if it cannot be moved now."""
    if not os.path.exists(path):
        return None
    root, ext = os.path.splitext(path)
    closed = f"{root}_{_stamp()}_closed{ext}"
    try:
        os.replace(path, closed)
    except OSError as exc:               # e.g. open in another program on Windows
        logger.warning("Rollover of %s postponed (%s): %s", path, reason, exc)
        return None
    logger.info("Rolled over %s -> %s (%s)", os.path.basename(path), os.path.basename(closed), reason)
    return closed


def _collect_pending(min_cursor: Callable[[Tuple[str, str]], int]) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """Copy the rows newer than ``min_cursor(key)`` of every table (one lock hold).

    Returns {(ip, table): batch}; a batch holds ip, tname, ts, epoch, seq,
    columns {param: values}, units and params (store column order).
    """
    pending: Dict[Tuple[str, str], Dict[str, Any]] = {}
    with _TS_LOCK:
        for ip, tables in TIME_SERIES_STORE.items():
            for tname, tstore in tables.items():
                meta = _aligned_meta(ip, tname, tstore)
                seqs = meta["seq"]
                start = bisect.bisect_right(seqs, min_cursor((ip, tname)))
                if start >= len(seqs):
                    continue
                pending[(ip, tname)] = {
                    "ip": ip, "tname": tname,
                    "ts": tstore["timestamps_local"][start:],
                    "epoch": meta["epoch"][start:],
                    "seq": seqs[start:],
                    "columns": {p: col[start:] for p, col in tstore["columns"].items()},
                    "units": dict(tstore["units"]),
                    "params": list(tstore["columns"].keys()),
                }
    return pending


def _slice_batch(batch: Dict[str, Any], cursor: int) -> Optional[Dict[str, Any]]:
    """The rows of ``batch`` newer than ``cursor``, or None when there are none."""
    start = bisect.bisect_right(batch["seq"], cursor)
    if start >= len(batch["seq"]):
        return None
    if start == 0:
        return batch
    sliced = dict(batch)
    for key in ("ts", "epoch", "seq"):
        sliced[key] = batch[key][start:]
    sliced["columns"] = {p: values[start:] for p, values in batch["columns"].items()}
    return sliced


def _writer_specs(cfg: Dict[str, Any], full: bool) -> List[Tuple[str, str, Callable[..., Dict]]]:
    """(writer name, folder, function) of the outputs to persist in this call.

    CSV, Markdown and .vszh5 append cheaply and run every call (the .vszh5
    only when VEUSZ_APPEND_EVERY_CYCLE or ``full``); FITS and XLSX rewrite
    their files and run only when ``full`` (cadence, flush, Save Now, stop).
    """
    specs: List[Tuple[str, str, Callable[..., Dict]]] = []
    if cfg.get("enable_csv"):
        specs.append(("csv", cfg.get("csv_dir", CSV_DIR), _write_csv_batches))
    if cfg.get("enable_log_append"):
        specs.append(("md", cfg.get("log_dir", LOG_DIR), _write_md_batches))
    if cfg.get("enable_veusz") and (full or cfg.get("veusz_append_every_cycle", VEUSZ_APPEND_EVERY_CYCLE)):
        specs.append(("vszh5", cfg.get("veusz_dir", VEUSZ_DIR), _write_vszh5_batches))
    if full and cfg.get("enable_fits"):
        specs.append(("fits", cfg.get("fits_dir", FITS_DIR), _write_fits_batches))
    if full and cfg.get("enable_xlsx"):
        specs.append(("xlsx", cfg.get("xlsx_dir", XLSX_DIR), _write_xlsx_batches))
    return specs


def _advance_cursors(ckey: Tuple[str, str], done: Optional[Dict[Tuple[str, str], int]]) -> None:
    """Move a writer's cursors forward to the last sequence numbers it persisted."""
    with _CURSOR_LOCK:
        cursors = _CURSORS.setdefault(ckey, {})
        for table_key, seq in (done or {}).items():
            if seq > cursors.get(table_key, 0):
                cursors[table_key] = seq


def _run_writer(func: Callable[..., Dict], ckey: Tuple[str, str], writer: str, batches: Dict,
                out_dir: str, cfg: Dict[str, Any], must_persist: bool) -> Dict[Tuple[str, str], int]:
    """Writer task: write, then advance the cursors BEFORE the future completes.

    1.6.1 race fix. 1.6.0 advanced the cursors in a future done-callback, but
    concurrent.futures wakes the threads waiting in result() BEFORE it runs the
    done-callbacks. A forced persist (flush / Save Now / stop) that waited for a
    busy writer could therefore read the old cursor and hand the same rows to
    the writer a second time - duplicate rows in every output (all writers
    append or merge without de-duplication). Seen on Windows Python where
    writes are slower; possible on any OS. Advancing the cursors inside the task
    makes "future done" imply "cursors up to date".
    """
    try:
        done = func(batches, out_dir, cfg, must_persist)
    except Exception as exc:
        logger.error("%s writer failed: %s", writer.upper(), exc, exc_info=True)
        raise
    _advance_cursors(ckey, done)
    return done


def _submit_writer(func: Callable[..., Dict], ckey: Tuple[str, str], writer: str, batches: Dict,
                   out_dir: str, cfg: Dict[str, Any], must_persist: bool) -> "concurrent.futures.Future":
    """Run a writer task on the pool; inline when the pool refuses (interpreter exit)."""
    args = (func, ckey, writer, batches, out_dir, cfg, must_persist)
    try:
        return _writer_pool().submit(_run_writer, *args)
    except RuntimeError:
        future: concurrent.futures.Future = concurrent.futures.Future()
        try:
            future.set_result(_run_writer(*args))
        except Exception as exc:
            future.set_exception(exc)
        return future


def persist_outputs(cfg: Optional[Dict[str, Any]] = None, full: bool = False,
                    must_persist: bool = False, reason: str = "cycle",
                    only: Optional[Sequence[str]] = None) -> Dict[str, str]:
    """Write the rows each enabled output has not persisted yet.

    Parameters
    ----------
    cfg : dict, optional
        Run configuration (default_config() when None).
    full : bool
        Include the rewrite outputs (FITS, XLSX) and the .vszh5 even when
        VEUSZ_APPEND_EVERY_CYCLE is off.
    must_persist : bool
        The data must reach the disk now (RAM flush, Save Now, stop): wait for
        busy writers and for FITS/XLSX, and use rescue files when the normal
        file is locked. Otherwise FITS/XLSX run in the background and a locked
        or busy output simply catches up on a later poll.
    reason : str
        For the log.
    only : sequence of str, optional
        Restrict to these writer names.

    Returns
    -------
    dict
        writer -> short status text ("12 rows", "up to date", "busy", "error: ...").
    """
    cfg = default_config() if cfg is None else cfg
    specs = _writer_specs(cfg, full)
    if only is not None:
        specs = [spec for spec in specs if spec[0] in only]
    if not specs:
        return {}
    status: Dict[str, str] = {}
    with _PERSIST_LOCK:
        ckeys = [_cursor_key(name, out_dir) for name, out_dir, _ in specs]
        with _CURSOR_LOCK:
            cursor_maps = [dict(_CURSORS.get(key, {})) for key in ckeys]
        pending = _collect_pending(lambda tk: min(m.get(tk, 0) for m in cursor_maps))
        if not pending:
            return {name: "up to date" for name, _, _ in specs}
        awaited: Dict[str, Tuple["concurrent.futures.Future", int]] = {}
        # cursor_maps only bound ``pending`` (a superset of what each writer needs).
        # Each writer's own cursors are re-read below AFTER confirming that no
        # write of it is running: a write that finished after the snapshot has
        # already advanced them, and none can advance them now (_PERSIST_LOCK is
        # held, so nothing new is submitted).
        for name, out_dir, func in specs:
            ckey = _cursor_key(name, out_dir)
            previous = _INFLIGHT.get(ckey)
            if previous is not None and not previous.done():
                if not must_persist:
                    status[name] = "busy (previous write still running)"
                    continue
                try:                           # data must land: wait for it
                    previous.result(timeout=WRITER_TIMEOUT_SEC)
                except Exception:
                    pass                       # a failed write was logged by _run_writer
                if not previous.done():
                    # Never start a second write to the same output while one is
                    # running (duplicates, two writers on one file). The rows stay
                    # in memory (the flush keeps rows a writer has not confirmed).
                    status[name] = "still running"
                    logger.error("%s writer still running after %.0f s (%s) - not started twice",
                                 name.upper(), WRITER_TIMEOUT_SEC, reason)
                    continue
            with _CURSOR_LOCK:
                cursors = dict(_CURSORS.get(ckey, {}))
            batches = {}
            for table_key, batch in pending.items():
                sliced = _slice_batch(batch, cursors.get(table_key, 0))
                if sliced is not None:
                    batches[table_key] = sliced
            if not batches:
                status[name] = "up to date"
                continue
            rows = sum(len(b["seq"]) for b in batches.values())
            # The task itself advances this writer's cursors before it completes.
            future = _submit_writer(func, ckey, name, batches, out_dir, cfg, must_persist)
            _INFLIGHT[ckey] = future
            if must_persist or name in _INCREMENTAL_WRITERS:
                awaited[name] = (future, rows)
            else:
                status[name] = f"{rows} rows dispatched"
        for name, (future, rows) in awaited.items():
            try:
                done = future.result(timeout=WRITER_TIMEOUT_SEC)
                written = sum(1 for _ in done) if done else 0
                status[name] = f"{rows} rows" if written else "nothing written (see log)"
            except concurrent.futures.TimeoutError:
                status[name] = "still running"
                logger.error("%s writer still running after %.0f s (%s)", name.upper(),
                             WRITER_TIMEOUT_SEC, reason)
            except Exception as exc:
                status[name] = f"error: {exc}"
    logger.debug("persist_outputs(%s, full=%s): %s", reason, full, status)
    return status


def _group_by_ip(batches: Dict[Tuple[str, str], Dict[str, Any]]) -> Dict[str, Dict[Tuple[str, str], Dict[str, Any]]]:
    """{ip: {(ip, table): batch}} in table page order."""
    out: Dict[str, Dict[Tuple[str, str], Dict[str, Any]]] = {}
    for key in sorted(batches, key=lambda k: (k[0], _table_order(k[1]))):
        out.setdefault(key[0], {})[key] = batches[key]
    return out


# %% ===========================================================================
#    OUTPUT MODULE: CSV  (one file per device per table; appended)
# =============================================================================
# Layout (unchanged): row 1 "Timestamp_Local, <params>", row 2 "(units), <units>",
# then one row per poll. New parameters become new right-hand columns (the file
# is rewritten once, streamed, with blanks for earlier rows).
def _csv_read_header(path: str) -> Optional[Tuple[List[str], List[str]]]:
    """(params, units) from the first two lines, or None if not our CSV layout."""
    with open(path, "r", newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader, None)
        units = next(reader, None) or []
    if not header or header[0].strip() != "Timestamp_Local":
        return None
    params = [cell.strip() for cell in header[1:]]
    unit_row = list(units[1:1 + len(params)])
    unit_row += [""] * (len(params) - len(unit_row))
    return params, unit_row


def _csv_rows(batch: Dict[str, Any], params: Sequence[str]) -> Iterable[List[Any]]:
    """Rows of a batch in ``params`` column order; None becomes an empty cell."""
    columns = [batch["columns"].get(p) for p in params]
    for i, stamp in enumerate(batch["ts"]):
        row: List[Any] = [stamp]
        for column in columns:
            value = column[i] if column is not None else None
            row.append("" if value is None else value)
        yield row


def _csv_add_columns(path: str, params: List[str], units: List[str], new: List[str],
                     new_units: Dict[str, str]) -> None:
    """Rewrite ``path`` (streamed, atomic) with the ``new`` columns appended."""
    temp = _temp_sibling(path)
    try:
        with open(path, "r", newline="", encoding="utf-8") as src, \
                open(temp, "w", newline="", encoding="utf-8") as dst:
            reader, writer = csv.reader(src), csv.writer(dst)
            next(reader, None)                 # old header
            next(reader, None)                 # old units row
            writer.writerow(["Timestamp_Local"] + params + new)
            writer.writerow(["(units)"] + units + [new_units.get(p, "") for p in new])
            width = 1 + len(params)
            for row in reader:
                row = row[:width] + [""] * (width - len(row))
                writer.writerow(row + [""] * len(new))
        os.replace(temp, path)
    finally:
        _remove_quietly(temp)


def _csv_write_new(path: str, batch: Dict[str, Any], cfg: Dict[str, Any]) -> None:
    """Create ``path`` with the header, the units row and the rows of ``batch``."""
    params = list(batch["params"])
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Timestamp_Local"] + params)
        writer.writerow(["(units)"] + [batch["units"].get(p, "") for p in params])
        writer.writerows(_csv_rows(batch, params))
        _fsync(handle, cfg)


def _csv_batch_from_file(path: str) -> Optional[Dict[str, Any]]:
    """The rows of one of our CSV files (a rescue file) as a writer batch.

    Values stay text, so they are copied into the main file exactly as written.
    None when the file does not have our header layout.
    """
    with open(path, "r", newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header, units = next(reader, None), next(reader, None) or []
        if not header or header[0].strip() != "Timestamp_Local":
            return None
        params = [cell.strip() for cell in header[1:]]
        stamps: List[str] = []
        columns: Dict[str, List[Optional[str]]] = {p: [] for p in params}
        for row in reader:
            if not row:
                continue
            stamps.append(row[0])
            for j, p in enumerate(params):
                cell = row[j + 1] if len(row) > j + 1 else ""
                columns[p].append(cell if cell != "" else None)
    return {"params": params, "ts": stamps, "columns": columns,
            "units": {p: (units[j + 1] if len(units) > j + 1 else "") for j, p in enumerate(params)}}


def _csv_append_rows(path: str, batch: Dict[str, Any], cfg: Dict[str, Any], fresh: bool) -> None:
    """Append one batch to one CSV (creating, extending or - ``fresh`` - restarting it)."""
    header = None
    if not fresh and os.path.exists(path) and os.path.getsize(path) > 0:
        header = _csv_read_header(path)
        if header is None:                     # not our layout: keep it, start anew
            aside = f"{os.path.splitext(path)[0]}_{_stamp()}_unreadable.csv"
            os.replace(path, aside)
            logger.warning("CSV header not recognised - moved %s aside to %s", path, aside)
    if header is None:
        _csv_write_new(path, batch, cfg)
        if fresh:
            _mark_fresh_done(path)
        return
    params, units = header
    new = [p for p in batch["params"] if p not in params]
    if new:
        _csv_add_columns(path, params, units, new, batch["units"])
        params = params + new
    with open(path, "a", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows(_csv_rows(batch, params))
        _fsync(handle, cfg)


def _restore_rescue(claimed: str, original: str) -> None:
    """Undo _claim_rescue() after a failed merge (the file is merged on a later write)."""
    try:
        os.replace(claimed, original)
    except OSError:
        logger.error("Rescue file %s could not be renamed back - merge it by hand", claimed)


def _csv_append(path: str, batch: Dict[str, Any], cfg: Dict[str, Any]) -> None:
    """Append one batch to one CSV, merging its rescue files first (oldest first).

    Rescue rows are older than ``batch``, so they go in before it; in fresh
    mode the first of them starts the new file.
    """
    fresh = _is_fresh_pending(path, cfg)
    for rescue in _rescue_files(path):
        claimed = _claim_rescue(rescue)
        if claimed is None:
            continue                           # open elsewhere: merged on a later write
        try:
            rescued = _csv_batch_from_file(claimed)
            if rescued is None:
                logger.error("CSV rescue file %s has an unknown layout - left on disk", claimed)
                continue
            if rescued["ts"]:
                _csv_append_rows(path, rescued, cfg, fresh)
                fresh = False
        except Exception:
            _restore_rescue(claimed, rescue)
            raise
        _remove_quietly(claimed)
        logger.info("CSV: rescue file merged into %s (%d row(s)): %s", os.path.basename(path),
                    len(rescued["ts"]), os.path.basename(rescue))
    _csv_append_rows(path, batch, cfg, fresh)


def _write_csv_batches(batches: Dict, out_dir: str, cfg: Dict[str, Any],
                       must_persist: bool) -> Dict[Tuple[str, str], int]:
    """CSV writer: returns {(ip, table): last seq persisted}.

    A file open in another program (Excel) keeps its rows in memory; with
    ``must_persist`` they go to a rescue file that is merged back later.
    """
    os.makedirs(out_dir, exist_ok=True)
    done: Dict[Tuple[str, str], int] = {}
    for ip, ip_batches in _group_by_ip(batches).items():
        for key, batch in ip_batches.items():
            path = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_{batch['tname']}.csv")
            try:
                _csv_append(path, batch, cfg)
                done[key] = batch["seq"][-1]
                _IN_USE_WARNED.pop(path, None)
            except Exception as exc:
                if not _file_in_use(exc):
                    logger.error("CSV write failed for %s / %s: %s", ip, batch["tname"], exc)
                elif not must_persist:
                    _warn_in_use(path, exc)
                else:
                    rescue = _rescue_name(path)
                    try:
                        _csv_write_new(rescue, batch, cfg)
                        done[key] = batch["seq"][-1]
                        logger.warning("CSV: %s is in use - %d row(s) written to %s (merged back "
                                       "when the file is free)", path, len(batch["seq"]), rescue)
                    except Exception as exc2:
                        logger.error("CSV rescue write failed for %s: %s", rescue, exc2)
    if done:
        logger.info("CSV appended: %d table file(s) in %s", len(done), out_dir)
    return done


# %% ===========================================================================
#    OUTPUT MODULE: MARKDOWN LOG TABLES  (one .md per device per table)
# =============================================================================
# A GitHub-flavoured Markdown document per device and table. When the columns
# change, a new heading and table header start a new section in the same file
# (1.5.x appended rows that no longer matched the header).
_MD_HEADER_CACHE: Dict[str, List[str]] = {}


def _md_labels(params: Sequence[str], units: Dict[str, str]) -> List[str]:
    """Column labels: Timestamp_Local, then "Param (unit)" or "Param"."""
    return ["Timestamp_Local"] + [f"{p} ({units[p]})" if units.get(p) else p for p in params]


def _md_last_header(path: str) -> Optional[List[str]]:
    """Labels of the last table header in ``path`` (cached after the first read)."""
    key = os.path.abspath(path)
    if key in _MD_HEADER_CACHE:
        return _MD_HEADER_CACHE[key]
    last = None
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            stripped = line.strip()
            if stripped.startswith("| Timestamp_Local"):
                last = [cell.strip() for cell in stripped.strip("|").split("|")]
    _MD_HEADER_CACHE[key] = last or []
    return last


def _md_append(path: str, batch: Dict[str, Any], cfg: Dict[str, Any]) -> None:
    """Append one batch to one Markdown log, merging its rescue files first.

    A rescue file is a complete section (heading, table header, rows); it is
    copied to the end of the log, oldest first, before the new rows.
    """
    fresh = _is_fresh_pending(path, cfg)
    merged = False
    for rescue in _rescue_files(path):
        claimed = _claim_rescue(rescue)
        if claimed is None:
            continue                           # open elsewhere: merged on a later write
        try:
            exists = os.path.exists(path) and os.path.getsize(path) > 0
            with open(claimed, "r", encoding="utf-8") as src:
                text = src.read()
            with open(path, "w" if (fresh or not exists) else "a", encoding="utf-8") as dst:
                if exists and not fresh:
                    dst.write("\n")
                dst.write(text)
                _fsync(dst, cfg)
            _MD_HEADER_CACHE.pop(os.path.abspath(path), None)    # last header changed
            if fresh:
                _mark_fresh_done(path)
                fresh = False
        except Exception:
            _restore_rescue(claimed, rescue)
            raise
        _remove_quietly(claimed)
        merged = True
        logger.info("Markdown: rescue file merged into %s: %s", os.path.basename(path),
                    os.path.basename(rescue))
    # After a merged rescue section the new rows start their own "(continued)"
    # section, so they do not appear under the rescue section's heading.
    _md_append_rows(path, batch, cfg, fresh, new_section=" (continued)" if merged else "")


def _md_append_rows(path: str, batch: Dict[str, Any], cfg: Dict[str, Any], fresh: bool,
                    note: str = "", new_section: str = "") -> None:
    """Append one batch to one Markdown log (new file / new section as needed).

    ``note`` is added to the heading of a NEW file; ``new_section`` forces a new
    section (heading + table header) with that text even when the columns match.
    """
    params = list(batch["params"])
    labels = _md_labels(params, batch["units"])
    rows = [[stamp] + ["" if batch["columns"][p][i] is None else str(batch["columns"][p][i])
                       for p in params]
            for i, stamp in enumerate(batch["ts"])]
    widths = [max(3, len(label)) for label in labels]
    for row in rows:
        for j, cell in enumerate(row):
            widths[j] = max(widths[j], len(cell))

    def _line(cells: Sequence[str]) -> str:
        return "| " + " | ".join(str(c).ljust(widths[k]) for k, c in enumerate(cells)) + " |"

    separator = "|" + "|".join("-" * (w + 2) for w in widths) + "|"
    exists = os.path.exists(path) and os.path.getsize(path) > 0
    previous = None if (fresh or not exists) else _md_last_header(path)
    with open(path, "w" if (fresh or not exists) else "a", encoding="utf-8") as handle:
        if fresh or not exists:
            handle.write(f"## {batch['ip']} — {batch['tname']}{note}\n\n")
            handle.write(_line(labels) + "\n" + separator + "\n")
        elif previous != labels or new_section:
            heading = new_section or f" (columns changed {batch['ts'][0]})"
            handle.write(f"\n## {batch['ip']} — {batch['tname']}{heading}\n\n")
            handle.write(_line(labels) + "\n" + separator + "\n")
        handle.writelines(_line(row) + "\n" for row in rows)
        _fsync(handle, cfg)
    _MD_HEADER_CACHE[os.path.abspath(path)] = labels
    if fresh:
        _mark_fresh_done(path)


def _write_md_batches(batches: Dict, out_dir: str, cfg: Dict[str, Any],
                      must_persist: bool) -> Dict[Tuple[str, str], int]:
    """Markdown writer: returns {(ip, table): last seq persisted}.

    A log open in another program keeps its rows in memory; with
    ``must_persist`` they go to a rescue section file that is merged back later.
    """
    os.makedirs(out_dir, exist_ok=True)
    done: Dict[Tuple[str, str], int] = {}
    for ip, ip_batches in _group_by_ip(batches).items():
        for key, batch in ip_batches.items():
            path = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_{batch['tname']}.md")
            try:
                _md_append(path, batch, cfg)
                done[key] = batch["seq"][-1]
                _IN_USE_WARNED.pop(path, None)
            except Exception as exc:
                if not _file_in_use(exc):
                    logger.error("Markdown log write failed for %s / %s: %s", ip, batch["tname"], exc)
                elif not must_persist:
                    _warn_in_use(path, exc)
                else:
                    rescue = _rescue_name(path)
                    try:
                        _md_append_rows(rescue, batch, cfg, fresh=False,
                                        note=" (rows written while the log was open elsewhere)")
                        _MD_HEADER_CACHE.pop(os.path.abspath(rescue), None)
                        done[key] = batch["seq"][-1]
                        logger.warning("Markdown: %s is in use - %d row(s) written to %s (merged "
                                       "back when the file is free)", path, len(batch["seq"]), rescue)
                    except Exception as exc2:
                        logger.error("Markdown rescue write failed for %s: %s", rescue, exc2)
    if done:
        logger.info("Markdown logs appended: %d table file(s) in %s", len(done), out_dir)
    return done


# %% ===========================================================================
#    OUTPUT MODULE: VEUSZ .vszh5  (native, appended in place; vszh5_writer)
# =============================================================================
# One document per device: ABMeter_<ip>.vszh5. Per table a date dataset
# "<octet>_dt_<table>" (x axis) and one numeric dataset per parameter
# "<octet>_<table>_<param>" (collision-free, never truncated). Rows are
# appended every poll; the Veusz command script (pages) is rewritten only when
# a dataset is added. The table/parameter -> dataset index is kept inside the
# file under /VszAppend (ignored by Veusz), so appending resumes after a restart.
_VSZ_INDEX_KEY = "abmeter_index"
_VSZ_MIN_WRITER = (0, 0, 18)


def _vszh5_module() -> Any:
    """Import vszh5_writer (>= 0.0.18) from this folder; None (logged) if unusable."""
    try:
        import vszh5_writer as vw
    except ImportError:
        here = os.path.dirname(os.path.abspath(__file__))
        if here not in sys.path:
            sys.path.insert(0, here)
        try:
            import vszh5_writer as vw
        except ImportError as exc:
            logger.error("vszh5_writer.py not importable (needs numpy + h5py) - "
                         "Veusz output skipped: %s", exc)
            return None
    version = tuple(int(x) for x in re.findall(r"\d+", getattr(vw, "__version__", "0"))[:3])
    if version < _VSZ_MIN_WRITER or not hasattr(vw, "VeuszAppender"):
        logger.error("vszh5_writer %s is too old (need >= 0.0.18) - Veusz output skipped",
                     getattr(vw, "__version__", "?"))
        return None
    return vw


def _unique_name(base: str, used: Set[str]) -> str:
    """``base``, or ``base_2``, ``base_3`` ... - the first not in ``used`` (added to it)."""
    name, counter = base, 2
    while name in used:
        name = f"{base}_{counter}"
        counter += 1
    used.add(name)
    return name


def _veusz_seconds(vw: Any, epochs: Sequence[float], cfg: Dict[str, Any]) -> Any:
    """Unix times -> Veusz date seconds, local wall clock (DST-aware) or UTC."""
    values = np.asarray(epochs, dtype=float)
    if str(cfg.get("vszh5_time_axis", VSZH5_TIME_AXIS)).lower() == "utc":
        return vw.epoch_to_veusz(values)
    return vw.epoch_to_veusz(values, vw.local_utc_offsets(values))


def _vszh5_document(vw: Any, ip: str, index: Dict[str, Any], cfg: Dict[str, Any]) -> str:
    """Veusz command script: per-table grid pages and per-unit overlay pages."""
    doc = vw.VszDoc(generator=f"ab_power_meter_monitor {__version__} ({ip})")
    octet = _octet(ip)
    per_page = max(1, int(cfg.get("vszh5_panels_per_page", VSZH5_PANELS_PER_PAGE)))
    time_label = "UTC" if str(cfg.get("vszh5_time_axis", VSZH5_TIME_AXIS)).lower() == "utc" else "Local time"
    tables = sorted(index.get("tables", {}).items(), key=lambda kv: _table_order(kv[0]))
    for tname, tinfo in tables:
        params = list(tinfo.get("params", {}).items())
        chunks = [params[i:i + per_page] for i in range(0, len(params), per_page)]
        for number, chunk in enumerate(chunks, start=1):
            panels = []
            for param, dataset in chunk:
                unit = tinfo.get("units", {}).get(param, "")
                panels.append(vw.Panel(
                    [vw.Trace(tinfo["dt"], dataset, "", marker="circle", marker_size="1.2pt")],
                    xlabel=time_label, ylabel=f"{param} [{unit}]" if unit else param, title=param,
                    xmode="datetime", x_autorange="exact", key=False))
            vw.grid_page(doc, _ident(f"{octet}_{tname}_p{number}"), panels,
                         title=f"{ip}   {tname}   ({number}/{len(chunks)})",
                         columns=2, width="29.7cm", height="21cm")
    for group, substrings in VEUSZ_OVERLAY_GROUPS.items():
        traces = []
        for tname, tinfo in tables:
            for param, dataset in tinfo.get("params", {}).items():
                if any(s in param.lower() for s in substrings):
                    short = tname.replace("_Table", "").replace("_", " ")
                    traces.append(vw.Trace(tinfo["dt"], dataset, f"{short}: {param}"))
        if traces:
            vw.xy_page(doc, _ident(f"{octet}_overlay_{group}"), traces,
                       title=f"{ip}   Overlay: {group}", xlabel=time_label, ylabel=group,
                       xmode="datetime", x_autorange="exact", key=True,
                       width="29.7cm", height="21cm")
    return doc.text()


def _vszh5_open(vw: Any, path: str, ip: str, cfg: Dict[str, Any],
                must_persist: bool) -> Tuple[Optional[Any], str]:
    """Open (or create) the device document; (appender, path) or (None, path).

    A document locked by a reader is retried on the next poll, unless the rows
    must land now - then they go to a standalone rescue document. A document
    that cannot be read (e.g. truncated by a power cut) is moved aside to
    <name>.corrupt_<stamp> and a new one is started; nothing ever raises.
    """
    locking = cfg.get("vszh5_file_locking", VSZH5_FILE_LOCKING)
    tag = f"ab_power_meter_monitor {__version__}"
    rescue = os.path.join(os.path.dirname(path), f"ABMeter_{_safe_ip(ip)}_{_stamp()}_rescue.vszh5")
    try:
        return vw.VeuszAppender(path, locking=locking, writer_tag=tag), path
    except vw.AppendLockedError as exc:
        if not must_persist:
            logger.info("%s is open in another program (HDF5 lock); rows kept for the next poll.",
                        os.path.basename(path))
            return None, path
        logger.warning("%s locked (%s) - writing these rows to %s", path, exc, rescue)
        return vw.VeuszAppender(rescue, locking=locking, writer_tag=tag), rescue
    except (OSError, ValueError, KeyError) as exc:
        aside = f"{path}.corrupt_{_stamp()}"
        try:
            os.replace(path, aside)
            logger.error("%s unreadable (%s) - moved aside to %s; starting a new document.",
                         path, exc, aside)
            return vw.VeuszAppender(path, locking=locking, writer_tag=tag), path
        except OSError:
            if not must_persist:
                logger.error("%s unreadable (%s) and cannot be moved; retrying next poll.", path, exc)
                return None, path
            return vw.VeuszAppender(rescue, locking=locking, writer_tag=tag), rescue


def _vszh5_append_ip(vw: Any, ip: str, ip_batches: Dict, out_dir: str, cfg: Dict[str, Any],
                     must_persist: bool, path: Optional[str] = None) -> Dict[Tuple[str, str], int]:
    """Append the batches of one device to its .vszh5; returns persisted cursors."""
    path = path or os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}.vszh5")
    if _is_fresh_pending(path, cfg):
        try:
            _remove_quietly(path)
        finally:
            _mark_fresh_done(path)
    limit_mb = float(cfg.get("vszh5_max_file_mb", VSZH5_MAX_FILE_MB))
    if limit_mb > 0 and os.path.exists(path) and os.path.getsize(path) > limit_mb * 1024 * 1024:
        _rollover(path, f"larger than {limit_mb:.0f} MB")
    appender, used_path = _vszh5_open(vw, path, ip, cfg, must_persist)
    if appender is None:
        return {}
    done: Dict[Tuple[str, str], int] = {}
    octet = _octet(ip)
    with appender as app:
        raw = app.get_meta(_VSZ_INDEX_KEY)
        index = json.loads(raw) if raw else {}
        if not isinstance(index.get("tables"), dict):
            index = {"format": 1, "ip": ip, "tables": {}}
        used: Set[str] = set(app.names())
        for tinfo in index["tables"].values():
            used.add(tinfo.get("dt", ""))
            used.update(tinfo.get("params", {}).values())
        changed = raw is None
        for key, batch in ip_batches.items():
            tname = batch["tname"]
            tinfo = index["tables"].get(tname)
            if tinfo is None or not app.has(tinfo.get("dt", "")):
                tinfo = {"dt": _unique_name(f"{octet}_dt_{_ident(tname)}", used),
                         "params": {}, "units": {}}
                app.create_date(tinfo["dt"], tags=(f"dev_{octet}", _ident(tname)))
                index["tables"][tname] = tinfo
                changed = True
            start = app.length(tinfo["dt"])
            # Realign columns an interrupted append left longer than the time axis.
            for dataset in tinfo["params"].values():
                if app.has(dataset):
                    app.truncate(dataset, start)
            app.append(tinfo["dt"], _veusz_seconds(vw, batch["epoch"], cfg))
            stop = start + len(batch["seq"])
            for param in batch["params"]:
                values = [float("nan") if v is None else float(v) for v in batch["columns"][param]]
                dataset = tinfo["params"].get(param)
                if dataset is None or not app.has(dataset):
                    dataset = _unique_name(f"{octet}_{_ident(tname)}_{_ident(param)}", used)
                    app.create_1d(dataset, values, prefix_nan=start,
                                  tags=(f"dev_{octet}", _ident(tname)))
                    tinfo["params"][param] = dataset
                    tinfo["units"][param] = batch["units"].get(param, "")
                    changed = True
                else:
                    app.pad_to(dataset, start)
                    app.append(dataset, values)
            for dataset in tinfo["params"].values():   # columns absent from this batch
                app.pad_to(dataset, stop)
            done[key] = batch["seq"][-1]
        if changed:
            app.set_meta(_VSZ_INDEX_KEY, json.dumps(index, sort_keys=True))
            app.set_document(_vszh5_document(vw, ip, index, cfg))
    logger.info("Veusz .vszh5 appended (%d table(s)): %s", len(done), used_path)
    return done


def _write_vszh5_batches(batches: Dict, out_dir: str, cfg: Dict[str, Any],
                         must_persist: bool) -> Dict[Tuple[str, str], int]:
    """.vszh5 writer: returns {(ip, table): last seq persisted}."""
    vw = _vszh5_module()
    if vw is None or np is None:
        return {}
    os.makedirs(out_dir, exist_ok=True)
    done: Dict[Tuple[str, str], int] = {}
    for ip, ip_batches in _group_by_ip(batches).items():
        try:
            done.update(_vszh5_append_ip(vw, ip, ip_batches, out_dir, cfg, must_persist))
        except Exception as exc:
            logger.error("Veusz .vszh5 write failed for %s: %s\n%s", ip, exc, traceback.format_exc())
    return done


def _batches_from_snapshot(ip: str, tables: Dict[str, Dict[str, Any]]) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """Writer batches from a TIME_SERIES_STORE-shaped snapshot of one device."""
    out: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for tname, tdata in tables.items():
        stamps = list(tdata.get("timestamps_local", []))
        if not stamps:
            continue
        epochs = list(tdata.get("epoch_utc") or [_epoch_from_local_text(t) for t in stamps])
        out[(ip, tname)] = {"ip": ip, "tname": tname, "ts": stamps, "epoch": epochs,
                            "seq": list(range(1, len(stamps) + 1)),
                            "columns": {p: list(v) for p, v in tdata.get("columns", {}).items()},
                            "units": dict(tdata.get("units", {})),
                            "params": list(tdata.get("columns", {}).keys())}
    return out


# %% ===========================================================================
#    OUTPUT MODULE: FITS  (one file per device; rewritten atomically)
# =============================================================================
# Layout: primary HDU with site keywords; one BinTableHDU per table with
# TIMESTAMP_LOCAL (A19), EPOCH_UTC (D, s; new in 1.6.0) and one D column per
# parameter. EXTNAME is the FULL table name (1.5.x used tname[:8], which made
# two pairs of tables collide and their rows multiply on every rewrite); tables
# are read back by their TBLNAME keyword, so 1.5.x files are separated again.
def _fits_column_names(params: Sequence[str]) -> Dict[str, str]:
    """param -> unique FITS column name (printable ASCII, <= 68 characters)."""
    names: Dict[str, str] = {}
    used = {"TIMESTAMP_LOCAL", "EPOCH_UTC"}
    for param in params:
        base = _fits_ascii(param)[:68] or "PARAM"
        name, counter = base, 2
        while name in used:
            name = f"{base[:63]}_{counter}"
            counter += 1
        used.add(name)
        names[param] = name
    return names


def _fits_table_from_batch(batch: Dict[str, Any]) -> Dict[str, Any]:
    """Batch -> {"ts": [...], "cols": {name: float64 array}, "units": {name: unit}}."""
    names = _fits_column_names(batch["params"])
    columns = {"EPOCH_UTC": np.asarray(batch["epoch"], dtype=np.float64)}
    units = {"EPOCH_UTC": "s"}
    for param in batch["params"]:
        name = names[param]
        columns[name] = np.array([np.nan if v is None else float(v) for v in batch["columns"][param]],
                                 dtype=np.float64)
        units[name] = batch["units"].get(param, "") or "dimensionless"
    return {"ts": list(batch["ts"]), "cols": columns, "units": units}


def _fits_concat(target: Dict[str, Any], extra: Dict[str, Any], unique: bool = False) -> None:
    """Append ``extra`` rows to ``target`` (NaN-padding columns either side lacks).

    unique=True drops rows whose timestamp ``target`` already has (used only
    when merging rescue files, which may overlap the canonical file).
    """
    rows = range(len(extra["ts"]))
    if unique:
        seen = set(target["ts"])
        rows = [i for i in rows if extra["ts"][i] not in seen]
        if not rows:
            return
    n_old, n_new = len(target["ts"]), len(rows)
    for name in list(target["cols"]) + [c for c in extra["cols"] if c not in target["cols"]]:
        old = target["cols"].get(name)
        if old is None:
            old = np.full(n_old, np.nan)
            target["units"][name] = extra["units"].get(name, "")
        new = extra["cols"].get(name)
        new = np.full(n_new, np.nan) if new is None else np.asarray(new, dtype=np.float64)[list(rows)]
        target["cols"][name] = np.concatenate([old, new])
    target["ts"].extend(extra["ts"][i] for i in rows)


def _fits_read(astrofits: Any, path: str) -> Dict[str, Dict[str, Any]]:
    """Read every table HDU of ``path`` into numpy arrays, keyed by TBLNAME."""
    tables: Dict[str, Dict[str, Any]] = {}
    with astrofits.open(path, memmap=False) as hdul:
        for hdu in hdul[1:]:
            data = getattr(hdu, "data", None)
            if data is None or len(data) == 0 or "TIMESTAMP_LOCAL" not in data.names:
                continue
            tname = str(hdu.header.get("TBLNAME") or hdu.header.get("EXTNAME") or "").strip()
            if not tname:
                continue
            table = {"ts": [str(v).strip() for v in data["TIMESTAMP_LOCAL"]], "cols": {}, "units": {}}
            for number, name in enumerate(data.names, start=1):
                if name == "TIMESTAMP_LOCAL":
                    continue
                table["cols"][name] = np.array(data[name], dtype=np.float64)
                table["units"][name] = str(hdu.header.get(f"TUNIT{number}", "") or "").strip()
            if tname in tables:
                _fits_concat(tables[tname], table)
            else:
                tables[tname] = table
    return tables


def _fits_build(astrofits: Any, ip: str, tables: Dict[str, Dict[str, Any]], filename: str) -> Any:
    """HDUList: primary header plus one BinTableHDU per table (page order)."""
    primary = astrofits.PrimaryHDU()
    header = primary.header
    now = _utc_now()
    header["TELESCOP"] = (_fits_ascii("GBT"), _fits_ascii("Green Bank Telescope facility"))
    header["INSTRUME"] = (_fits_ascii("ABPowerMeter"), _fits_ascii("Allen-Bradley 1403 Site Power Meter"))
    header["ORIGIN"] = (_fits_ascii("NRAO-GBO"), _fits_ascii("National Radio Astronomy Observatory"))
    header["OBSERVER"] = (_fits_ascii("WWallace"), _fits_ascii("W. Wallace"))
    header["DATE-WRT"] = (_fits_ascii(_utc_iso_z(now)), _fits_ascii("UTC time of this write"))
    header["FILENAME"] = (_fits_ascii(filename), _fits_ascii("FITS file name"))
    header["DEVIP"] = (_fits_ascii(ip), _fits_ascii("Source device IP address"))
    header["SWVERS"] = (_fits_ascii(__version__), _fits_ascii("ab_power_meter_monitor version"))
    header["COMMENT"] = _fits_ascii("Allen-Bradley power meter telemetry - NRAO GBO site infrastructure",
                                    comment=True)
    header["HISTORY"] = _fits_ascii(f"Written by ab_power_meter_monitor.py {__version__} on {now.date()}",
                                    comment=True)
    stamps = [t for table in tables.values() for t in (table["ts"][:1] + table["ts"][-1:])]
    if stamps:
        header["DATE-OBS"] = (_fits_ascii(min(stamps).replace(" ", "T")), _fits_ascii("Local time of first sample"))
        header["DATE-END"] = (_fits_ascii(max(stamps).replace(" ", "T")), _fits_ascii("Local time of last sample"))
        header["NSAMP"] = (max(len(t["ts"]) for t in tables.values()), _fits_ascii("Max rows over all tables"))
    hdus = [primary]
    for tname in sorted(tables, key=_table_order):
        table = tables[tname]
        if not table["ts"]:
            continue
        order = (["EPOCH_UTC"] if "EPOCH_UTC" in table["cols"] else []) + \
            [c for c in table["cols"] if c != "EPOCH_UTC"]
        columns = [astrofits.Column(name="TIMESTAMP_LOCAL", format="A19", unit="local time",
                                    array=np.array(table["ts"], dtype="U19"))]
        for name in order:
            columns.append(astrofits.Column(
                name=name, format="D",
                unit=_fits_ascii(table["units"].get(name) or "dimensionless"),
                array=np.asarray(table["cols"][name], dtype=np.float64)))
        hdu = astrofits.BinTableHDU.from_columns(columns)
        hdu.header["EXTNAME"] = _fits_ascii(tname)
        hdu.header["TBLNAME"] = _fits_ascii(tname)
        hdu.header["SRCIP"] = _fits_ascii(ip)
        hdu.header["NSAMP"] = (len(table["ts"]), _fits_ascii("Number of rows (poll cycles)"))
        hdu.header["DATE-OBS"] = _fits_ascii(table["ts"][0].replace(" ", "T"))
        hdu.header["DATE-END"] = _fits_ascii(table["ts"][-1].replace(" ", "T"))
        hdu.header["COMMENT"] = _fits_ascii(f"AB meter table: {tname}", comment=True)
        hdu.header["COMMENT"] = _fits_ascii(
            f"{len(table['ts'])} rows; TIMESTAMP_LOCAL local wall clock, EPOCH_UTC Unix s", comment=True)
        hdus.append(hdu)
    return astrofits.HDUList(hdus)


def _fits_write_ip(astrofits: Any, ip: str, ip_batches: Dict, out_dir: str, cfg: Dict[str, Any],
                   must_persist: bool) -> Dict[Tuple[str, str], int]:
    """Merge new rows into the device FITS file and replace it atomically."""
    canonical = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}.fits")
    if _is_fresh_pending(canonical, cfg):
        _remove_quietly(canonical)
        _mark_fresh_done(canonical)
    readable, disk = True, {}
    if os.path.exists(canonical):
        try:
            disk = _fits_read(astrofits, canonical)
        except Exception as exc:
            readable = False
            logger.warning("FITS: cannot read %s (%s) - new rows go to a rescue file", canonical, exc)
    new_tables = {b["tname"]: _fits_table_from_batch(b) for b in ip_batches.values()}
    max_rows = int(cfg.get("fits_max_rows_per_file", FITS_MAX_ROWS_PER_FILE))
    if readable and disk and max_rows > 0:
        have = max(len(t["ts"]) for t in disk.values())
        adding = max(len(t["ts"]) for t in new_tables.values())
        if have + adding > max_rows and _rollover(canonical, f"more than {max_rows} rows"):
            disk = {}
    rescued: List[str] = []
    if readable:
        pattern = glob.escape(os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_")) + "*_rescue.fits"
        for rescue_path in sorted(glob.glob(pattern)):
            try:
                for tname, table in _fits_read(astrofits, rescue_path).items():
                    if tname in disk:
                        _fits_concat(disk[tname], table, unique=True)
                    else:
                        disk[tname] = table
                rescued.append(rescue_path)
            except Exception as exc:
                logger.warning("FITS: rescue file %s unreadable (%s) - left on disk", rescue_path, exc)
        for tname, table in new_tables.items():
            if tname in disk:
                _fits_concat(disk[tname], table)
            else:
                disk[tname] = table
        target = canonical
    else:
        disk = new_tables
        target = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_{_stamp()}_rescue.fits")
    temp = _temp_sibling(target)
    try:
        _fits_build(astrofits, ip, disk, os.path.basename(target)).writeto(temp, overwrite=True)
        try:
            os.replace(temp, target)
        except OSError as exc:               # canonical held open by another program
            if not must_persist:
                logger.warning("FITS: %s is in use (%s); retrying later", target, exc)
                return {}
            target = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_{_stamp()}_rescue.fits")
            _fits_build(astrofits, ip, new_tables, os.path.basename(target)).writeto(target, overwrite=True)
            rescued = []
    finally:
        _remove_quietly(temp)
    for rescue_path in rescued:
        _remove_quietly(rescue_path)
        logger.info("FITS: rescue file merged and removed: %s", rescue_path)
    logger.info("FITS written (%d table HDUs, %s): %s", len(disk), ip, target)
    return {key: batch["seq"][-1] for key, batch in ip_batches.items()}


def _write_fits_batches(batches: Dict, out_dir: str, cfg: Dict[str, Any],
                        must_persist: bool) -> Dict[Tuple[str, str], int]:
    """FITS writer: returns {(ip, table): last seq persisted}."""
    try:
        from astropy.io import fits as astrofits
    except ImportError as exc:
        logger.error("astropy not available - FITS output skipped: %s", exc)
        return {}
    if np is None:
        logger.error("numpy not available - FITS output skipped")
        return {}
    os.makedirs(out_dir, exist_ok=True)
    done: Dict[Tuple[str, str], int] = {}
    for ip, ip_batches in _group_by_ip(batches).items():
        try:
            done.update(_fits_write_ip(astrofits, ip, ip_batches, out_dir, cfg, must_persist))
        except Exception as exc:
            logger.error("FITS write failed for %s: %s\n%s", ip, exc, traceback.format_exc())
    return done


# %% ===========================================================================
#    OUTPUT MODULE: EXCEL XLSX  (one workbook per device; streamed rewrite)
# =============================================================================
# Layout (unchanged): one sheet per table; row 1 header, row 2 units, rows 3+
# data; a line chart of up to 12 parameters. The old workbook is read with a
# read-only workbook and copied row by row into a write-only workbook, so the
# memory use does not grow with the file [R7] (1.5.x held every cell in RAM).
def _xlsx_sheet_title(tname: str) -> str:
    """Excel sheet title: no []:*?/\\ characters, at most 31 characters."""
    return re.sub(r"[\[\]:*?/\\]", "_", tname)[:31]


def _xlsx_read_small(path: str) -> Dict[str, Dict[str, Any]]:
    """{sheet: {"params", "units", "rows"}} of a (small) rescue workbook."""
    import openpyxl
    out: Dict[str, Dict[str, Any]] = {}
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header, units = next(rows, None), next(rows, None)
            if not header:
                continue
            params = ["" if h is None else str(h) for h in header[1:]]
            out[sheet.title] = {"params": params,
                                "units": ["" if u is None else str(u) for u in (units or ())[1:]],
                                "rows": [tuple(r) for r in rows if r and r[0] is not None]}
    finally:
        book.close()
    return out


def _xlsx_write_ip(ip: str, ip_batches: Dict, out_dir: str, cfg: Dict[str, Any],
                   must_persist: bool) -> Dict[Tuple[str, str], int]:
    """Stream the old workbook plus new rows into a new workbook; replace atomically."""
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.data_source import StrRef
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    canonical = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}.xlsx")
    if _is_fresh_pending(canonical, cfg):
        _remove_quietly(canonical)
        _mark_fresh_done(canonical)
    excel_rows = 1048576 - 2                       # header + units rows [R6]
    max_rows = int(cfg.get("xlsx_max_rows_per_file", XLSX_MAX_ROWS_PER_FILE))
    max_rows = min(max_rows, excel_rows) if max_rows > 0 else excel_rows
    new_by_sheet = {_xlsx_sheet_title(b["tname"]): b for b in ip_batches.values()}
    source, readable = None, True
    if os.path.exists(canonical):
        try:
            source = openpyxl.load_workbook(canonical, read_only=True, data_only=True)
        except Exception as exc:
            readable = False
            logger.warning("XLSX: cannot read %s (%s) - new rows go to a rescue file", canonical, exc)
    if source is not None:
        current = max([max(0, (ws.max_row or 0) - 2) for ws in source.worksheets] or [0])
        adding = max(len(b["seq"]) for b in ip_batches.values())
        if current + adding > max_rows:
            source.close()
            source = None
            if not _rollover(canonical, f"more than {max_rows} rows"):
                if not must_persist:
                    return {}
                readable = False               # cannot roll: rows go to a rescue file
    rescue_sheets: Dict[str, Dict[str, Any]] = {}
    rescued: List[str] = []
    if readable:
        pattern = glob.escape(os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_")) + "*_rescue.xlsx"
        for rescue_path in sorted(glob.glob(pattern)):
            try:
                for title, data in _xlsx_read_small(rescue_path).items():
                    if title in rescue_sheets:
                        rescue_sheets[title]["rows"].extend(data["rows"])
                    else:
                        rescue_sheets[title] = data
                rescued.append(rescue_path)
            except Exception as exc:
                logger.warning("XLSX: rescue file %s unreadable (%s) - left on disk", rescue_path, exc)
    target = canonical if readable else os.path.join(
        out_dir, f"ABMeter_{_safe_ip(ip)}_{_stamp()}_rescue.xlsx")

    head_font = Font(name="Calibri", bold=True, color="FFFFFF")
    head_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
    unit_font = Font(name="Calibri", italic=True, color="FFFFFF")
    unit_fill = PatternFill(fill_type="solid", fgColor="2E75B6")
    centre = Alignment(horizontal="center")

    book = openpyxl.Workbook(write_only=True)
    titles = list(source.sheetnames) if source is not None else []
    for title in list(rescue_sheets) + list(new_by_sheet):
        if title not in titles:
            titles.append(title)
    title_to_table = {_xlsx_sheet_title(t): t for t in TABLE_NAMES.values()}
    titles.sort(key=lambda t: _table_order(title_to_table.get(t, t)))
    try:
        for title in titles:
            old_rows: Iterable[Any] = ()
            params: List[str] = []
            units: Dict[str, str] = {}
            width_old = 0
            if source is not None and title in source.sheetnames:
                iterator = source[title].iter_rows(values_only=True)
                header, unit_row = next(iterator, None), next(iterator, None)
                if header:
                    params = ["" if h is None else str(h) for h in header[1:]]
                    while params and params[-1] == "":
                        params.pop()
                    width_old = len(params)
                    for j, p in enumerate(params):
                        cell = unit_row[j + 1] if unit_row and len(unit_row) > j + 1 else None
                        units[p] = "" if cell is None else str(cell)
                    old_rows = iterator
            rescue = rescue_sheets.get(title)
            batch = new_by_sheet.get(title)
            for extra, extra_units in (((rescue or {}).get("params", []), dict(zip(
                    (rescue or {}).get("params", []), (rescue or {}).get("units", [])))),
                    ((batch or {}).get("params", []), (batch or {}).get("units", {}))):
                for p in extra:
                    if p and p not in params:
                        params.append(p)
                        units[p] = extra_units.get(p, "")
            sheet = book.create_sheet(title=title)
            sheet.freeze_panes = "A3"           # must precede the first row [R7]
            for col, label in enumerate(["Timestamp_Local"] + params, start=1):
                sheet.column_dimensions[get_column_letter(col)].width = min(max(len(label), 10) + 4, 40)
            cells = []
            for label in ["Timestamp_Local"] + params:
                cell = WriteOnlyCell(sheet, value=label)
                cell.font, cell.fill, cell.alignment = head_font, head_fill, centre
                cells.append(cell)
            sheet.append(cells)
            cells = []
            for label in ["(units)"] + [units.get(p, "") for p in params]:
                cell = WriteOnlyCell(sheet, value=label)
                cell.font, cell.fill, cell.alignment = unit_font, unit_fill, centre
                cells.append(cell)
            sheet.append(cells)
            count = 0
            seen: Optional[Set[str]] = set() if rescue else None
            for row in old_rows:
                if not row or row[0] is None:
                    continue
                values = list(row[1:1 + width_old])
                values += [None] * (len(params) - len(values))
                sheet.append([row[0]] + values)
                count += 1
                if seen is not None:
                    seen.add(str(row[0]))
            if rescue:
                index = {p: j for j, p in enumerate(rescue["params"])}
                for row in rescue["rows"]:
                    if str(row[0]) in seen:
                        continue
                    sheet.append([row[0]] + [row[1 + index[p]] if p in index and len(row) > 1 + index[p]
                                             else None for p in params])
                    count += 1
            if batch:
                columns = [batch["columns"].get(p) for p in params]
                for i, stamp in enumerate(batch["ts"]):
                    sheet.append([stamp] + [c[i] if c is not None else None for c in columns])
                    count += 1
            if count >= 2 and params:
                chart = LineChart()
                chart.title = f"{title} - {ip}"
                chart.style = 10
                chart.y_axis.title = "Value"
                chart.x_axis.title = "Sample (poll cycle)"
                chart.width, chart.height = 24, 14
                for col in range(2, 2 + min(12, len(params))):
                    chart.add_data(Reference(sheet, min_col=col, max_col=col, min_row=3,
                                             max_row=count + 2), titles_from_data=False)
                    chart.series[-1].tx = SeriesLabel(
                        strRef=StrRef(f=f"'{title}'!${get_column_letter(col)}$1"))
                chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=count + 2))
                sheet.add_chart(chart, f"A{count + 5}")
    finally:
        if source is not None:
            source.close()
    temp = _temp_sibling(target)
    try:
        book.save(temp)
        try:
            os.replace(temp, target)
        except OSError as exc:                 # workbook open in Excel (Windows)
            if not must_persist:
                logger.warning("XLSX: %s is in use (%s); retrying later", target, exc)
                return {}
            rescue_target = os.path.join(out_dir, f"ABMeter_{_safe_ip(ip)}_{_stamp()}_rescue.xlsx")
            os.replace(temp, rescue_target)
            logger.warning("XLSX: %s in use - wrote %s instead", target, rescue_target)
            rescued = []
    finally:
        _remove_quietly(temp)
    for rescue_path in rescued:
        _remove_quietly(rescue_path)
        logger.info("XLSX: rescue file merged and removed: %s", rescue_path)
    logger.info("XLSX written (%d sheet(s), %s): %s", len(titles), ip, target)
    return {key: batch["seq"][-1] for key, batch in ip_batches.items()}


def _write_xlsx_batches(batches: Dict, out_dir: str, cfg: Dict[str, Any],
                        must_persist: bool) -> Dict[Tuple[str, str], int]:
    """XLSX writer: returns {(ip, table): last seq persisted}."""
    try:
        import openpyxl  # noqa: F401  (availability check)
    except ImportError as exc:
        logger.error("openpyxl not available - XLSX output skipped: %s", exc)
        return {}
    os.makedirs(out_dir, exist_ok=True)
    done: Dict[Tuple[str, str], int] = {}
    for ip, ip_batches in _group_by_ip(batches).items():
        try:
            done.update(_xlsx_write_ip(ip, ip_batches, out_dir, cfg, must_persist))
        except Exception as exc:
            logger.error("XLSX write failed for %s: %s\n%s", ip, exc, traceback.format_exc())
    return done


# %% ===========================================================================
#    PUBLIC WRITER FUNCTIONS (1.5.x names and signatures, now incremental)
# =============================================================================
# Each writes the rows that output has not persisted yet (to that folder) and
# waits for the result. ``all_device_data`` is accepted for compatibility; the
# data always comes from TIME_SERIES_STORE.
_WRITER_SWITCH = {"csv": ("enable_csv", "csv_dir"), "md": ("enable_log_append", "log_dir"),
                  "vszh5": ("enable_veusz", "veusz_dir"), "fits": ("enable_fits", "fits_dir"),
                  "xlsx": ("enable_xlsx", "xlsx_dir")}


def _persist_single(writer: str, out_dir: str, append: bool) -> Dict[str, str]:
    """Run exactly one output writer into ``out_dir`` and wait for it."""
    cfg = default_config()
    for name, (flag, _key) in _WRITER_SWITCH.items():
        cfg[flag] = 1 if name == writer else 0
    cfg[_WRITER_SWITCH[writer][1]] = out_dir
    cfg["append_files"] = int(bool(append))
    return persist_outputs(cfg, full=True, must_persist=True, reason=f"write_{writer}", only=(writer,))


def write_csv(all_device_data: Any, csv_dir: str, append: bool = True) -> None:
    """Append the not-yet-written rows to the per-table CSV files in ``csv_dir``."""
    _persist_single("csv", csv_dir, append)


def write_log_text(all_device_data: Any, log_dir: str, append: bool = True) -> None:
    """Append the not-yet-written rows to the per-table Markdown logs in ``log_dir``."""
    _persist_single("md", log_dir, append)


def write_fits(all_device_data: Any, fits_dir: str, append: bool = True) -> None:
    """Merge the not-yet-written rows into the per-device FITS files in ``fits_dir``."""
    _persist_single("fits", fits_dir, append)


def write_xlsx(all_device_data: Any, xlsx_dir: str, append: bool = True) -> None:
    """Merge the not-yet-written rows into the per-device workbooks in ``xlsx_dir``."""
    _persist_single("xlsx", xlsx_dir, append)


def write_veusz(all_device_data: Any, veusz_dir: str, show_window: bool = False,
                append: bool = True, timestamp_suffix: Optional[str] = None,
                ts_snapshot: Optional[Dict[str, Any]] = None) -> List[str]:
    """Write Veusz .vszh5 documents natively (vszh5_writer; no Veusz needed).

    Without ``timestamp_suffix`` / ``ts_snapshot`` the not-yet-written rows are
    appended to the live per-device documents ABMeter_<ip>.vszh5. With either
    one, a standalone snapshot document ABMeter_<ip>_<suffix>.vszh5 is written
    from ``ts_snapshot`` (TIME_SERIES_STORE layout) or the current store.
    ``show_window`` opens the result(s) in the Veusz program (if installed).
    Returns the document paths.
    """
    paths: List[str] = []
    if timestamp_suffix is not None or ts_snapshot is not None:
        vw = _vszh5_module()
        if vw is None or np is None:
            return paths
        os.makedirs(veusz_dir, exist_ok=True)
        source = ts_snapshot if ts_snapshot is not None else store_snapshot()
        suffix = timestamp_suffix or _stamp()
        cfg = default_config()
        cfg["append_files"] = 1
        for ip, tables in source.items():
            path = os.path.join(veusz_dir, f"ABMeter_{_safe_ip(ip)}_{suffix}.vszh5")
            batches = _batches_from_snapshot(ip, tables)
            if batches:
                _vszh5_append_ip(vw, ip, batches, veusz_dir, cfg, True, path=path)
                paths.append(path)
    else:
        _persist_single("vszh5", veusz_dir, append)
        with _TS_LOCK:
            devices = list(TIME_SERIES_STORE)
        for ip in dedupe_ips(list(devices) + list(all_device_data or {})):
            path = os.path.join(veusz_dir, f"ABMeter_{_safe_ip(ip)}.vszh5")
            if os.path.exists(path):
                paths.append(path)
    if show_window and paths:
        open_in_veusz(paths)
    return paths


def open_in_veusz(paths: Sequence[str]) -> bool:
    """Open documents in the Veusz program (VEUSZ_EXE / ABMETER_VEUSZ / PATH)."""
    vw = _vszh5_module()
    exe = vw.find_veusz_exe(env_vars=("ABMETER_VEUSZ", "VEUSZ_EXE", "QUICKSYN_VEUSZ")) if vw else ""
    if not exe:
        logger.warning("Veusz program not found - set VEUSZ_EXE to veusz(.exe) to open %s",
                       ", ".join(os.path.basename(p) for p in paths))
        return False
    detach: Dict[str, Any] = ({"creationflags": 0x00000008 | 0x00000200}   # DETACHED | NEW_GROUP
                              if os.name == "nt" else {"start_new_session": True})
    for path in paths:
        try:
            subprocess.Popen([exe, str(path)], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                             stderr=subprocess.DEVNULL, close_fds=True, **detach)
            logger.info("Opened in Veusz: %s", path)
        except OSError as exc:
            logger.error("Could not start Veusz (%s): %s", exe, exc)
            return False
    return True


def open_veusz_preview(all_device_data: Any, veusz_dir: str) -> List[str]:
    """GUI "Open in Veusz": bring the documents up to date and open them."""
    return write_veusz(all_device_data, veusz_dir, show_window=True)


def _any_output_enabled(cfg: Dict[str, Any]) -> bool:
    """True when any file output (including ab_monitor.log) is enabled."""
    return any(cfg.get(k) for k in ("enable_fits", "enable_csv", "enable_xlsx",
                                    "enable_log_append", "enable_log_file", "enable_veusz"))


def _any_data_output(cfg: Dict[str, Any]) -> bool:
    """True when any output that persists the samples is enabled."""
    return any(cfg.get(k) for k in ("enable_fits", "enable_csv", "enable_xlsx",
                                    "enable_log_append", "enable_veusz"))


def flush_outputs_parallel(cfg: Dict[str, Any]) -> "concurrent.futures.Future":
    """1.5.x API: persist every enabled output in the background; returns a Future."""
    return _flush_manager().submit(persist_outputs, cfg, True, False, "flush_outputs_parallel")


def finalize_outputs(cfg: Optional[Dict[str, Any]] = None) -> Dict[str, str]:
    """Write everything not yet on disk to every enabled output and wait for it."""
    cfg = default_config() if cfg is None else cfg
    if not _any_data_output(cfg):
        return {}
    return persist_outputs(cfg, full=True, must_persist=True, reason="final")


# %% ===========================================================================
#    RAM FLUSH
# =============================================================================
def _rate_limited_warning(text: str, every_sec: float = 600.0) -> None:
    """logger.warning(text) at most once every ``every_sec`` seconds."""
    now = time.monotonic()
    last = _FLUSH_STATE.get("last_warn_mono")
    if last is None or now - last >= every_sec:
        _FLUSH_STATE["last_warn_mono"] = now
        logger.warning(text)


def ram_flush_reason(cfg: Optional[Dict[str, Any]] = None) -> Optional[str]:
    """Why an automatic RAM flush is due now, or None (see MEMORY AND PERSISTENCE).

    Honours ENABLE_RAM_FLUSH (module switch, read now) AND cfg["enable_ram_flush"],
    and MEM_FLUSH_MIN_INTERVAL_SEC since the previous flush.
    """
    cfg = default_config() if cfg is None else cfg
    if not (bool(ENABLE_RAM_FLUSH) and bool(cfg.get("enable_ram_flush", 1))):
        return None
    last = _FLUSH_STATE.get("last_mono")
    interval = float(cfg.get("mem_flush_min_interval_sec", MEM_FLUSH_MIN_INTERVAL_SEC))
    if last is not None and time.monotonic() - last < interval:
        return None
    store_mb = ts_store_size_estimate_mb()
    limit = float(cfg.get("mem_flush_threshold_mb", MEM_FLUSH_THRESHOLD_MB))
    if limit > 0 and store_mb > limit:
        return f"store {store_mb:.1f} MB > {limit:.0f} MB"
    big_enough = store_mb >= float(cfg.get("mem_min_store_mb_for_clear", MEM_MIN_STORE_MB_FOR_CLEAR))
    rss_limit = float(cfg.get("mem_proc_rss_limit_mb", MEM_PROC_RSS_LIMIT_MB))
    if rss_limit > 0:
        rss = process_rss_mb()
        if rss > rss_limit:
            if big_enough:
                return f"process RSS {rss:.0f} MB > {rss_limit:.0f} MB"
            _rate_limited_warning(f"Process RSS {rss:.0f} MB exceeds {rss_limit:.0f} MB but the "
                                  f"store holds only {store_mb:.2f} MB - a flush would not help.")
    pct_limit = float(cfg.get("mem_ram_pct_limit", MEM_RAM_PCT_LIMIT))
    if pct_limit > 0:
        pct = system_ram_used_pct()
        if pct >= pct_limit:
            if big_enough:
                return f"system RAM {pct:.1f} % >= {pct_limit:.0f} %"
            _rate_limited_warning(f"System RAM {pct:.1f} % >= {pct_limit:.0f} %, but this program's "
                                  f"store holds only {store_mb:.2f} MB - other processes use the RAM.")
    free_min = float(cfg.get("mem_free_min_mb", MEM_FREE_MIN_MB))
    if free_min > 0:
        free = system_free_ram_mb()
        if free < free_min:
            if big_enough:
                return f"free RAM {free:.0f} MB < {free_min:.0f} MB"
            _rate_limited_warning(f"Free RAM {free:.0f} MB < {free_min:.0f} MB, but the store holds "
                                  f"only {store_mb:.2f} MB.")
    return None


def should_flush(cfg: Dict[str, Any]) -> bool:
    """1.5.x API: True when ram_flush_reason() names a reason."""
    return ram_flush_reason(cfg) is not None


def _spill_store(cfg: Dict[str, Any]) -> Optional[str]:
    """Pickle the whole store (with epochs) to SPILL_DIR atomically; path or None."""
    folder = cfg.get("spill_dir", SPILL_DIR)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"ABMeter_spill_{_stamp()}.pkl")
    temp = _temp_sibling(path)
    try:
        with open(temp, "wb") as handle:
            pickle.dump({"format": 1, "monitor_version": __version__, "created_utc": _utc_iso_z(),
                         "store": store_snapshot()}, handle, protocol=pickle.HIGHEST_PROTOCOL)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp, path)
        logger.info("[RAM Flush] store spilled to %s", path)
        return path
    except Exception as exc:
        logger.error("[RAM Flush] spill to %s failed: %s", path, exc)
        return None
    finally:
        _remove_quietly(temp)


def load_spill(path: str) -> Dict[str, Any]:
    """Load a RAM_FLUSH_SPILL_PICKLE file: {"store": {ip: {table: ...}}, ...}.

    Only load files this program wrote - unpickling runs code from the file.
    """
    with open(path, "rb") as handle:
        return pickle.load(handle)


def perform_ram_flush(cfg: Optional[Dict[str, Any]] = None, reason: str = "manual",
                      keep_samples: Optional[int] = None) -> Dict[str, Any]:
    """Persist everything, then drop the rows every enabled output has persisted.

    Order of operations
    -------------------
    1. persist_outputs(full=True, must_persist=True) - waits for every write.
    2. Per table, drop the oldest rows whose sequence number every enabled
       output has confirmed, keeping the newest ``keep_samples`` rows
       (default cfg["ram_flush_keep_samples"]). Rows an output failed to
       write are kept and retried. Without data outputs the rows are spilled
       to a pickle (RAM_FLUSH_SPILL_PICKLE) or dropped.
    3. Tables left empty are removed (store == {} after a full flush, as 1.5.x).
    4. gc.collect() and, with MEM_RETURN_TO_OS, return_memory_to_os():
       Linux malloc_trim; Windows heap decommit (HeapOptimizeResources) and
       (with MEM_TRIM_WORKING_SET) a working-set trim.

    Returns
    -------
    dict
        reason, rows_removed, rows_kept_unpersisted, store_mb_before/after,
        rss_mb_before/after, writers (persist status), spill (path or None),
        os_release (return_memory_to_os() report), os_release_text.
    """
    global _VEUSZ_FLUSH_COUNT
    cfg = default_config() if cfg is None else cfg
    keep = int(cfg.get("ram_flush_keep_samples", RAM_FLUSH_KEEP_SAMPLES)
               if keep_samples is None else keep_samples)
    store_before, rss_before = ts_store_size_estimate_mb(), process_rss_mb()
    removed = unpersisted = 0
    spill = None
    with _PERSIST_LOCK:
        specs = _writer_specs(cfg, full=True)
        writer_status = persist_outputs(cfg, full=True, must_persist=True, reason="ram_flush")
        unlimited = False
        if not specs:
            if bool(cfg.get("ram_flush_spill_pickle", RAM_FLUSH_SPILL_PICKLE)):
                spill = _spill_store(cfg)
                unlimited = spill is not None
            else:
                unlimited = True
                logger.warning("[RAM Flush] no file output enabled - in-memory samples are discarded "
                               "(enable an output or RAM_FLUSH_SPILL_PICKLE to keep them).")
        with _CURSOR_LOCK:
            maps = [dict(_CURSORS.get(_cursor_key(n, d), {})) for n, d, _ in specs]
        with _TS_LOCK:
            for ip in list(TIME_SERIES_STORE):
                tables = TIME_SERIES_STORE[ip]
                for tname in list(tables):
                    meta = _aligned_meta(ip, tname, tables[tname])
                    n_rows = len(meta["seq"])
                    if unlimited:
                        safe_rows = n_rows
                    elif maps:
                        safe = min(m.get((ip, tname), 0) for m in maps)
                        safe_rows = bisect.bisect_right(meta["seq"], safe)
                    else:
                        safe_rows = 0              # spill failed: keep everything
                    unpersisted += n_rows - safe_rows
                    removed += _drop_head_rows(ip, tname, min(safe_rows, max(0, n_rows - keep)))
                    if not tables[tname]["timestamps_local"]:
                        del tables[tname]
                        _TS_META.get(ip, {}).pop(tname, None)
                if not tables:
                    del TIME_SERIES_STORE[ip]
                    _TS_META.pop(ip, None)
        _FLUSH_STATE["count"] += 1
        _FLUSH_STATE["last_mono"] = time.monotonic()
        _FLUSH_STATE["last_reason"] = reason
        _VEUSZ_FLUSH_COUNT += 1
    # 4. Give the freed memory back to the OS (outside _PERSIST_LOCK: the
    #    working-set trim may take a moment and must not delay the next poll's
    #    writes). gc.collect() always runs; the OS calls only with MEM_RETURN_TO_OS.
    if bool(cfg.get("mem_return_to_os", MEM_RETURN_TO_OS)):
        os_release = return_memory_to_os(bool(cfg.get("mem_trim_working_set", MEM_TRIM_WORKING_SET)))
    else:
        gc.collect()
        os_release = {}
    result = {"reason": reason, "rows_removed": removed, "rows_kept_unpersisted": unpersisted,
              "store_mb_before": store_before, "store_mb_after": ts_store_size_estimate_mb(),
              "rss_mb_before": rss_before, "rss_mb_after": process_rss_mb(),
              "writers": writer_status, "spill": spill, "count": _FLUSH_STATE["count"],
              "os_release": os_release, "os_release_text": describe_memory_release(os_release)}
    logger.warning("[RAM Flush] %s: removed %d row(s), kept %d unpersisted; store %.2f -> %.2f MB; "
                   "RSS %.0f -> %.0f MB; %s", reason, removed, unpersisted, result["store_mb_before"],
                   result["store_mb_after"], result["rss_mb_before"], result["rss_mb_after"],
                   result["os_release_text"])
    if unpersisted:
        logger.error("[RAM Flush] %d row(s) could not be persisted and stay in memory - see the "
                     "writer errors above.", unpersisted)
    return result


# %% ===========================================================================
#    RUN LIFECYCLE: default_config, begin_run, poll_cycle, end_run
# =============================================================================
def _follow_base(current: str, name: str) -> str:
    """A sub-folder that follows a changed OUTPUT_BASE_DIR unless set explicitly."""
    if os.path.normcase(current) == os.path.normcase(os.path.join(_IMPORT_BASE_DIR, name)):
        return os.path.join(OUTPUT_BASE_DIR, name)
    return current


def stop_signal_file() -> str:
    """Path of the STOP_COLLECTION file (follows a changed OUTPUT_BASE_DIR)."""
    return _follow_base(STOP_SIGNAL_FILE, "STOP_COLLECTION")


def default_config() -> Dict[str, Any]:
    """Run configuration built from the module switches at call time."""
    return {
        "ip_list": dedupe_ips(IP_LIST),
        "sample_period": float(SAMPLE_PERIOD_SEC),
        "headless_loop_count": HEADLESS_LOOP_COUNT,
        "headless_console_dicts_only": HEADLESS_CONSOLE_DICTS_ONLY,
        "headless_print_each_sample": HEADLESS_PRINT_EACH_SAMPLE,
        "headless_print_cumulative": HEADLESS_PRINT_CUMULATIVE,
        "headless_silent": HEADLESS_SILENT,
        "headless_max_consec_fails": HEADLESS_MAX_CONSEC_FAILS,
        "enable_gui": ENABLE_GUI,
        "enable_fits": ENABLE_FITS,
        "enable_csv": ENABLE_CSV,
        "enable_xlsx": ENABLE_XLSX,
        "enable_log_append": ENABLE_LOG_APPEND,
        "enable_log_file": ENABLE_LOG_FILE,
        "enable_veusz": ENABLE_VEUSZ,
        "veusz_append_every_cycle": VEUSZ_APPEND_EVERY_CYCLE,
        "veusz_write_on_flush": VEUSZ_WRITE_ON_FLUSH,          # deprecated, no effect
        "enable_ram_flush": ENABLE_RAM_FLUSH,
        "mem_ram_pct_limit": MEM_RAM_PCT_LIMIT,
        "mem_flush_threshold_mb": MEM_FLUSH_THRESHOLD_MB,
        "mem_free_min_mb": MEM_FREE_MIN_MB,
        "mem_proc_rss_limit_mb": MEM_PROC_RSS_LIMIT_MB,
        "mem_flush_min_interval_sec": MEM_FLUSH_MIN_INTERVAL_SEC,
        "mem_min_store_mb_for_clear": MEM_MIN_STORE_MB_FOR_CLEAR,
        "ram_flush_keep_samples": RAM_FLUSH_KEEP_SAMPLES,
        "ram_flush_spill_pickle": RAM_FLUSH_SPILL_PICKLE,
        "mem_return_to_os": MEM_RETURN_TO_OS,
        "mem_trim_working_set": MEM_TRIM_WORKING_SET,
        "full_write_every_n_cycles": FULL_WRITE_EVERY_N_CYCLES,
        "fsync_outputs": FSYNC_OUTPUTS,
        "fits_max_rows_per_file": FITS_MAX_ROWS_PER_FILE,
        "xlsx_max_rows_per_file": XLSX_MAX_ROWS_PER_FILE,
        "vszh5_max_file_mb": VSZH5_MAX_FILE_MB,
        "vszh5_time_axis": VSZH5_TIME_AXIS,
        "vszh5_panels_per_page": VSZH5_PANELS_PER_PAGE,
        "vszh5_file_locking": VSZH5_FILE_LOCKING,
        "output_base_dir": OUTPUT_BASE_DIR,
        "fits_dir": _follow_base(FITS_DIR, "fits"),
        "csv_dir": _follow_base(CSV_DIR, "csv"),
        "xlsx_dir": _follow_base(XLSX_DIR, "xlsx"),
        "veusz_dir": _follow_base(VEUSZ_DIR, "veusz"),
        "log_dir": _follow_base(LOG_DIR, "logs"),
        "spill_dir": _follow_base(SPILL_DIR, "spill"),
        "http_timeout_sec": HTTP_TIMEOUT_SEC,
        "http_retry_count": HTTP_RETRY_COUNT,
        "http_retry_delay_sec": HTTP_RETRY_DELAY_SEC,
        "append_files": APPEND_OUTPUT_FILES,
    }


def _ensure_output_dirs(cfg: Dict[str, Any]) -> None:
    """Create only the folders of the enabled outputs (none when all are off)."""
    if not _any_output_enabled(cfg):
        return
    os.makedirs(cfg["output_base_dir"], exist_ok=True)
    for flag, key in (("enable_fits", "fits_dir"), ("enable_csv", "csv_dir"),
                      ("enable_xlsx", "xlsx_dir"), ("enable_veusz", "veusz_dir")):
        if cfg.get(flag):
            os.makedirs(cfg[key], exist_ok=True)
    if cfg.get("enable_log_append") or cfg.get("enable_log_file"):
        os.makedirs(cfg["log_dir"], exist_ok=True)


def begin_run(cfg: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Start a run: output folders, log file, fresh-file bookkeeping, counters.

    With append_files = 0 every output file is started afresh once per run.
    Cursors persist across runs (sequence numbers never repeat), so rows
    already written are never written again. Returns the configuration used.
    """
    global _VEUSZ_FLUSH_COUNT
    cfg = default_config() if cfg is None else cfg
    with _PERSIST_LOCK:
        _RUN_STATE["active"] = True
        _RUN_STATE["run_id"] += 1
        _RUN_STATE["cycle"] = 0
        _RUN_STATE["started_utc"] = _utc_iso_z()
        with _FRESH_LOCK:
            _RUN_STATE["fresh"] = set()
        _FLUSH_STATE["count"] = 0
        _FLUSH_STATE["last_mono"] = None
        _VEUSZ_FLUSH_COUNT = 0
    _ensure_output_dirs(cfg)
    if cfg.get("enable_log_file"):
        _setup_logging(cfg["log_dir"], append=bool(cfg.get("append_files", 1)), enable_log_file=True)
    logger.info("Run %d started (monitor %s, Python %s) - devices %s",
                _RUN_STATE["run_id"], __version__, ".".join(map(str, sys.version_info[:3])),
                cfg.get("ip_list"))
    return cfg


def poll_cycle(cfg: Optional[Dict[str, Any]] = None,
               stop_event: Optional[threading.Event] = None) -> Dict[str, Any]:
    """One complete poll: read every meter, store, persist, check memory.

    Starts a run first when none is active. Persists the incremental outputs
    every call and FITS/XLSX every FULL_WRITE_EVERY_N_CYCLES calls; performs an
    automatic RAM flush when ram_flush_reason() says so. Never raises for
    device, file or memory problems (they are logged and reported).

    Returns
    -------
    dict
        cycle, time_utc, ok_ips, failed_ips, all_failed, cancelled,
        poll_seconds, persist (writer status), flush (perform_ram_flush
        result or None), store_mb, rss_mb, ram_pct.
    """
    cfg = default_config() if cfg is None else cfg
    if not _RUN_STATE["active"]:
        begin_run(cfg)
    _RUN_STATE["cycle"] += 1
    cycle = _RUN_STATE["cycle"]
    started = time.monotonic()
    data = poll_all_devices(cfg.get("ip_list", IP_LIST), TABLE_NAMES,
                            retries=cfg.get("http_retry_count"),
                            retry_delay=cfg.get("http_retry_delay_sec"),
                            stop_event=stop_event, timeout=cfg.get("http_timeout_sec"))
    meta = data.pop("__poll_meta__", {})
    if not meta.get("all_failed", True):
        # An all-failed poll is NOT stored (it would add gap rows to the store).
        update_named_dicts(data, ip_list=cfg.get("ip_list", IP_LIST))
    data.clear()                           # parsed pages are not needed any more
    poll_seconds = time.monotonic() - started
    persist: Dict[str, str] = {}
    if _any_data_output(cfg):
        every = int(cfg.get("full_write_every_n_cycles", FULL_WRITE_EVERY_N_CYCLES))
        try:
            persist = persist_outputs(cfg, full=every > 0 and cycle % every == 0, reason="cycle")
        except Exception as exc:
            logger.error("Persisting outputs failed: %s", exc, exc_info=True)
    flush = None
    try:
        reason = ram_flush_reason(cfg)
        if reason:
            flush = perform_ram_flush(cfg, reason)
    except Exception as exc:
        logger.error("RAM flush failed: %s", exc, exc_info=True)
    return {"cycle": cycle, "time_utc": _utc_iso_z(), "ok_ips": meta.get("ok_ips", []),
            "failed_ips": meta.get("failed_ips", []), "all_failed": bool(meta.get("all_failed", True)),
            "cancelled": bool(meta.get("cancelled", False)), "poll_seconds": poll_seconds,
            "persist": persist, "flush": flush, "store_mb": ts_store_size_estimate_mb(),
            "rss_mb": process_rss_mb(), "ram_pct": system_ram_used_pct()}


def end_run(cfg: Optional[Dict[str, Any]] = None) -> Dict[str, str]:
    """Finish a run: write everything not yet on disk (all outputs) and wait."""
    cfg = default_config() if cfg is None else cfg
    status: Dict[str, str] = {}
    try:
        status = finalize_outputs(cfg)
    finally:
        _RUN_STATE["active"] = False
    if status:
        logger.info("Run %d finished - final writes: %s", _RUN_STATE["run_id"], status)
    return status


def run_status() -> Dict[str, Any]:
    """Run/flush counters and memory figures (GUI status line, diagnostics)."""
    with _TS_LOCK:
        rows = sum(len(t["timestamps_local"]) for d in TIME_SERIES_STORE.values() for t in d.values())
        values = sum(len(t["timestamps_local"]) * len(t["columns"])
                     for d in TIME_SERIES_STORE.values() for t in d.values())
    return {"run_id": _RUN_STATE["run_id"], "active": _RUN_STATE["active"],
            "cycle": _RUN_STATE["cycle"], "flushes": _FLUSH_STATE["count"],
            "last_flush_reason": _FLUSH_STATE["last_reason"], "rows": rows, "values": values,
            "store_mb": ts_store_size_estimate_mb(), "rss_mb": process_rss_mb(),
            "commit_mb": process_commit_mb(), "ram_pct": system_ram_used_pct()}


# %% ===========================================================================
#    PREVIEW DATA AND RENDERING (used by the GUI; no pyplot anywhere)
# =============================================================================
#: Tableau-10 line colours (distinguishable for common colour-vision deficiencies).
_PREVIEW_COLOURS = ("#1f77b4", "#d62728", "#2ca02c", "#9467bd", "#ff7f0e",
                    "#17becf", "#8c564b", "#e377c2", "#7f7f7f", "#bcbd22")


def _utc_offsets(epochs: Any) -> Any:
    """Local UTC offset (s) per Unix time, DST-aware; one lookup per distinct hour."""
    values = np.asarray(epochs, dtype=float)
    out = np.zeros_like(values)
    good = np.isfinite(values)
    if good.any():
        hours = np.floor(values[good] / 3600.0)
        unique, inverse = np.unique(hours, return_inverse=True)
        table = np.array([datetime.datetime.fromtimestamp(h * 3600.0 + 1800.0, datetime.timezone.utc)
                          .astimezone().utcoffset().total_seconds() for h in unique])
        out[good] = table[inverse]
    return out


def _minmax_decimate(x: Any, y: Any, max_points: int) -> Tuple[Any, Any]:
    """Reduce a series to about ``max_points`` points keeping every bucket's min and max.

    Unlike plain striding, a single-sample spike (e.g. an inrush current) is
    never dropped from the preview. The newest point is always kept.
    """
    n_points = len(x)
    buckets = max(1, int(max_points) // 2)
    size = int(math.ceil(n_points / float(buckets)))
    if size <= 1:
        return x, y
    n_buckets = int(math.ceil(n_points / float(size)))
    padded = np.concatenate([y, np.full(n_buckets * size - n_points, np.nan)]).reshape(n_buckets, size)
    low = np.where(np.isnan(padded), np.inf, padded).argmin(axis=1)
    high = np.where(np.isnan(padded), -np.inf, padded).argmax(axis=1)
    base = np.arange(n_buckets) * size
    index = np.unique(np.concatenate([base + low, base + high, [n_points - 1]]))
    index = index[index < n_points]
    return x[index], y[index]


def preview_plot_keys(ip_order: Optional[Sequence[str]] = None) -> List[Tuple[str, Optional[str], Optional[str], str]]:
    """(ip, table, overlay group, title) of every preview plot that has data."""
    keys: List[Tuple[str, Optional[str], Optional[str], str]] = []
    with _TS_LOCK:
        order = dedupe_ips([ip for ip in (ip_order or IP_LIST) if ip in TIME_SERIES_STORE]
                           + list(TIME_SERIES_STORE))
        for ip in order:
            tables = TIME_SERIES_STORE.get(ip, {})
            names = sorted((t for t, s in tables.items() if s["columns"] and s["timestamps_local"]),
                           key=_table_order)
            for tname in names:
                keys.append((ip, tname, None, f"{ip} — {tname}"))
            for group, substrings in VEUSZ_OVERLAY_GROUPS.items():
                if any(any(s in p.lower() for s in substrings) for t in names for p in tables[t]["columns"]):
                    keys.append((ip, None, group, f"{ip} — Overlay: {group}"))
    return keys


def preview_series(ip: str, tname: Optional[str] = None, group: Optional[str] = None,
                   max_points: Optional[int] = None) -> List[Tuple[str, Any, Any]]:
    """[(label, unix_times, values)] of one table or one overlay group, decimated.

    Only the requested table(s) are copied, briefly, under the store lock;
    conversion and min/max decimation happen outside it.
    """
    max_points = PREVIEW_MAX_POINTS if max_points is None else max_points
    raw: List[Tuple[str, List[float], List[Any]]] = []
    substrings = VEUSZ_OVERLAY_GROUPS.get(group, []) if group else []
    with _TS_LOCK:
        tables = TIME_SERIES_STORE.get(ip, {})
        for name in ([tname] if tname else sorted(tables, key=_table_order)):
            tstore = tables.get(name)
            if not tstore:
                continue
            epochs = list(_aligned_meta(ip, name, tstore)["epoch"])
            for param, values in tstore["columns"].items():
                if group and not any(s in param.lower() for s in substrings):
                    continue
                unit = tstore["units"].get(param, "")
                label = f"{param} ({unit})" if unit else param
                if not tname:
                    label = f"{name.replace('_Table', '').replace('_', ' ')}: {label}"
                raw.append((label, epochs, list(values)))
    out: List[Tuple[str, Any, Any]] = []
    for label, epochs, values in raw:
        x = np.asarray(epochs, dtype=float)
        y = np.array([np.nan if v is None else v for v in values], dtype=float)
        if max_points and len(x) > max_points:
            x, y = _minmax_decimate(x, y, max_points)
        out.append((label, x, y))
    return out


def _theme_colours(dark: bool) -> Tuple[str, str, str]:
    """(foreground, background, grid) colours of a plot for the GUI theme."""
    return ("#cdd6f4", "#1e1e2e", "#45475a") if dark else ("#202020", "#ffffff", "#b0b0b0")


def _draw_series(fig: Any, series: List[Tuple[str, Any, Any]], title: str, ylabel: str,
                 dark: bool, interactive: bool = False) -> None:
    """Draw ``series`` on a fresh axes of ``fig`` with a local-time date axis."""
    import matplotlib.dates as mdates
    fg, bg, grid = _theme_colours(dark)
    fig.set_facecolor(bg)
    ax = fig.add_subplot(111, facecolor=bg)
    plotted = 0
    for index, (label, x, y) in enumerate(series):
        ok = np.isfinite(x)
        if not ok.any():
            continue
        local_ms = ((x[ok] + _utc_offsets(x[ok])) * 1000.0).astype("int64")
        times = local_ms.astype("datetime64[ms]")     # naive local wall clock
        few = len(times) <= 200
        # Only pass picker for the interactive window: Line2D treats any
        # non-callable picker (including None) as a pick radius and rejects None.
        extra = {"picker": True, "pickradius": 5} if interactive else {}
        ax.plot(times, y[ok], label=label, color=_PREVIEW_COLOURS[index % len(_PREVIEW_COLOURS)],
                linewidth=1.2, marker="o" if few else None, markersize=2.5, **extra)
        plotted += 1
    if not plotted:
        ax.text(0.5, 0.5, "No numeric data yet", ha="center", va="center",
                transform=ax.transAxes, color=fg, fontsize=12)
    locator = mdates.AutoDateLocator()
    ax.xaxis.set_major_locator(locator)
    ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(locator))
    ax.set_title(title, fontsize=9, color=fg)
    ax.set_xlabel("Local time", color=fg)
    ax.set_ylabel(ylabel, color=fg)
    ax.tick_params(colors=fg, labelsize=7)
    for spine in ax.spines.values():
        spine.set_color(grid)
    ax.grid(color=grid, alpha=0.5)
    if plotted:
        legend = ax.legend(fontsize=6, loc="upper left", bbox_to_anchor=(1.01, 1.0), borderaxespad=0,
                           facecolor=bg, edgecolor=grid)
        for text in legend.get_texts():
            text.set_color(fg)
    fig.tight_layout(rect=(0, 0, 0.80, 1))


def render_preview_png(ip: str, tname: Optional[str] = None, group: Optional[str] = None,
                       dark: bool = False, max_points: Optional[int] = None,
                       size: Tuple[float, float] = (10.0, 4.0), dpi: int = 100) -> bytes:
    """PNG bytes of one preview plot (Agg canvas; the figure is freed before return)."""
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    series = preview_series(ip, tname, group, max_points)
    title = f"{tname or 'Overlay: ' + str(group)}\n{ip}  ({PREVIEW_MAX_POINTS} pts max per series)"
    fig = Figure(figsize=size, dpi=dpi)
    canvas = FigureCanvasAgg(fig)
    try:
        _draw_series(fig, series, title, group or "Value", dark)
        buffer = io.BytesIO()
        canvas.print_png(buffer)
        return buffer.getvalue()
    finally:
        fig.clf()                 # drop axes and artists now, not at the next GC
        del canvas


def build_preview_pngs(_unused: Any = None) -> List[Dict[str, Any]]:
    """1.5.x API: render EVERY preview plot (the GUI now renders only the visible one)."""
    return [{"png_bytes": render_preview_png(ip, tname, group), "title": title,
             "ip": ip, "tname": tname, "group": group}
            for ip, tname, group, title in preview_plot_keys()]


def build_preview_figures(all_device_data: Any = None) -> List[Dict[str, Any]]:
    """1.5.x API shim: same as build_preview_pngs()."""
    return build_preview_pngs()


# %% ===========================================================================
#    GUI (PySide6 through qtpy)
# =============================================================================
_GUI_DARK_STYLE = """
    QMainWindow, QWidget, QDialog { background-color: #1e1e2e; color: #cdd6f4; }
    QGroupBox { background-color: #181825; border: 1px solid #45475a; border-radius: 6px;
                margin-top: 10px; padding: 8px; color: #89b4fa; font-weight: bold; }
    QGroupBox::title { subcontrol-origin: margin; left: 10px; }
    QCheckBox, QLabel { color: #cdd6f4; background-color: transparent; }
    QPushButton { background-color: #313244; color: #cdd6f4; border: 1px solid #45475a;
                  border-radius: 4px; padding: 4px 10px; }
    QPushButton:hover { background-color: #45475a; }
    QPushButton:pressed { background-color: #585b70; }
    QPushButton:disabled { color: #6c7086; }
    QSpinBox, QDoubleSpinBox, QLineEdit { background-color: #313244; color: #cdd6f4;
                  border: 1px solid #45475a; border-radius: 3px; padding: 2px; }
    QTabWidget::pane { border: 1px solid #45475a; }
    QTabBar::tab { background: #313244; color: #cdd6f4; padding: 6px 14px; }
    QTabBar::tab:selected { background: #45475a; color: #89b4fa; }
    QPlainTextEdit { background-color: #181825; color: #a6e3a1; border: 1px solid #45475a;
                     font-family: "Courier New", monospace; font-size: 9pt; }
    QScrollArea { border: none; }
    QScrollBar:vertical { background: #1e1e2e; width: 10px; }
    QScrollBar::handle:vertical { background: #45475a; border-radius: 5px; }
    QStatusBar { background: #181825; color: #a6adc8; }
    QMenuBar { background: #181825; color: #cdd6f4; }
    QMenuBar::item:selected { background: #313244; }
    QMenu { background: #1e1e2e; color: #cdd6f4; border: 1px solid #45475a; }
    QMenu::item:selected { background: #313244; }
"""
_GUI_LIGHT_STYLE = ""            # Qt default palette

#: Console colour per log level (dark and light themes share these).
_GUI_LEVEL_COLOUR = {logging.DEBUG: "#7f849c", logging.INFO: "#4c8bf5", logging.WARNING: "#e5a50a",
                     logging.ERROR: "#e64553", logging.CRITICAL: "#ff0000"}

#: The running main window (set by launch_gui; used by tests and diagnostics).
_GUI_WINDOW: Any = None


def launch_gui(initial_switches: Dict[str, Any], initial_figures: Optional[List[Any]] = None,
               run_event_loop: bool = True) -> Any:
    """Create the main window and (by default) run the Qt event loop.

    Threads
    -------
    * GUI thread: widgets only - status, tabs, rendering of the VISIBLE tab.
    * AcquisitionWorker (QThread): begin_run(), poll_cycle() every period
      (poll, store, persist, RAM flush), end_run(); also serves Save Now,
      Flush RAM Now and Open in Veusz while running, so those never race
      the poll for the store or the files.
    * TaskThread (QThread): the same requests when no acquisition is running.

    Parameters
    ----------
    initial_switches : dict
        Starting values (default_config() keys); the GUI edits a copy.
    initial_figures : list, optional
        Ignored (1.5.x signature).
    run_event_loop : bool
        False returns the window without running the loop (tests). True runs
        it and ends with sys.exit() like 1.5.x (main() catches SystemExit).
    """
    global _GUI_WINDOW
    os.environ.setdefault("QT_API", "pyside6")
    try:
        from qtpy import QtCore, QtGui, QtWidgets
        from qtpy.QtCore import QObject, Qt, QThread, QTimer, Signal
        from qtpy.QtWidgets import (QCheckBox, QDialog, QDoubleSpinBox, QFileDialog, QGridLayout,
                                    QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMainWindow,
                                    QMessageBox, QPlainTextEdit, QPushButton, QScrollArea,
                                    QSpinBox, QSplitter, QStatusBar, QTabWidget, QVBoxLayout,
                                    QWidget)
    except ImportError as exc:
        logger.critical("GUI dependencies missing: %s\nInstall: mamba install -c conda-forge "
                        "pyside6 qtpy matplotlib-base", exc)
        return None
    try:
        from qtpy.QtGui import QAction           # Qt 6 location
    except ImportError:                          # pragma: no cover - Qt 5 location
        from qtpy.QtWidgets import QAction

    # -------------------------------------------------------------------------
    # %%% Log bridge: any thread -> status console, through a queued signal
    # -------------------------------------------------------------------------
    class _LogBridge(QObject):
        """Carries formatted log records to the GUI thread (queued connection)."""
        record = Signal(str, int)

    class _GuiLogHandler(logging.Handler):
        """logging.Handler that emits into _LogBridge; safe from every thread."""

        def __init__(self, bridge: Any, level: int = logging.INFO) -> None:
            super().__init__(level)
            self.bridge = bridge
            self.setFormatter(logging.Formatter("%(asctime)s  %(levelname)-8s  %(message)s", "%H:%M:%S"))

        def emit(self, record: logging.LogRecord) -> None:
            try:
                self.bridge.record.emit(self.format(record), record.levelno)
            except Exception:                    # bridge already deleted at shutdown
                pass

    # -------------------------------------------------------------------------
    # %%% Worker threads
    # -------------------------------------------------------------------------
    class AcquisitionWorker(QThread):
        """Runs a whole acquisition run off the GUI thread (see launch_gui)."""
        cycle_done = Signal(object)          # poll_cycle() status dict
        request_done = Signal(str, object)   # ("save" | "flush" | "veusz", result)
        run_finished = Signal(object)        # end_run() status

        def __init__(self, cfg: Dict[str, Any], count: int = 0, parent: Any = None) -> None:
            super().__init__(parent)
            self.cfg = dict(cfg)
            self.count = int(count)
            self._stop = threading.Event()
            self._requests: "queue.Queue[str]" = queue.Queue()

        def request_stop(self) -> None:
            """Stop after the current page; final writes still happen."""
            self._stop.set()

        def submit(self, what: str) -> None:
            """Queue "save", "flush" or "veusz" to run between polls."""
            self._requests.put(what)

        def _serve_requests(self) -> None:
            while True:
                try:
                    what = self._requests.get_nowait()
                except queue.Empty:
                    return
                self.request_done.emit(what, _run_request(what, self.cfg))

        def run(self) -> None:  # noqa: D401 - QThread entry point
            final: Dict[str, Any] = {}
            consecutive, polls = 0, 0
            max_fails = int(self.cfg.get("headless_max_consec_fails", HEADLESS_MAX_CONSEC_FAILS))
            period = max(0.5, float(self.cfg.get("sample_period", SAMPLE_PERIOD_SEC)))
            try:
                begin_run(self.cfg)
                while not self._stop.is_set():
                    self._serve_requests()
                    started = time.monotonic()
                    status = poll_cycle(self.cfg, stop_event=self._stop)
                    polls += 1
                    if status["all_failed"] and not status["cancelled"]:
                        consecutive += 1
                        status["consecutive_failures"] = consecutive
                        if max_fails > 0 and consecutive >= max_fails:
                            status["auto_stop"] = True
                    else:
                        consecutive = 0
                    self.cycle_done.emit(status)
                    if status.get("auto_stop") or (self.count and polls >= self.count):
                        break
                    while not self._stop.is_set():
                        remaining = period - (time.monotonic() - started)
                        if remaining <= 0:
                            break
                        self._stop.wait(min(0.25, remaining))
                        self._serve_requests()
            except Exception as exc:
                logger.error("Acquisition worker error: %s", exc, exc_info=True)
            finally:
                try:
                    final = end_run(self.cfg)
                except Exception as exc:
                    logger.error("Final writes failed: %s", exc, exc_info=True)
                    final = {"error": str(exc)}
                self._serve_requests()
                self.run_finished.emit(final)

    class TaskThread(QThread):
        """One request (save / flush / veusz) when no acquisition is running."""
        done = Signal(str, object)

        def __init__(self, what: str, cfg: Dict[str, Any], parent: Any = None) -> None:
            super().__init__(parent)
            self.what, self.cfg = what, dict(cfg)

        def run(self) -> None:  # noqa: D401 - QThread entry point
            self.done.emit(self.what, _run_request(self.what, self.cfg))

    def _run_request(what: str, cfg: Dict[str, Any]) -> Dict[str, Any]:
        """Execute one GUI request in a worker thread; never raises."""
        try:
            if what == "save":
                return {"writers": persist_outputs(cfg, full=True, must_persist=True, reason="Save Now")}
            if what == "flush":
                return perform_ram_flush(cfg, "manual (Flush RAM Now)")
            if what == "veusz":
                return {"paths": open_veusz_preview(None, cfg.get("veusz_dir", VEUSZ_DIR))}
        except Exception as exc:
            logger.error("%s failed: %s", what, exc, exc_info=True)
            return {"error": str(exc)}
        return {"error": f"unknown request {what!r}"}

    # -------------------------------------------------------------------------
    # %%% Main window
    # -------------------------------------------------------------------------
    class MainWindow(QMainWindow):
        """AB Power Meter Monitor main window."""

        def __init__(self, switches: Dict[str, Any]) -> None:
            super().__init__()
            self.setWindowTitle(f"AB Power Meter Monitor {__version__} — NRAO / GBO")
            self.resize(1360, 860)
            self._switches = dict(default_config(), **dict(switches or {}))
            self._worker: Optional[AcquisitionWorker] = None
            self._tasks: Set[Any] = set()
            self._tabs: Dict[Tuple[str, Optional[str], Optional[str]], Dict[str, Any]] = {}
            self._current_key: Optional[Tuple[str, Optional[str], Optional[str]]] = None
            self._iplot_open: Set[Tuple[str, Optional[str], Optional[str]]] = set()
            self._dark_mode = False
            self._closing = False
            self._close_ok = False
            self._close_started = 0.0
            self._close_asked = False
            self._build_menu()
            self._build_central()
            self.setStatusBar(QStatusBar())
            self.statusBar().showMessage("Ready - choose options and press Poll Now or Start Auto.")
            # Log records from every thread arrive through the bridge (queued).
            self._bridge = _LogBridge()
            self._bridge.record.connect(self._on_log_record)
            self._log_handler = _GuiLogHandler(self._bridge, logging.INFO)
            logging.getLogger("ABMonitor").addHandler(self._log_handler)
            # Memory line refresh between polls (cheap: no store walk).
            self._status_timer = QTimer(self)
            self._status_timer.timeout.connect(self._refresh_memory_line)
            self._status_timer.start(5000)
            if QtWidgets.QApplication.instance().palette().window().color().lightness() < 128:
                self._apply_theme(True)

        # ---------------------------------------------------------------
        # %%% Menus and layout
        # ---------------------------------------------------------------
        def _build_menu(self) -> None:
            bar = self.menuBar()
            file_menu = bar.addMenu("&File")
            for text, slot in (("&Save Now", self._save_now), ("&Flush RAM Now", self._flush_now),
                               ("Open &output folder", self._open_output_folder), (None, None),
                               ("&Quit", self.close)):
                if text is None:
                    file_menu.addSeparator()
                    continue
                action = QAction(text, self)
                action.triggered.connect(slot)
                file_menu.addAction(action)
            view_menu = bar.addMenu("&View")
            self._act_theme = QAction("Switch to &Dark Theme", self)
            self._act_theme.triggered.connect(lambda: self._apply_theme(not self._dark_mode))
            view_menu.addAction(self._act_theme)
            help_menu = bar.addMenu("&Help")
            about = QAction("&About", self)
            about.triggered.connect(self._show_about)
            help_menu.addAction(about)

        def _build_central(self) -> None:
            splitter = QSplitter(Qt.Orientation.Horizontal)
            controls = QWidget()
            column = QVBoxLayout(controls)
            column.setSpacing(6)
            column.addWidget(self._build_device_group())
            column.addWidget(self._build_output_group())
            column.addWidget(self._build_memory_group())
            column.addWidget(self._build_buttons())
            column.addStretch(1)
            scroller = QScrollArea()
            scroller.setWidget(controls)
            scroller.setWidgetResizable(True)
            scroller.setMinimumWidth(380)
            left = QWidget()
            left_layout = QVBoxLayout(left)
            left_layout.setContentsMargins(0, 0, 0, 0)
            left_layout.addWidget(scroller, 3)
            left_layout.addWidget(QLabel("Status console (newest %d lines)" % GUI_LOG_MAX_LINES))
            self._console = QPlainTextEdit()
            self._console.setReadOnly(True)
            self._console.setMaximumBlockCount(int(GUI_LOG_MAX_LINES))   # bounded memory [R2]
            left_layout.addWidget(self._console, 2)
            splitter.addWidget(left)
            self._tab_widget = QTabWidget()
            self._placeholder = QLabel("No data yet - press Poll Now or Start Auto.")
            self._placeholder.setAlignment(Qt.AlignmentFlag.AlignCenter)
            self._tab_widget.addTab(self._placeholder, "Waiting …")
            self._tab_widget.currentChanged.connect(self._on_tab_changed)
            splitter.addWidget(self._tab_widget)
            splitter.setStretchFactor(1, 1)
            self.setCentralWidget(splitter)

        def _build_device_group(self) -> Any:
            group = QGroupBox("Devices and timing")
            grid = QGridLayout(group)
            grid.addWidget(QLabel("IP list (comma-separated):"), 0, 0, 1, 2)
            self._le_ips = QLineEdit(", ".join(self._switches.get("ip_list", IP_LIST)))
            self._le_ips.setToolTip("Full IP addresses; duplicates are removed.")
            grid.addWidget(self._le_ips, 1, 0, 1, 2)
            grid.addWidget(QLabel("Sample period (s):"), 2, 0)
            self._spin_period = QDoubleSpinBox()
            self._spin_period.setRange(5.0, 3600.0)
            self._spin_period.setDecimals(1)
            self._spin_period.setSingleStep(5.0)
            self._spin_period.setValue(float(self._switches.get("sample_period", SAMPLE_PERIOD_SEC)))
            grid.addWidget(self._spin_period, 2, 1)
            grid.addWidget(QLabel("HTTP retries / delay (s):"), 3, 0)
            row = QHBoxLayout()
            self._spin_retries = QSpinBox()
            self._spin_retries.setRange(1, 10)
            self._spin_retries.setValue(int(self._switches.get("http_retry_count", HTTP_RETRY_COUNT)))
            self._spin_delay = QDoubleSpinBox()
            self._spin_delay.setRange(0.5, 30.0)
            self._spin_delay.setSingleStep(0.5)
            self._spin_delay.setValue(float(self._switches.get("http_retry_delay_sec", HTTP_RETRY_DELAY_SEC)))
            row.addWidget(self._spin_retries)
            row.addWidget(self._spin_delay)
            grid.addLayout(row, 3, 1)
            grid.addWidget(QLabel("Stop after N all-failed polls:"), 4, 0)
            self._spin_fails = QSpinBox()
            self._spin_fails.setRange(0, 1000)
            self._spin_fails.setValue(int(self._switches.get("headless_max_consec_fails",
                                                             HEADLESS_MAX_CONSEC_FAILS)))
            self._spin_fails.setToolTip("0 = never stop automatically.")
            grid.addWidget(self._spin_fails, 4, 1)
            return group

        def _build_output_group(self) -> Any:
            group = QGroupBox("Outputs")
            column = QVBoxLayout(group)
            self._cb = {}
            for key, text in (("enable_csv", "CSV (appended every poll)"),
                              ("enable_log_append", "Markdown log tables (.md)"),
                              ("enable_veusz", "Veusz .vszh5 (native, appended)"),
                              ("veusz_append_every_cycle", "   └ append every poll (else flush/stop)"),
                              ("enable_fits", "FITS"), ("enable_xlsx", "Excel XLSX"),
                              ("enable_log_file", "ab_monitor.log file"),
                              ("append_files", "Append to existing files (else start fresh)")):
                box = QCheckBox(text)
                box.setChecked(bool(self._switches.get(key, 0)))
                column.addWidget(box)
                self._cb[key] = box
            self._cb["veusz_append_every_cycle"].setEnabled(self._cb["enable_veusz"].isChecked())
            self._cb["enable_veusz"].toggled.connect(self._cb["veusz_append_every_cycle"].setEnabled)
            row = QHBoxLayout()
            row.addWidget(QLabel("FITS/XLSX every N polls:"))
            self._spin_full = QSpinBox()
            self._spin_full.setRange(0, 10000)
            self._spin_full.setValue(int(self._switches.get("full_write_every_n_cycles",
                                                            FULL_WRITE_EVERY_N_CYCLES)))
            self._spin_full.setToolTip("0 = only at RAM flush, Save Now and Stop.")
            row.addWidget(self._spin_full)
            column.addLayout(row)
            column.addWidget(QLabel("Output root folder:"))
            row = QHBoxLayout()
            self._le_out = QLineEdit(self._switches.get("output_base_dir", OUTPUT_BASE_DIR))
            browse = QPushButton("Browse …")
            browse.clicked.connect(self._choose_out_dir)
            row.addWidget(self._le_out)
            row.addWidget(browse)
            column.addLayout(row)
            return group

        def _build_memory_group(self) -> Any:
            group = QGroupBox("Memory protection")
            grid = QGridLayout(group)
            self._cb_flush = QCheckBox("Automatic RAM flush")
            self._cb_flush.setChecked(bool(self._switches.get("enable_ram_flush", ENABLE_RAM_FLUSH)))
            grid.addWidget(self._cb_flush, 0, 0, 1, 2)
            self._spin_pct = QSpinBox()
            self._spin_pct.setRange(10, 99)
            self._spin_pct.setSuffix(" %")
            self._spin_pct.setValue(int(self._switches.get("mem_ram_pct_limit", MEM_RAM_PCT_LIMIT)))
            self._spin_store = QSpinBox()
            self._spin_store.setRange(1, 65536)
            self._spin_store.setSuffix(" MB")
            self._spin_store.setValue(int(self._switches.get("mem_flush_threshold_mb", MEM_FLUSH_THRESHOLD_MB)))
            self._spin_rss = QSpinBox()
            self._spin_rss.setRange(0, 262144)
            self._spin_rss.setSuffix(" MB")
            self._spin_rss.setValue(int(self._switches.get("mem_proc_rss_limit_mb", MEM_PROC_RSS_LIMIT_MB)))
            self._spin_rss.setToolTip("Flush when this program's RSS exceeds this (0 = off).")
            self._spin_keep = QSpinBox()
            self._spin_keep.setRange(0, 100000)
            self._spin_keep.setValue(int(self._switches.get("gui_keep_samples_after_flush",
                                                            GUI_KEEP_SAMPLES_AFTER_FLUSH)))
            self._spin_keep.setToolTip("Newest rows per table kept in memory after a flush "
                                       "(preview continuity). They are already on disk.")
            for row_number, (text, widget) in enumerate((("System RAM limit:", self._spin_pct),
                                                        ("Store limit:", self._spin_store),
                                                        ("Process RSS limit:", self._spin_rss),
                                                        ("Keep after flush (rows):", self._spin_keep)), 1):
                grid.addWidget(QLabel(text), row_number, 0)
                grid.addWidget(widget, row_number, 1)
            self._cb_spill = QCheckBox("Spill to pickle when no output is enabled")
            self._cb_spill.setChecked(bool(self._switches.get("ram_flush_spill_pickle", RAM_FLUSH_SPILL_PICKLE)))
            grid.addWidget(self._cb_spill, 5, 0, 1, 2)
            # 1.6.1: give the freed memory back to the OS after every flush.
            self._cb_return_os = QCheckBox("Return freed memory to the OS after a flush")
            self._cb_return_os.setChecked(bool(self._switches.get("mem_return_to_os", MEM_RETURN_TO_OS)))
            self._cb_return_os.setToolTip("Linux: malloc_trim. Windows: decommit free heap memory "
                                          "(HeapSetInformation HeapOptimizeResources).")
            grid.addWidget(self._cb_return_os, 6, 0, 1, 2)
            self._cb_trim_ws = QCheckBox("   └ trim the working set (Windows)")
            self._cb_trim_ws.setChecked(bool(self._switches.get("mem_trim_working_set", MEM_TRIM_WORKING_SET)))
            self._cb_trim_ws.setToolTip("EmptyWorkingSet: the physical RAM goes back to Windows at once and "
                                        "Task Manager's 'Memory' drops; used pages fault back in on access.")
            self._cb_trim_ws.setEnabled(sys.platform == "win32" and self._cb_return_os.isChecked())
            self._cb_return_os.toggled.connect(
                lambda on: self._cb_trim_ws.setEnabled(sys.platform == "win32" and on))
            grid.addWidget(self._cb_trim_ws, 7, 0, 1, 2)
            self._lbl_memory = QLabel("")
            self._lbl_memory.setWordWrap(True)
            grid.addWidget(self._lbl_memory, 8, 0, 1, 2)
            return group

        def _build_buttons(self) -> Any:
            box = QWidget()
            grid = QGridLayout(box)
            self._btn_poll = QPushButton("Poll Now")
            self._btn_start = QPushButton("Start Auto")
            self._btn_stop = QPushButton("Stop")
            self._btn_save = QPushButton("Save Now")
            self._btn_flush = QPushButton("Flush RAM Now")
            self._btn_veusz = QPushButton("Open in Veusz")
            self._btn_stop.setEnabled(False)
            self._btn_save.setToolTip("Write every enabled output now (including FITS/XLSX).")
            self._btn_flush.setToolTip("Write every enabled output, then free the in-memory rows.")
            self._btn_veusz.setToolTip("Bring the .vszh5 documents up to date and open them in Veusz.")
            for slot, button in ((self._poll_now, self._btn_poll), (self._start, self._btn_start),
                                 (self._stop, self._btn_stop), (self._save_now, self._btn_save),
                                 (self._flush_now, self._btn_flush), (self._open_veusz, self._btn_veusz)):
                button.clicked.connect(slot)
            for index, button in enumerate((self._btn_poll, self._btn_start, self._btn_stop,
                                            self._btn_save, self._btn_flush, self._btn_veusz)):
                grid.addWidget(button, index // 3, index % 3)
            return box

        # ---------------------------------------------------------------
        # %%% Configuration from the widgets
        # ---------------------------------------------------------------
        def _runtime_config(self) -> Dict[str, Any]:
            """default_config() updated with the current widget values."""
            cfg = default_config()
            base = self._le_out.text().strip() or OUTPUT_BASE_DIR
            ips = dedupe_ips(self._le_ips.text().split(",")) or dedupe_ips(IP_LIST)
            cfg.update({
                "ip_list": ips,
                "sample_period": float(self._spin_period.value()),
                "http_retry_count": int(self._spin_retries.value()),
                "http_retry_delay_sec": float(self._spin_delay.value()),
                "headless_max_consec_fails": int(self._spin_fails.value()),
                "full_write_every_n_cycles": int(self._spin_full.value()),
                "enable_ram_flush": int(self._cb_flush.isChecked()),
                "mem_ram_pct_limit": float(self._spin_pct.value()),
                "mem_flush_threshold_mb": float(self._spin_store.value()),
                "mem_proc_rss_limit_mb": float(self._spin_rss.value()),
                "ram_flush_keep_samples": int(self._spin_keep.value()),
                "ram_flush_spill_pickle": int(self._cb_spill.isChecked()),
                "mem_return_to_os": int(self._cb_return_os.isChecked()),
                "mem_trim_working_set": int(self._cb_trim_ws.isChecked()),
                "output_base_dir": base,
                "fits_dir": os.path.join(base, "fits"), "csv_dir": os.path.join(base, "csv"),
                "xlsx_dir": os.path.join(base, "xlsx"), "veusz_dir": os.path.join(base, "veusz"),
                "log_dir": os.path.join(base, "logs"), "spill_dir": os.path.join(base, "spill"),
            })
            for key, box in self._cb.items():
                cfg[key] = int(box.isChecked())
            return cfg

        def _choose_out_dir(self) -> None:
            path = QFileDialog.getExistingDirectory(self, "Output root folder", self._le_out.text())
            if path:
                self._le_out.setText(path)

        # ---------------------------------------------------------------
        # %%% Acquisition control
        # ---------------------------------------------------------------
        def _busy(self) -> bool:
            worker_running = self._worker is not None and self._worker.isRunning()
            return worker_running or any(t.isRunning() for t in list(self._tasks))

        def _set_running(self, running: bool) -> None:
            for button in (self._btn_poll, self._btn_start):
                button.setEnabled(not running and not self._closing)
            self._btn_stop.setEnabled(running and not self._closing)

        def _launch_worker(self, count: int) -> None:
            if self._worker is not None and self._worker.isRunning():
                return
            cfg = self._runtime_config()
            self._worker = AcquisitionWorker(cfg, count=count, parent=self)
            self._worker.cycle_done.connect(self._on_cycle_done)
            self._worker.request_done.connect(self._on_request_done)
            self._worker.run_finished.connect(self._on_run_finished)
            self._set_running(True)
            self._worker.start()
            self.statusBar().showMessage("Polling …" if count == 1 else "Auto-polling started …")

        def _poll_now(self) -> None:
            self._launch_worker(count=1)

        def _start(self) -> None:
            self._launch_worker(count=0)

        def _stop(self) -> None:
            if self._worker is not None and self._worker.isRunning():
                self._worker.request_stop()
                self._btn_stop.setEnabled(False)
                self.statusBar().showMessage("Stopping - finishing the final writes …")

        def _request(self, what: str) -> None:
            """Save / flush / Veusz: in the worker if running, else in a task thread."""
            if self._worker is not None and self._worker.isRunning():
                self._worker.submit(what)
                return
            task = TaskThread(what, self._runtime_config(), parent=self)
            task.done.connect(self._on_request_done)
            task.finished.connect(lambda _t=task: self._tasks.discard(_t))
            self._tasks.add(task)
            task.start()

        def _save_now(self) -> None:
            self._request("save")

        def _flush_now(self) -> None:
            self._request("flush")

        def _open_veusz(self) -> None:
            self._request("veusz")

        # ---------------------------------------------------------------
        # %%% Worker results (GUI thread)
        # ---------------------------------------------------------------
        def _on_cycle_done(self, status: Dict[str, Any]) -> None:
            n_ok, n_failed = len(status.get("ok_ips", [])), len(status.get("failed_ips", []))
            text = (f"Poll {status.get('cycle')}: {n_ok}/{n_ok + n_failed} meter(s) in "
                    f"{status.get('poll_seconds', 0):.1f} s")
            if status.get("consecutive_failures"):
                text += f" - all failed ({status['consecutive_failures']} in a row)"
            if status.get("flush"):
                flush = status["flush"]
                self._log(f"[RAM flush] {flush['reason']}: {flush['rows_removed']} row(s) freed, "
                          f"store {flush['store_mb_before']:.2f} -> {flush['store_mb_after']:.2f} MB; "
                          f"{flush.get('os_release_text', '')}", logging.WARNING)
            if status.get("auto_stop"):
                self._log("[Auto-stop] consecutive all-device failures reached the limit.",
                          logging.ERROR)
            self.statusBar().showMessage(text)
            self._refresh_memory_line()
            self._sync_tabs()

        def _on_request_done(self, what: str, result: Dict[str, Any]) -> None:
            if "error" in (result or {}):
                self._log(f"{what}: ERROR {result['error']}", logging.ERROR)
            elif what == "flush":
                self._log(f"Flush RAM Now: {result.get('rows_removed', 0)} row(s) freed, "
                          f"RSS {result.get('rss_mb_before', 0):.0f} -> {result.get('rss_mb_after', 0):.0f} MB; "
                          f"{result.get('os_release_text', '')}")
                self._sync_tabs()
            elif what == "save":
                self._log(f"Save Now: {result.get('writers') or 'no file output enabled'}")
            elif what == "veusz":
                paths = result.get("paths") or []
                self._log("Veusz documents: " + (", ".join(paths) if paths else
                                                 "none yet (enable Veusz output and poll first)"))
            self._refresh_memory_line()

        def _on_run_finished(self, final: Dict[str, Any]) -> None:
            self._set_running(False)
            if final:
                self._log(f"Run finished - final writes: {final}")
            self.statusBar().showMessage("Stopped.")
            self._refresh_memory_line()

        # ---------------------------------------------------------------
        # %%% Preview tabs: created once per plot, updated in place
        # ---------------------------------------------------------------
        def _sync_tabs(self) -> None:
            """Add/remove tabs to match the plots with data; re-render the visible one.

            1.5.x called QTabWidget.clear() and rebuilt every tab each poll;
            clear() does not delete the page widgets [R1], so each poll leaked
            dozens of widgets and their pixmaps. Here a removed tab's widget is
            deleteLater()'d and kept tabs are reused.
            """
            keys = preview_plot_keys(self._runtime_ips())
            wanted = [(ip, tname, group) for ip, tname, group, _title in keys]
            titles = {(ip, tname, group): title for ip, tname, group, title in keys}
            for key in [k for k in self._tabs if k not in titles]:
                widget = self._tabs.pop(key)["widget"]
                self._tab_widget.removeTab(self._tab_widget.indexOf(widget))
                widget.deleteLater()
            if wanted and self._tab_widget.indexOf(self._placeholder) >= 0:
                self._tab_widget.removeTab(self._tab_widget.indexOf(self._placeholder))
            for position, key in enumerate(wanted):
                if key not in self._tabs:
                    self._tabs[key] = self._make_tab(key, titles[key])
                    self._tab_widget.insertTab(position, self._tabs[key]["widget"], titles[key][:34])
                    self._tab_widget.setTabToolTip(position, titles[key])
            if not wanted and self._tab_widget.indexOf(self._placeholder) < 0:
                self._tab_widget.addTab(self._placeholder, "Waiting …")
            for info in self._tabs.values():
                info["dirty"] = True
            self._render_current()

        def _runtime_ips(self) -> List[str]:
            return dedupe_ips(self._le_ips.text().split(",")) or dedupe_ips(IP_LIST)

        def _make_tab(self, key: Tuple[str, Optional[str], Optional[str]], title: str) -> Dict[str, Any]:
            page = QWidget()
            layout = QVBoxLayout(page)
            layout.setContentsMargins(0, 0, 0, 0)
            layout.setSpacing(0)
            label = QLabel("rendering …")
            label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            scroller = QScrollArea()
            scroller.setWidget(label)
            scroller.setWidgetResizable(True)
            button = QPushButton("▶  Open interactive plot (zoom / pan / pick)")
            button.setMinimumHeight(32)
            button.clicked.connect(lambda _checked=False, _key=key, _t=title: self._open_interactive(_key, _t))
            layout.addWidget(scroller, 1)
            layout.addWidget(button, 0)
            return {"widget": page, "label": label, "dirty": True, "has_pixmap": False}

        def _key_of(self, widget: Any) -> Optional[Tuple[str, Optional[str], Optional[str]]]:
            return next((k for k, info in self._tabs.items() if info["widget"] is widget), None)

        def _on_tab_changed(self, _index: int) -> None:
            # Free the pixmap of the tab just left: only one large image lives at a time.
            previous = self._tabs.get(self._current_key) if self._current_key else None
            if previous is not None:
                previous["label"].clear()
                previous["has_pixmap"] = False
                previous["dirty"] = True
            self._render_current()

        def _render_current(self) -> None:
            key = self._key_of(self._tab_widget.currentWidget())
            self._current_key = key
            if key is None:
                return
            info = self._tabs[key]
            if not info["dirty"] and info["has_pixmap"]:
                return
            try:
                png = render_preview_png(key[0], key[1], key[2], dark=self._dark_mode)
            except Exception as exc:
                logger.error("Preview rendering failed: %s", exc)
                return
            pixmap = QtGui.QPixmap()
            pixmap.loadFromData(png, "PNG")
            info["label"].setPixmap(pixmap)
            info["dirty"], info["has_pixmap"] = False, True

        def _open_interactive(self, key: Tuple[str, Optional[str], Optional[str]], title: str) -> None:
            """Frozen interactive matplotlib window for one plot (non-modal)."""
            if key in self._iplot_open:
                self._log(f"An interactive window for '{title}' is already open.")
                return
            try:
                from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as Canvas
                from matplotlib.backends.backend_qtagg import NavigationToolbar2QT as Toolbar
                from matplotlib.figure import Figure
            except ImportError as exc:
                self._log(f"Interactive plot unavailable (matplotlib Qt backend): {exc}", logging.ERROR)
                return
            series = preview_series(key[0], key[1], key[2], max_points=50000)
            figure = Figure(figsize=(13, 5))
            _draw_series(figure, series, f"{title}  [snapshot {_local_text(time.time())}]",
                         key[2] or "Value", self._dark_mode, interactive=True)
            dialog = QDialog(self)
            dialog.setWindowTitle(f"Interactive — {title}")
            dialog.resize(1150, 540)
            dialog.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose, True)
            canvas = Canvas(figure)
            layout = QVBoxLayout(dialog)
            layout.setContentsMargins(4, 4, 4, 4)
            layout.addWidget(Toolbar(canvas, dialog))
            layout.addWidget(canvas, 1)
            self._iplot_open.add(key)

            def _closed(_result: int, _key: Any = key, _figure: Any = figure) -> None:
                self._iplot_open.discard(_key)
                _figure.clf()                    # release the Agg buffer now
            dialog.finished.connect(_closed)
            dialog.show()

        # ---------------------------------------------------------------
        # %%% Status console, memory line, theme
        # ---------------------------------------------------------------
        def _on_log_record(self, text: str, level: int) -> None:
            self._log(text, level, already_formatted=True)

        def _log(self, text: str, level: int = logging.INFO, already_formatted: bool = False) -> None:
            if not already_formatted:
                text = f"{datetime.datetime.now():%H:%M:%S}  {text}"
            safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            colour = _GUI_LEVEL_COLOUR.get(level, "#4c8bf5")
            self._console.appendHtml(f'<span style="color:{colour};">{safe}</span>')

        def _refresh_memory_line(self) -> None:
            status = run_status()
            self._lbl_memory.setText(
                f"Store {status['store_mb']:.2f} MB ({status['values']:,} values, {status['rows']:,} rows) | "
                f"RSS {status['rss_mb']:.0f} MB" +
                (f" (commit {status['commit_mb']:.0f} MB)" if status.get("commit_mb") else "") +
                f" | system RAM {status['ram_pct']:.1f} % | "
                f"flushes {status['flushes']}" + ("" if psutil is not None else " | psutil missing!"))

        def _apply_theme(self, dark: bool) -> None:
            self._dark_mode = bool(dark)
            QtWidgets.QApplication.instance().setStyleSheet(_GUI_DARK_STYLE if dark else _GUI_LIGHT_STYLE)
            self._act_theme.setText("Switch to &Light Theme" if dark else "Switch to &Dark Theme")
            for info in self._tabs.values():
                info["dirty"] = True
            self._render_current()

        def _open_output_folder(self) -> None:
            folder = self._le_out.text().strip() or OUTPUT_BASE_DIR
            os.makedirs(folder, exist_ok=True)
            QtGui.QDesktopServices.openUrl(QtCore.QUrl.fromLocalFile(folder))

        def _show_about(self) -> None:
            QMessageBox.information(
                self, "About AB Power Meter Monitor",
                f"Allen-Bradley Site Power Meter Monitor {__version__}\nNRAO / Green Bank Observatory\n\n"
                "Polls AB 1403 meters (pages 0-10) and writes CSV, Markdown, Veusz .vszh5 (native),\n"
                "FITS and XLSX with RAM-flush memory protection.\n\n"
                f"Author: W. Wallace\nPython {sys.version.split()[0]}  |  Qt binding via qtpy")

        # ---------------------------------------------------------------
        # %%% Closing: final writes always complete first
        # ---------------------------------------------------------------
        def closeEvent(self, event: Any) -> None:  # noqa: N802 - Qt override
            if not self._close_ok and self._busy():
                if not self._closing:
                    self._closing = True
                    self._close_started = time.monotonic()
                    self._set_running(True)
                    for button in (self._btn_save, self._btn_flush, self._btn_veusz, self._btn_stop):
                        button.setEnabled(False)
                    if self._worker is not None:
                        self._worker.request_stop()
                    self._log("Closing - finishing the final writes first …", logging.WARNING)
                    self.statusBar().showMessage("Closing - finishing the final writes …")
                    timer = QTimer(self)
                    timer.timeout.connect(self._try_finish_close)
                    timer.start(200)
                event.ignore()
                return
            for dialog in self.findChildren(QDialog):
                dialog.close()
            self._status_timer.stop()
            logging.getLogger("ABMonitor").removeHandler(self._log_handler)
            event.accept()

        def _try_finish_close(self) -> None:
            if not self._busy():
                self._close_ok = True
                self.close()
                return
            if not self._close_asked and time.monotonic() - self._close_started > 120.0:
                self._close_asked = True
                answer = QMessageBox.question(
                    self, "Final writes still running",
                    "The final file writes have not finished after 2 minutes.\n"
                    "Keep waiting (recommended)?  Choosing No quits now and may lose unsaved rows.")
                if answer == QMessageBox.StandardButton.No:
                    logger.critical("Forced quit while final writes were running.")
                    logging.shutdown()
                    os._exit(2)

    # -------------------------------------------------------------------------
    # %%% Application entry
    # -------------------------------------------------------------------------
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication(sys.argv)
    window = MainWindow(initial_switches)
    _GUI_WINDOW = window
    window.show()
    if not run_event_loop:
        return window
    exit_code = app.exec() if hasattr(app, "exec") else app.exec_()
    sys.exit(exit_code)


# %% ===========================================================================
#    HEADLESS CONSOLE OUTPUT HELPERS
# =============================================================================
def _get_named_dicts() -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Latest snapshot as {ip: {table: dict}} (copy), or the named dicts if no poll yet."""
    snapshot = all_device_data_snapshot()
    if snapshot:
        return snapshot
    return {"(no data)": {
        "Device_Configuration_Table": Device_Configuration_Table,
        "Communications_Configuration_Table": Communications_Configuration_Table,
        "Voltage_Current_Table": Voltage_Current_Table,
        "Real_Time_Power_Table": Real_Time_Power_Table,
        "Cumulative_Power_Table": Cumulative_Power_Table,
        "Demand_Data_Table": Demand_Data_Table,
        "Diagnostic_Table": Diagnostic_Table,
        "Voltage_Current_Snapshot_Log_Table": Voltage_Current_Snapshot_Log_Table,
        "Power_Snapshot_Log_Table": Power_Snapshot_Log_Table,
        "MinMax_Log_Table": MinMax_Log_Table,
        "Diagnostic_Table_Extended": Diagnostic_Table_Extended,
    }}


def _print_named_dicts(cycle: int, label: str = "current snapshot") -> None:
    """Print every device's 11 table dicts (latest poll) as JSON."""
    devices = _get_named_dicts()
    print("\n" + "=" * 72)
    print(f"  NAMED TABLE DICTS - cycle {cycle} ({label})")
    print(f"  Devices: {list(devices.keys())}")
    print("=" * 72)
    for ip, tables in devices.items():
        print(f"\n{'-' * 72}\n  Device: {ip}\n{'-' * 72}")
        for name, table in tables.items():
            print(f"\n  {name}:")
            if table:
                print(json.dumps({k: str(v) for k, v in table.items() if k != "_meta"}, indent=4))
            else:
                print("    (empty - no data fetched yet)")


def _print_cumulative_store(cycle: int, reason: str = "end of loop") -> None:
    """Print the whole in-memory store as aligned text tables."""
    snapshot = store_snapshot()
    print("\n" + "=" * 72)
    print(f"  CUMULATIVE TIME-SERIES STORE - cycle {cycle}, trigger: {reason}")
    print("=" * 72)
    if not snapshot:
        print("  (store is empty - no data accumulated yet)")
        return
    for ip, tables in snapshot.items():
        print(f"\n{'-' * 72}\n  Device: {ip}\n{'-' * 72}")
        for tname, tdata in tables.items():
            stamps, columns, units = tdata["timestamps_local"], tdata["columns"], tdata["units"]
            params = list(columns)
            print(f"\n  Table: {tname}  ({len(stamps)} sample(s))")
            if not stamps:
                continue
            labels = ["Timestamp_Local"] + [f"{p} ({units[p]})" if units.get(p) else p for p in params]
            rows = [[t] + ["" if columns[p][i] is None else str(columns[p][i]) for p in params]
                    for i, t in enumerate(stamps)]
            widths = [max([len(labels[j])] + [len(r[j]) for r in rows]) for j in range(len(labels))]
            print("  | " + " | ".join(l.ljust(widths[j]) for j, l in enumerate(labels)) + " |")
            print("  |" + "|".join("-" * (w + 2) for w in widths) + "|")
            for row in rows:
                print("  | " + " | ".join(c.ljust(widths[j]) for j, c in enumerate(row)) + " |")


# %% ===========================================================================
#    HEADLESS MODE
# =============================================================================
# Set by the signal handlers (or by a caller) to stop the headless loop; also
# passed to poll_all_devices(), so a stop cuts a running poll short.
_HEADLESS_STOP = threading.Event()
_ROOT_NULL_HANDLER = logging.NullHandler()   # added to the root logger at most once


def _install_signal_handlers() -> None:
    """SIGINT/SIGTERM (+ SIGBREAK on Windows) -> _HEADLESS_STOP; remove a stale stop file.

    SIGBREAK (1.6.2): Ctrl+Break in a Windows console, or CTRL_BREAK_EVENT sent
    by a supervisor, used to kill the process without the final writes.
    """
    def _handler(signum: int, _frame: Any) -> None:
        logger.info("Signal %s received - stopping after the current page.", signum)
        _HEADLESS_STOP.set()

    names = ("SIGINT", "SIGTERM", "SIGBREAK")       # SIGBREAK exists on Windows only
    try:
        for name in names:
            if hasattr(signal, name):
                signal.signal(getattr(signal, name), _handler)
    except (OSError, ValueError):
        pass                                 # not the main thread: cannot register
    _remove_stop_file("stale stop-signal file from a previous run")


def _remove_stop_file(why: str) -> None:
    """Delete the STOP_COLLECTION file if present (logged)."""
    path = stop_signal_file()
    try:
        if os.path.exists(path):
            os.remove(path)
            logger.info("Removed %s: %s", why, path)
    except OSError:
        pass


def _silence_console() -> None:
    """Remove console StreamHandlers of root and ABMonitor (files keep logging)."""
    for target in (logging.getLogger(), logging.getLogger("ABMonitor")):
        for handler in list(target.handlers):
            if isinstance(handler, logging.StreamHandler) and getattr(handler, "stream", None) in (
                    sys.__stdout__, sys.__stderr__, sys.stdout, sys.stderr):
                target.removeHandler(handler)
    root = logging.getLogger()
    if _ROOT_NULL_HANDLER not in root.handlers:      # 1.5.x added one per run
        root.addHandler(_ROOT_NULL_HANDLER)


def run_headless(cfg: Dict[str, Any]) -> None:
    """Headless loop: HEADLESS_LOOP_COUNT polls (0 = until stopped), then final writes.

    Stop with Ctrl-C / SIGTERM (when ENABLE_SIGNAL_HANDLERS = 1), the
    STOP_COLLECTION file (ab_stop.py), or _HEADLESS_STOP.set(). Every poll
    persists the incremental outputs and checks memory (poll_cycle); the end
    of the run writes everything that is not on disk yet (end_run).
    The HEADLESS_* switches are read from the module at call time.
    """
    if ENABLE_SIGNAL_HANDLERS:
        _install_signal_handlers()
    _HEADLESS_STOP.clear()
    loop_count = int(HEADLESS_LOOP_COUNT)
    infinite = loop_count == 0
    sample_sec = float(cfg.get("sample_period", SAMPLE_PERIOD_SEC))
    max_fails = int(cfg.get("headless_max_consec_fails", HEADLESS_MAX_CONSEC_FAILS))
    silent = bool(HEADLESS_SILENT)
    dicts_only = bool(HEADLESS_CONSOLE_DICTS_ONLY) and not silent
    print_each = bool(HEADLESS_PRINT_EACH_SAMPLE) and not silent
    print_cumulative = bool(HEADLESS_PRINT_CUMULATIVE) and not silent
    stop_file = stop_signal_file()

    devnull_out = devnull_err = None
    saved_streams = (sys.stdout, sys.stderr)
    if silent or dicts_only:
        _silence_console()
    if silent:
        devnull_out = open(os.devnull, "w", encoding="utf-8")
        devnull_err = open(os.devnull, "w", encoding="utf-8")
        sys.stdout, sys.stderr = devnull_out, devnull_err
    if not bool(ENABLE_RAM_FLUSH):
        logger.info("Automatic RAM flushing DISABLED (ENABLE_RAM_FLUSH=0) - "
                    "the store accumulates for the whole run.")
    logger.info("Headless loop starting - %s, period %.1f s, output %s",
                "until stopped" if infinite else f"{loop_count} poll(s)", sample_sec,
                cfg.get("output_base_dir", OUTPUT_BASE_DIR))

    cycle, consecutive = 0, 0
    try:
        begin_run(cfg)
        while True:
            if _HEADLESS_STOP.is_set():
                logger.info("Stop requested - leaving the loop after poll %d.", cycle)
                if print_cumulative:
                    _print_cumulative_store(cycle, reason="stop signal")
                break
            if os.path.exists(stop_file):
                logger.info("Stop-signal file detected (%s) - leaving the loop.", stop_file)
                _remove_stop_file("stop-signal file")
                if print_cumulative:
                    _print_cumulative_store(cycle, reason="ab_stop.py stop-signal file")
                break
            if not infinite and cycle >= loop_count:
                logger.info("Reached %d poll(s) - loop complete.", loop_count)
                break
            started = time.monotonic()
            status = poll_cycle(cfg, stop_event=_HEADLESS_STOP)
            cycle = status["cycle"]
            if status["all_failed"] and not status["cancelled"]:
                consecutive += 1
                logger.warning("Poll %d: ALL meters failed (%d in a row, limit %s).", cycle,
                               consecutive, max_fails if max_fails > 0 else "none")
                if max_fails > 0 and consecutive >= max_fails:
                    logger.error("Consecutive failure limit %d reached - exiting.", max_fails)
                    break
            else:
                if consecutive:
                    logger.info("Poll %d: connectivity restored after %d failure(s).", cycle, consecutive)
                consecutive = 0
            logger.info("Poll %d done in %.2f s - store %.2f MB, RSS %.0f MB, RAM %.1f %%",
                        cycle, status["poll_seconds"], status["store_mb"], status["rss_mb"],
                        status["ram_pct"])
            if status.get("flush") and print_cumulative:
                _print_cumulative_store(cycle, reason="memory flush")
            if print_each:
                _print_named_dicts(cycle, label="current snapshot")
            if not infinite and cycle >= loop_count:
                continue                       # last poll: no sleep, exit at loop top
            while not _HEADLESS_STOP.is_set() and not os.path.exists(stop_file):
                remaining = sample_sec - (time.monotonic() - started)
                if remaining <= 0:
                    break
                _HEADLESS_STOP.wait(min(0.5, remaining))
    except KeyboardInterrupt:
        logger.info("KeyboardInterrupt - stopping the headless loop.")
        if print_cumulative:
            _print_cumulative_store(cycle, reason="KeyboardInterrupt")
    except Exception as exc:                   # noqa: BLE001 - must reach the final writes
        logger.error("Headless loop aborted by an unexpected error: %s", exc, exc_info=True)
    try:
        final = end_run(cfg)
        if _any_data_output(cfg):
            logger.info("Final writes done (%s). Output folder: %s", final,
                        cfg.get("output_base_dir", OUTPUT_BASE_DIR))
        else:
            logger.info("No file output selected - nothing written.")
        _remove_stop_file("residual stop-signal file")
        if not silent and not dicts_only:
            _print_named_dicts(cycle, label="final snapshot")
        if print_cumulative:
            _print_cumulative_store(cycle, reason="end of loop")
    finally:
        if silent:
            sys.stdout, sys.stderr = saved_streams
            for handle in (devnull_out, devnull_err):
                try:
                    handle.close()
                except Exception:
                    pass
    gc.collect()


# %% ===========================================================================
#    MAIN ENTRY POINT
# =============================================================================
def main() -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Run the GUI (ENABLE_GUI = 1) or the headless loop; return TIME_SERIES_STORE.

    Return shape (unchanged since 1.x)::

        {"10.16.130.50": {"Real_Time_Power_Table": {
            "timestamps_local": ["2026-05-13 12:00:00", ...],
            "columns": {"L1 Voltage": [120.1, ...], ...},
            "units": {"L1 Voltage": "V", ...}}, ...}, ...}
    """
    global logger                              # noqa: PLW0603 - module logger re-bind
    cfg = default_config()
    _ensure_output_dirs(cfg)
    logger = _setup_logging(cfg["log_dir"], append=bool(cfg.get("append_files", 1)),
                            enable_log_file=bool(cfg.get("enable_log_file")))
    logger.info("AB Power Meter Monitor %s starting (Python %s).", __version__,
                sys.version.split()[0])
    logger.debug("Configuration: %s", json.dumps(cfg, indent=2, default=str))
    if ENABLE_GUI:
        try:
            launch_gui(initial_switches=cfg, initial_figures=[])
        except SystemExit:
            pass                               # launch_gui ends with sys.exit(app.exec())
        except Exception as exc:
            logger.critical("GUI launch failed: %s\n%s", exc, traceback.format_exc())
            logger.info("Falling back to headless mode.")
            run_headless(cfg)
    else:
        run_headless(cfg)
    return TIME_SERIES_STORE


def _parse_cli_and_apply(argv: Optional[Sequence[str]] = None) -> None:
    """Command-line overrides of the module switches (only when run as a script)."""
    import argparse
    parser = argparse.ArgumentParser(description="Allen-Bradley site power-meter monitor.",
                                     formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--headless", action="store_true", help="Force headless mode (ENABLE_GUI=0).")
    mode.add_argument("--gui", action="store_true", help="Force GUI mode (ENABLE_GUI=1).")
    parser.add_argument("--count", type=int, default=None, help="Headless polls (0 = until stopped).")
    parser.add_argument("--interval", type=float, default=None, help="Seconds between polls.")
    parser.add_argument("--ips", type=str, default=None, help="Comma-separated meter IPs.")
    parser.add_argument("--ram-pct", type=float, default=None, dest="ram_pct",
                        help="System RAM %% flush threshold (10-99).")
    parser.add_argument("--no-ram-flush", action="store_true", dest="no_ram_flush",
                        help="Disable automatic RAM flushing (ENABLE_RAM_FLUSH=0).")
    parser.add_argument("--output-dir", type=str, default=None, dest="output_dir",
                        help="Output root folder (OUTPUT_BASE_DIR).")
    parser.add_argument("--url-template", type=str, default=None, dest="url_template",
                        help='Page URL template, default "http://{ip}/{page}".')
    args = parser.parse_args(argv)

    global ENABLE_GUI, HEADLESS_LOOP_COUNT, SAMPLE_PERIOD_SEC, IP_LIST
    global MEM_RAM_PCT_LIMIT, ENABLE_RAM_FLUSH, OUTPUT_BASE_DIR, METER_URL_TEMPLATE
    if args.headless:
        ENABLE_GUI = 0
    if args.gui:
        ENABLE_GUI = 1
    if args.count is not None:
        HEADLESS_LOOP_COUNT = max(0, args.count)
    if args.interval is not None:
        SAMPLE_PERIOD_SEC = max(0.5, args.interval)
    if args.ips:
        IP_LIST = dedupe_ips(args.ips.split(","))
    if args.ram_pct is not None:
        MEM_RAM_PCT_LIMIT = max(10.0, min(99.0, float(args.ram_pct)))
    if args.no_ram_flush:
        ENABLE_RAM_FLUSH = 0
    if args.output_dir:
        OUTPUT_BASE_DIR = os.path.abspath(args.output_dir)
    if args.url_template:
        METER_URL_TEMPLATE = args.url_template


# %% ===========================================================================
#    SCRIPT ENTRY
# =============================================================================
if __name__ == "__main__":
    _parse_cli_and_apply()
    main()


# %% ===========================================================================
#    INTERACTIVE SCRATCH CELL (Spyder) - run cell by cell, not as a script
# =============================================================================
# import ab_power_meter_monitor as abm
# abm.IP_LIST = ["10.16.130.50"]; abm.ENABLE_CSV = 1; abm.ENABLE_VEUSZ = 1
# cfg = abm.begin_run()
# status = abm.poll_cycle(cfg)          # one poll + persist + memory check
# print(status["ok_ips"], status["store_mb"], abm.run_status())
# abm.perform_ram_flush(cfg, "manual test")
# abm.end_run(cfg)
